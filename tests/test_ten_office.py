"""10.0 documents: PDF edits keep text as text, Excel edits and formulas,
mail merge, the text checks, and the Documents page."""

from __future__ import annotations

import pytest


@pytest.fixture
def jarvis(base):
    from jarvis.assistant import Jarvis

    return Jarvis(voice_enabled=False)


@pytest.fixture
def pdf(tmp_path):
    from fpdf import FPDF

    doc = FPDF()
    for i in range(4):
        doc.add_page()
        doc.set_font("helvetica", size=30)
        doc.cell(0, 30, f"Page {i + 1}")
    path = tmp_path / "sample.pdf"
    doc.output(str(path))
    return path


def _texts(path):
    from pypdf import PdfReader

    return [p.extract_text() for p in PdfReader(str(path)).pages]


def test_page_lists():
    from jarvis.ten.office import OfficeError, page_list

    assert page_list("1-3,7", 8) == [1, 2, 3, 7]
    assert page_list("", 3) == [1, 2, 3]
    assert page_list("3-1", 3) == [3, 2, 1]
    with pytest.raises(OfficeError):
        page_list("9", 3)


def test_pdf_reorder_delete_rotate_number_watermark(jarvis, pdf):
    import re

    out = jarvis.process(f"/pdfpages {pdf} | 4,1").text
    path = re.search(r"(\S+\.pdf)", out.splitlines()[-1]).group(1)
    assert [t.strip() for t in _texts(path)] == ["Page 4", "Page 1"]
    out = jarvis.process(f"/pdfpages {pdf} | delete 2,3").text
    path = re.search(r"(\S+\.pdf)", out.splitlines()[-1]).group(1)
    assert len(_texts(path)) == 2
    out = jarvis.process(f"/pdfnumber {pdf} | bottom-center | of | 1").text
    path = re.search(r"(\S+\.pdf)", out.splitlines()[-1]).group(1)
    texts = _texts(path)
    assert "Page 1" in texts[0] and "1 / 4" in texts[0]
    out = jarvis.process(f"/pdfwatermark {pdf} | DRAFT").text
    path = re.search(r"(\S+\.pdf)", out.splitlines()[-1]).group(1)
    assert "DRAFT" in _texts(path)[2]
    out = jarvis.process(f"/pdfrotate {pdf} | 2 | 90").text
    from pypdf import PdfReader

    path = re.search(r"(\S+\.pdf)", out.splitlines()[-1]).group(1)
    assert PdfReader(path).pages[1].rotation == 90 and PdfReader(path).pages[0].rotation == 0


def test_pdf_signature_lands_on_the_last_page(base, pdf, tmp_path):
    from PIL import Image

    from jarvis.ten import office

    sig = tmp_path / "sig.png"
    Image.new("RGBA", (300, 100), (10, 30, 80, 255)).save(sig)
    target = office.pdf_sign(pdf, sig, 0, box=(0.6, 0.8))
    from pypdf import PdfReader

    pages = PdfReader(str(target)).pages
    assert len(pages[3].images) == 1 and len(pages[0].images) == 0


def test_pdf_compress_never_grows(base, pdf):
    from jarvis.ten import office

    target, before, after = office.pdf_compress(pdf)
    assert after <= before and target.exists()


def test_excel_from_ops_and_formulas(base, tmp_path):
    from openpyxl import load_workbook

    from jarvis.ten import office

    headers = ["Item", "Qty", "Price"]
    rows = [["pen", 3, 2.5], ["book", 1, 12.0], ["pen", 3, 2.5], ["bag", 2, 30.0]]
    h, r, notes, fmts, totals = office.apply_ops(headers, rows, [
        {"op": "delete_duplicates"}, {"op": "add_column", "name": "Total", "formula": "=B{r}*C{r}"},
        {"op": "sort", "by": "Price", "desc": True}, {"op": "total_row", "columns": ["Total"]}])
    assert h[-1] == "Total" and len(r) == 3 and r[0][0] == "bag" and totals == ["Total"]
    path = office.write_table(tmp_path / "x.xlsx", [{"name": "S", "headers": h, "rows": r, "totals": totals}])
    ws = load_workbook(path).active
    assert ws["D2"].value == "=B2*C2" and ws["D5"].value == "=SUM(D2:D4)"
    assert office.to_turkish_formula('=IFERROR(VLOOKUP(A2,B:C,2,FALSE),"yok")') == \
        '=EĞERHATA(DÜŞEYARA(A2;B:C;2;FALSE);"yok")'
    assert office.to_turkish_formula("=ROUND(A1*1.18,2)") == "=YUVARLA(A1*1,18;2)"


def test_excel_from_a_description(jarvis, monkeypatch):
    from jarvis import kit

    monkeypatch.setattr(kit, "ask_json", lambda brain, prompt, **kw: {"title": "Budget", "sheets": [
        {"name": "May", "headers": ["Category", "Planned", "Actual", "Left"],
         "rows": [["Rent", 1000, 1000, "=B{r}-C{r}"], ["Food", 400, 350, "=B{r}-C{r}"]], "totals": ["Planned"]}]})
    out = jarvis.process("/excel a monthly budget").text
    assert "Budget" in out and ".xlsx" in out


def test_mail_merge_fills_every_row(base, tmp_path):
    from docx import Document

    from jarvis.ten import office

    template = tmp_path / "t.docx"
    doc = Document()
    doc.add_paragraph("Dear {{Name}},")
    doc.add_paragraph("You owe {{ Amount }} TL.")
    doc.save(template)
    data = tmp_path / "d.csv"
    data.write_text("Name,Amount\nAyşe,100\nMehmet,250\n", encoding="utf-8")
    folder, combined, count = office.mail_merge(template, data)
    text = "\n".join(p.text for p in Document(str(combined)).paragraphs)
    assert count == 2 and "Dear Ayşe," in text and "You owe 250 TL." in text and "{{" not in text
    assert len(list(folder.glob("0*.docx"))) == 2


def test_text_checks():
    from jarvis.ten import office

    s = office.stats("One two three. Four five!\n\nSix.")
    assert s["words"] == 6 and s["sentences"] == 3 and s["paragraphs"] == 2
    easy = office.readability("The cat sat. The dog ran. We had fun.")
    hard = office.readability("Notwithstanding considerable organisational complexities, institutional "
                              "decision-making necessitates comprehensive interdisciplinary evaluation.")
    assert easy["score"] > hard["score"]
    assert office.readability("Bu çok güzel bir gün ve biz parka gittik.")["language"].startswith("Turkish")
    r = office.similarity("the quick brown fox jumps over the lazy dog every single day",
                          "yesterday the quick brown fox jumps over the lazy dog every single day")
    assert r["ratio"] > 0.8 and r["shared"]


def test_email_templates_fill_in(jarvis):
    out = jarvis.process("/emailtpl follow-up | name=Ayşe; topic=the budget; me=Ahmet").text
    assert "Hi Ayşe" in out and "the budget" in out and "Still to fill" in out


from tests.test_gui import _display_available, app  # noqa: E402,F401


@pytest.mark.skipif(not _display_available(), reason="no display")
def test_the_documents_page_views_a_spreadsheet(app, tmp_path):
    data = tmp_path / "v.csv"
    data.write_text("Name,Score\nAli,10\nVeli,30\nAyşe,20\n", encoding="utf-8")
    app._show_tab("documents")
    app.update()
    page = app.pages["documents"]
    page.open_sheet(str(data))
    assert len(page.tree.get_children()) == 3
    page.sort(1)
    first = page.tree.item(page.tree.get_children()[0])["values"]
    assert first[0] == "Ali"
    page.filter.insert(0, "ay")
    page.fill_tree()
    assert len(page.tree.get_children()) == 1
    assert "total 60" in page.sheet_stats.cget("text")
