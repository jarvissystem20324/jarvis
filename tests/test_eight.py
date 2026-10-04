"""8.0: ninety-six features, Mistral and Cloudflare.

No network, no model, no system changes: the web, the model, PowerShell and
the registry writes are stand-ins, and anything that would change Windows
runs with permission denied.
"""

from __future__ import annotations

import base64
import json
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

from jarvis import (devtools, drawing, guard, i18n, i18n_more, intents, kit, life, live, makers, pctools, providers,
                    registry, reminders, social, study)

windows_only = pytest.mark.skipif(sys.platform != "win32", reason="Windows-only")


@pytest.fixture
def jarvis(base, allow, monkeypatch):
    from jarvis.assistant import Jarvis

    j = Jarvis(voice_enabled=False)
    monkeypatch.setattr(j.voice, "speak", lambda *a, **k: None)
    monkeypatch.setattr(j, "_write_setting", lambda name, value: monkeypatch.setenv(name, value))
    return j


@pytest.fixture
def offline(monkeypatch):
    def nope(*a, **k):
        raise kit.KitError("offline in tests")

    monkeypatch.setattr(kit, "get_json", nope)
    monkeypatch.setattr(kit, "get_text", nope)


# --- the registry ------------------------------------------------------------------------

def test_every_command_has_a_method_and_help_lists_the_groups(base):
    from jarvis.assistant import Jarvis

    for name, method in registry.COMMANDS.items():
        assert callable(getattr(Jarvis, method)), name
    text = registry.help_text()
    # Every group that has commands; 10.0 also names groups for pages still
    # being built (Media…), which have none yet and are left out of help.
    for group in {g for g, _, _ in registry.HELP}:
        assert group in text
    assert len(registry.COMMANDS) >= 96


def test_every_command_answers_without_arguments(base, deny, offline, monkeypatch):
    """Nothing crashes and nothing changes the PC when called bare."""
    from jarvis.assistant import Jarvis, JarvisResponse
    from jarvis import pc

    j = Jarvis(voice_enabled=False)
    monkeypatch.setattr(j.brain, "ask_once", lambda *a, **k: "{}")
    monkeypatch.setattr(pctools, "powershell", lambda *a, **k: "")
    monkeypatch.setattr(pctools, "temp_files", lambda *a, **k: ([], 0))
    monkeypatch.setattr(pctools, "recycle_bin_size", lambda: (0, 0))
    monkeypatch.setattr(pctools, "wifi_info", lambda: {})
    monkeypatch.setattr(guard, "suspicious_processes", lambda: [])
    monkeypatch.setattr(pc, "known_folder", lambda name: base)
    # 10.0's commands have their own PowerShell helpers. Left real, a bare
    # /wifidevices scanned GitHub's network and hung the run for 20 minutes,
    # and a bare /avscan started a real Defender quick scan on every run.
    from jarvis.ten import shield10, system

    for module in (system, shield10):
        monkeypatch.setattr(module, "powershell", lambda *a, **k: "")
        monkeypatch.setattr(module, "ps_json", lambda *a, **k: [])
    monkeypatch.setattr(shield10, "scan_network", lambda: {"error": "offline in tests"})
    monkeypatch.setattr(shield10, "MPCMD", base / "no-defender.exe")
    from jarvis import weather

    monkeypatch.setattr(weather, "home_city", lambda: "")      # /wear and /uv would look the city up
    for name in sorted(set(registry.COMMANDS)):
        if name in {"20q", "twentyquestions", "typing", "typingtest"}:
            continue      # these start a game, tested below
        result = j.eight_command(name, "")
        assert isinstance(result, (str, JarvisResponse)), name
    assert not any(r.kind == "dnd" for r in reminders.board.items)


def test_new_commands_are_protected_from_addons():
    from jarvis import addons

    assert "slides" in registry.COMMANDS and "slides" not in addons.RESERVED_COMMANDS   # checked live instead
    assert set(registry.COMMANDS).isdisjoint({"task", "tasks", "todo", "review", "convert", "pdf", "img"})


# --- plain words ----------------------------------------------------------------------------

@pytest.mark.parametrize("said,command", [
    ("make a presentation about renewable energy", "/slides renewable energy"),
    ("solve x^2 - 4 = 0", "/solve x^2 - 4 = 0"),
    ("what does ubiquitous mean", "/define ubiquitous"),
    ("add milk to my list", "/mylist add milk"),
    ("split 450 between 3 tip 10%", "/split 450 3 tip 10%"),
    ("flip a coin", "/roll coin"),
    ("prayer times in ankara", "/prayer ankara"),
    ("gold price", "/gold"),
    ("alert me when usd goes above 35", "/alert usd above 35"),
    ("turn on dark mode", "/darkmode on"),
    ("is this link safe https://paypa1.com/login", "/checklink https://paypa1.com/login"),
    ("check my password", "/passcheck"),
    ("remind me every day at 9 to take vitamins", "/every every day at 9 to take vitamins"),
    ("whatsapp ali: late 10 min", "/whatsapp ali: late 10 min"),
    ("play 20 questions", "/20q"),
    ("screen time this week", "/screentime week"),
])
def test_plain_words(said, command):
    assert intents.route(said)[0] == command


@pytest.mark.parametrize("said", ["the news", "I like stock photos", "what is the derivative of x^2", "open spotify"])
def test_not_eight_commands(said):
    routed = intents.route(said)
    assert routed is None or routed[0].split()[0] not in {"/stock", "/news", "/solve"}


# --- everyday ---------------------------------------------------------------------------------

def test_dates_in_plain_words():
    now = datetime(2026, 10, 1, 10, 0)           # a Thursday
    assert life.parse_when("dentist friday 14:00", now) == (datetime(2026, 10, 2, 14, 0), "dentist")
    assert life.parse_when("report by monday", now) == (datetime(2026, 10, 5, 9, 0), "report")
    assert life.parse_when("call mum tomorrow at 5pm", now) == (datetime(2026, 10, 2, 17, 0), "call mum")
    assert life.parse_when("pay rent 15.10", now) == (datetime(2026, 10, 15, 9, 0), "pay rent")
    assert life.parse_when("water plants in 3 days", now)[0] == now + timedelta(days=3)
    assert life.parse_when("just a thing", now) == (None, "just a thing")


def test_repeating_rules():
    now = datetime(2026, 10, 1, 10, 0)
    repeat, first, what, until = life.parse_repeat("every monday at 18 to call mum", now)
    assert (repeat, first, what) == ("weekly", datetime(2026, 10, 5, 18, 0), "call mum")
    repeat, first, what, _ = life.parse_repeat("every 2 hours to drink water", now)
    assert repeat == "every:7200" and first == now + timedelta(hours=2)
    repeat, first, _, until = life.parse_repeat("every day at 21 to stretch for 30 days", now)
    assert repeat == "daily" and first == datetime(2026, 10, 1, 21, 0) and until > first.timestamp()
    assert life.parse_repeat("every weekday at 9:30 to stand up", datetime(2026, 10, 2, 12, 0))[1] == datetime(2026, 10, 5, 9, 30)


def test_next_occurrence():
    friday = datetime(2026, 10, 2, 9, 0).timestamp()
    assert datetime.fromtimestamp(reminders.next_due(friday, "weekdays", friday)) == datetime(2026, 10, 5, 9, 0)
    assert datetime.fromtimestamp(reminders.next_due(friday, "daily", friday)) == datetime(2026, 10, 3, 9, 0)
    assert reminders.next_due(friday, "every:3600", friday) == friday + 3600
    assert datetime.fromtimestamp(reminders.next_due(datetime(2026, 1, 31, 9).timestamp(), "monthly", 0)).month == 2
    assert reminders.next_due(friday, "", friday) is None


def test_a_missed_repeating_reminder_is_reported_once_and_rebooked(base):
    old = time.time() - 3 * 86400
    board = reminders.Board()
    board.items = [reminders.Reminder(1, "meds", "vitamin D", old, old, repeat="daily")]
    board.save()
    board.load()
    assert [r.text for r in board.missed] == ["vitamin D"]
    assert len(board.items) == 1 and board.items[0].due > time.time()


def test_the_list(jarvis):
    jarvis.process("add milk to my list")
    jarvis.process("/mylist add report friday 14:00")
    listing = jarvis.process("what's on my list").text
    assert "milk" in listing and "report" in listing
    assert any(r.kind == "task" for r in reminders.board.items)
    assert "Done: report" in jarvis.process("/done report").text
    assert "all clear" in jarvis.process("/done 1").text
    assert "empty" in jarvis.process("/mylist").text


def test_every_and_meds(jarvis):
    assert "every day" in jarvis.process("remind me every day at 9 to take vitamins").text
    assert "every 8 hours" in jarvis.process("/meds add amoxicillin every 8 hours for 7 days").text
    kinds = sorted(r.kind for r in reminders.board.items)
    assert kinds == ["meds", "reminder"]
    assert all(r.repeat for r in reminders.board.items)


def test_money_maths():
    assert "150.00 each" in life.split_bill("450 3")
    assert "165.00 each" in life.split_bill("450 3 tip 10%")
    by_person = life.split_bill("ali 120, ayşe 80, shared 150")
    assert "ali: 195.00" in by_person and "ayşe: 155.00" in by_person
    payment, total, interest = life.loan(100_000, 12, 12)
    assert round(payment, 2) == 8884.88 and round(interest, 2) == 6618.55
    payment, *_ = life.loan(100_000, 1, 12, monthly_rate=True)
    assert round(payment, 2) == 8884.88


def test_sleep_cycles():
    wake = datetime(2026, 10, 2, 7, 0)
    assert [t.strftime("%H:%M") for t in life.sleep_times(wake)] == ["21:45", "23:15", "00:45", "02:15"]


def test_chance(jarvis):
    assert jarvis.process("/roll coin").text.split()[-1] in {"Heads", "Tails"}
    assert "=" in jarvis.process("/roll 3d6").text
    assert jarvis.process("/pick tea, coffee").text.split()[-1] in {"tea", "coffee"}
    assert 1 <= int(jarvis.process("/random 1-3").text.split()[-1]) <= 3


def test_calendar_and_ics(jarvis, base):
    reply = jarvis.process("/cal add dentist friday 14:00 @ Kadıköy").text
    assert "dentist" in reply and "Kadıköy" in reply
    exported = jarvis.process("/cal export").text.splitlines()[1].strip()
    text = Path(exported).read_text(encoding="utf-8")
    assert "SUMMARY:dentist" in text and "LOCATION:Kadıköy" in text
    events = life.parse_ics(text)
    assert events[0]["title"] == "dentist"


def test_weekly_rule_expands():
    raw = ("BEGIN:VCALENDAR\r\nBEGIN:VEVENT\r\nUID:x\r\nDTSTART:20260928T100000\r\nDTEND:20260928T110000\r\n"
           "SUMMARY:Lecture\r\nRRULE:FREQ=WEEKLY;BYDAY=MO,WE\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n")
    events = social.expand_ics([], raw, datetime(2026, 10, 1), datetime(2026, 10, 9))
    assert [datetime.fromtimestamp(e["start"]).strftime("%a %d") for e in events] == ["Mon 05", "Wed 07"]


# --- study ------------------------------------------------------------------------------------

@pytest.mark.parametrize("question,answer", [
    ("2x + 3 = 11", "x = 4"), ("x + y = 10, x - y = 2", "x = 6, y = 4"), ("factor x^2-9", "(x - 3)*(x + 3)"),
    ("integrate x^2 from 0 to 3", "= 9"), ("limit of sin(x)/x as x to 0", "= 1"),
])
def test_exact_algebra(question, answer):
    assert answer in study.solve_math(question)


@pytest.mark.parametrize("attack", ["__import__('os')", "x.__class__", "open(1)", "eval(1)", "exec x", "lambda: 1"])
def test_algebra_cannot_run_code(attack):
    with pytest.raises(Exception):
        study.solve_math(attack)


def test_spaced_repetition():
    card = {"front": "a", "back": "b", "due": 0, "interval": 0, "ease": 2.5, "reps": 0}
    today = date(2026, 10, 1)
    study.grade(card, 3, today)
    assert card["interval"] == 1
    study.grade(card, 3, today)
    assert card["interval"] == 6
    study.grade(card, 4, today)
    assert card["interval"] >= 15
    study.grade(card, 1, today)
    assert card["interval"] == 1 and card["reps"] == 0


def test_card_review_runs_in_the_chat(jarvis, monkeypatch):
    study.add_cards("Spanish", [("hola", "hello"), ("gracias", "thank you")])
    monkeypatch.setattr(jarvis.brain, "chat", lambda *a, **k: pytest.fail("went to the AI"))
    first = jarvis.process("review my flashcards").text
    assert "2 left" in first
    assert "Answer:" in jarvis.process("no idea").text
    assert "left" in jarvis.process("3").text
    jarvis.process("x")
    assert "Review done" in jarvis.process("4").text
    assert all(c["due"] > date.today().toordinal() for c in study.CARDS.load()["Spanish"])


def test_typing_score():
    wpm, accuracy = study.typing_score("hello world", "hello world", 6)
    assert round(wpm) == 22 and accuracy == 100


def test_chapters_split_on_headings():
    text = "\n".join(f"Chapter {i}\n" + ("words " * 200) for i in range(1, 4))
    assert [h for h, _ in study.chapters_of(text)] == ["Chapter 1", "Chapter 2", "Chapter 3"]


def test_dictionary_and_wiki(jarvis, monkeypatch):
    def fake(url, params=None, **k):
        if "sozluk" in url:
            return [{"anlamlarListe": [{"anlam": "Okunmak için yazılmış eser", "orneklerListe": []}]}]
        if "dictionaryapi" in url:
            return [{"word": "serendipity", "phonetic": "/ˌsɛrənˈdɪpɪti/", "meanings": [
                {"partOfSpeech": "noun", "definitions": [{"definition": "luck in finding good things"}], "synonyms": []}]}]
        if "opensearch" in str(params):
            return ["x", ["Alan Turing"]]
        return {"title": "Alan Turing", "extract": "English mathematician.", "description": "computer scientist"}

    monkeypatch.setattr(kit, "get_json", fake)
    assert "luck in finding" in jarvis.process("define serendipity").text
    assert "Okunmak" in jarvis.process("/define kitap tr").text
    assert "English mathematician" in jarvis.process("/wiki turing").text


def test_tutor_sets_instructions(jarvis):
    jarvis.process("/tutor spanish A2")
    assert "Spanish" in jarvis.brain.instructions and "A2" in jarvis.brain.instructions
    jarvis.process("/tutor off")
    assert jarvis.brain.instructions == ""


# --- live information ----------------------------------------------------------------------------

def test_gold_maths(monkeypatch):
    from jarvis import calc

    monkeypatch.setattr(kit, "get_json", lambda url, *a, **k: {"price": 3110.34768 if "XAU" in url else 31.1034768})
    monkeypatch.setattr(calc, "rates", lambda base: ({"TRY": 40.0}, "test"))
    prices = live.gold_prices()
    assert round(prices["gram_try"], 2) == 4000.0 and round(prices["silver_gram_try"], 2) == 40.0


def test_an_alert_fires_once_and_is_removed(base, monkeypatch):
    values = iter([34.0, 36.0])
    monkeypatch.setattr(live, "quote_value", lambda subject: (next(values), "USD/TRY"))
    watch = live.add_watch({"type": "alert", "subject": "usd", "op": ">", "value": 35, "every": 900})
    assert live.check_watch(watch["id"]) == ""
    assert "above your 35" in live.check_watch(watch["id"])
    assert live.WATCHES.load() == []


def test_prices_found_on_product_pages():
    ld = '<script type="application/ld+json">{"@type":"Product","offers":{"price":"1.299,90","priceCurrency":"TRY"}}</script>'
    assert live.find_price(ld) == 1299.90
    assert live.find_price('<meta property="product:price:amount" content="249.99">') == 249.99
    assert live.find_price("<p>no price</p>") is None


def test_a_page_watch_notices_changes(base, monkeypatch):
    pages = iter(["Results are not out yet.", "Results are not out yet. Results are now published for all students."])
    monkeypatch.setattr(live, "page_text", lambda url: next(pages))
    first = live.page_text("x")
    watch = live.add_watch({"type": "page", "url": "https://example.com", "hash": __import__("hashlib").sha256(first.encode()).hexdigest(),
                            "text": first, "every": 3600})
    message = live.check_watch(watch["id"])
    assert "changed" in message and "published" in message


def test_next_prayer(monkeypatch):
    monkeypatch.setattr(live, "prayer_times", lambda lat, lon, day=None: {
        "Imsak": "05:20", "Fajr": "05:30", "Sunrise": "06:55", "Dhuhr": "13:00", "Asr": "16:10", "Maghrib": "18:50", "Isha": "20:10"})
    name, moment = live.next_prayer(41, 29, datetime(2026, 10, 1, 14, 0))
    assert name == "Asr" and moment.hour == 16
    name, moment = live.next_prayer(41, 29, datetime(2026, 10, 1, 21, 0))
    assert name == "Fajr" and moment.day == 2


# --- PC and files ---------------------------------------------------------------------------------

def test_duplicates(tmp_path):
    (tmp_path / "a.txt").write_text("same content", encoding="utf-8")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.txt").write_text("same content", encoding="utf-8")
    (tmp_path / "c.txt").write_text("different", encoding="utf-8")
    groups = pctools.find_duplicates(tmp_path)
    assert len(groups) == 1 and {p.name for p in groups[0]} == {"a.txt", "b.txt"}


def test_rename_rules_and_undo(jarvis, tmp_path):
    tmp_path = tmp_path / "photos"
    tmp_path.mkdir()
    for name in ("IMG_1.jpg", "IMG_2.jpg"):
        (tmp_path / name).write_bytes(b"x")
    plan = pctools.rename_plan(tmp_path, "replace IMG with Holiday")
    assert [b.name for _, b in plan] == ["Holiday_1.jpg", "Holiday_2.jpg"]
    with pytest.raises(ValueError):
        pctools.rename_plan(tmp_path, "replace IMG_1 with IMG_2")
    jarvis.process(f'/rename "{tmp_path}" number trip')
    jarvis.process("/rename go")
    assert sorted(p.name for p in tmp_path.iterdir()) == ["trip_001.jpg", "trip_002.jpg"]
    jarvis.process("/rename undo")
    assert sorted(p.name for p in tmp_path.iterdir()) == ["IMG_1.jpg", "IMG_2.jpg"]


def test_conversions(tmp_path):
    from PIL import Image

    csv = tmp_path / "d.csv"
    csv.write_text("name,score\nada,9\nbob,7.5\n", encoding="utf-8")
    xlsx = pctools.convert_file(csv, "xlsx", tmp_path)
    from openpyxl import load_workbook

    assert load_workbook(xlsx).active["B3"].value == 7.5
    png = tmp_path / "p.png"
    Image.new("RGBA", (10, 10), (255, 0, 0, 128)).save(png)
    assert pctools.convert_file(png, "jpg", tmp_path).suffix == ".jpg"
    md = tmp_path / "n.md"
    md.write_text("# Notes\n\n- one\n- two\n", encoding="utf-8")
    assert pctools.convert_file(md, "docx", tmp_path).exists()
    with pytest.raises(ValueError):
        pctools.convert_file(md, "mp3", tmp_path)


def test_pdf_merge_and_split(jarvis, tmp_path):
    from fpdf import FPDF

    for name in ("a", "b"):
        pdf = FPDF()
        for _ in range(2):
            pdf.add_page()
        pdf.output(str(tmp_path / f"{name}.pdf"))
    merged = jarvis.process(f"/pdfmerge {tmp_path / 'a.pdf'} {tmp_path / 'b.pdf'}").text
    assert "4 pages" in merged
    split = jarvis.process(f"/pdfsplit {tmp_path / 'a.pdf'} 2").text
    from pypdf import PdfReader

    assert len(PdfReader(split.splitlines()[-1].strip()).pages) == 1


def test_screen_time_counts_and_skips_idle(base):
    tracker = pctools.ScreenTime(every=5)
    tracker.sample("chrome", 1)
    tracker.sample("chrome", 1)
    tracker.sample("code", 500)        # away from the keyboard
    tracker.flush()
    assert pctools.SCREEN_TIME.load()[date.today().isoformat()] == {"chrome": 10}


def test_system_changes_ask_first(base, deny):
    from jarvis.assistant import Jarvis

    j = Jarvis(voice_enabled=False)
    if sys.platform == "win32":
        assert "Denied" in j.process("/dnd on").text or "already" in j.process("/darkmode").text
    assert "Nothing is listening on port 9" in j.process("/port kill 9").text


# --- security --------------------------------------------------------------------------------------

@pytest.mark.parametrize("password,verdict", [("password1", "very weak"), ("Galatasaray1905", "very weak"), ("Ab3$fg9", "fair"),
                                              ("lamba-kedi-river-orbit-91", "very strong")])
def test_password_strength(password, verdict):
    assert guard.password_strength(password)["verdict"] == verdict


def test_breach_check_sends_only_a_prefix(monkeypatch):
    import hashlib

    digest = hashlib.sha1(b"hunter2").hexdigest().upper()
    seen = {}

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def read(self):
            return f"{digest[5:]}:42\r\nABCDEF:1".encode()

    def fake(request, timeout=0):
        seen["url"] = request.full_url
        return Response()

    monkeypatch.setattr(guard.net, "urlopen", fake)
    assert guard.pwned_count("hunter2") == 42
    assert seen["url"].endswith(digest[:5]) and "hunter2" not in seen["url"]


@pytest.mark.parametrize("url,risky", [("https://paypa1.com/login", True), ("http://192.168.4.20/verify", True),
                                       ("https://xn--pple-43d.com", True), ("https://www.google.com/search?q=x", False),
                                       ("https://garanti-bonus-odul.xyz/kazan", True)])
def test_link_checks(url, risky):
    score, reasons, _ = guard.link_risks(url)
    assert (score >= 3) == risky, reasons


def test_hash_encrypt_decrypt_shred(jarvis, tmp_path):
    secret = tmp_path / "plan.txt"
    secret.write_text("abc", encoding="utf-8")
    assert guard.file_hashes(secret)["sha256"] == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    assert "Matches" in jarvis.process(f"/hash {secret} ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad").text
    locked = guard.encrypt_file(secret, "correct horse")
    assert b"abc" not in locked.read_bytes()
    secret.unlink()
    with pytest.raises(ValueError):
        guard.decrypt_file(locked, "wrong horse")
    assert guard.decrypt_file(locked, "correct horse").read_text(encoding="utf-8") == "abc"
    jarvis.process(f"/shred {tmp_path / 'plan.txt'}")
    assert not (tmp_path / "plan.txt").exists()


def test_totp_matches_rfc_6238():
    secret = base64.b32encode(b"12345678901234567890").decode()
    assert guard.totp(secret, at=59, digits=8) == "94287082"
    assert guard.totp(secret, at=1111111109, digits=8) == "07081804"
    assert guard.parse_secret("otpauth://totp/GitHub:me?secret=JBSWY3DPEHPK3PXP&issuer=GitHub")["secret"] == "JBSWY3DPEHPK3PXP"


def test_secrets_never_reach_the_chat(jarvis, monkeypatch):
    assert jarvis.mask("/pwned hunter2") == "/pwned " + "•" * 8
    prompts, finish = jarvis.secret_request("/pwned")
    assert prompts[0][1] is True
    monkeypatch.setattr(guard, "pwned_count", lambda p: 0)
    assert "Not found" in finish(["hunter2"])
    prompts, finish = jarvis.secret_request("/2fa add github")
    assert "github" in prompts[0][0]
    assert "Added github" in finish(["JBSWY3DPEHPK3PXP"])
    assert "github" in jarvis.process("/2fa").text
    assert "JBSWY3DPEHPK3PXP" not in json.dumps(jarvis.brain.history)
    assert jarvis.secret_request("hello") is None


# --- coding --------------------------------------------------------------------------------------

def test_docstrings_are_inserted_and_code_still_parses():
    import ast

    source = "def add(a, b):\n    return a + b\n\n\nclass Box:\n    size = 1\n\n    def grow(self):\n        self.size += 1\n"
    updated, added = devtools.insert_docstrings(source, {"add": "Add two numbers.", "Box": "A box.\n\nIt grows.", "grow": "Make it bigger."})
    tree = ast.parse(updated)
    assert added == 3
    assert ast.get_docstring(tree.body[0]) == "Add two numbers."
    assert ast.get_docstring(tree.body[1]) == "A box.\n\nIt grows."
    assert "return a + b" in updated


def test_jwt_b64_json_env(jarvis, tmp_path):
    token = ".".join(base64.urlsafe_b64encode(json.dumps(p).encode()).decode().rstrip("=") for p in
                     ({"alg": "HS256"}, {"sub": "42", "exp": 1})) + ".sig"
    reply = jarvis.process(f"/jwt {token}").text
    assert '"sub": "42"' in reply and "EXPIRED" in reply
    assert jarvis.process("/b64 decode " + base64.b64encode("merhaba".encode()).decode()).text == "merhaba"
    assert "line 1, column" in jarvis.process('/json {"a": 1,}').text
    (tmp_path / ".env").write_text("API_KEY=secret-value\nEMPTY=\n", encoding="utf-8")
    (tmp_path / ".env.example").write_text("API_KEY=\nDB_URL=\n", encoding="utf-8")
    check = jarvis.process(f"/envcheck {tmp_path}").text
    assert "DB_URL" in check and "EMPTY" in check and "secret-value" not in check


def test_python_scratchpad(jarvis):
    assert "285" in jarvis.process("/py print(sum(i*i for i in range(10)))").text


# --- make and write -----------------------------------------------------------------------------------

def test_slides_are_a_real_deck(jarvis, monkeypatch):
    from pptx import Presentation

    deck = {"title": "Energy", "subtitle": "A short talk", "slides": [
        {"title": "Solar", "bullets": ["Cheap", "Fast to build"], "notes": "Say this."},
        {"title": "Wind", "bullets": ["Strong at night"], "notes": ""}]}
    monkeypatch.setattr(jarvis.brain, "ask_once", lambda *a, **k: json.dumps(deck))
    reply = jarvis.process("make a presentation about energy").text
    path = next(line.strip() for line in reply.splitlines() if line.strip().endswith(".pptx"))
    prs = Presentation(path)
    assert len(prs.slides) == 3
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Say this."


def test_citations(monkeypatch):
    item = {"authors": ["Ada Lovelace", "Charles Babbage"], "year": "1843", "title": "Notes on the Engine",
            "site": "Scientific Memoirs", "url": "https://example.org", "kind": "web"}
    assert makers.format_citation(item, "apa").startswith("Lovelace, A., & Babbage, C. (1843). *Notes on the Engine*.")
    assert makers.format_citation(item, "mla").startswith("Lovelace, Ada, and Charles Babbage.")
    seen = {}
    monkeypatch.setattr(kit, "get_text", lambda url, headers=None, **k: (seen.update(url=url, accept=headers["Accept"]), "LeCun, Y. (2015). Deep learning.")[1])
    assert makers.cite("doi:10.1038/nature14539", "mla").startswith("LeCun")
    assert seen["url"] == "https://doi.org/10.1038/nature14539" and "modern-language-association" in seen["accept"]


def test_invoice_totals_and_pdf(base):
    rows, subtotal, tax, total = makers.invoice_totals([{"description": "Design", "quantity": 1, "unit_price": 12000},
                                                        {"description": "Hosting", "quantity": 12, "unit_price": 150}], 20)
    assert (subtotal, tax, total) == (13800.0, 2760.0, 16560.0)
    path = makers.make_invoice({"from": "Ahmed", "to": "Acme", "currency": "TRY", "tax_percent": 20,
                                "items": [{"description": "Design", "quantity": 1, "unit_price": 12000}]})
    assert path.name.startswith(f"INV-{date.today():%Y}-001") and path.stat().st_size > 1000
    assert makers.make_invoice({"items": [{"description": "x", "unit_price": 1}]}).name.endswith("-002.pdf")


def test_pictures(tmp_path):
    from PIL import Image

    drawing.mind_map({"center": "Energy", "branches": [{"name": "Solar", "children": ["Cheap", "Panels"]},
                                                         {"name": "Wind", "children": ["Turbines"]}]}, tmp_path / "m.png")
    drawing.flowchart({"nodes": [{"id": "a", "text": "Start", "type": "start"}, {"id": "b", "text": "Logged in?", "type": "decision"},
                                 {"id": "c", "text": "Done", "type": "end"}],
                       "edges": [{"from": "a", "to": "b"}, {"from": "b", "to": "c", "label": "yes"}, {"from": "b", "to": "a", "label": "no"}]},
                      tmp_path / "f.png")
    Image.new("RGB", (400, 300), "navy").save(tmp_path / "base.png")
    drawing.meme(tmp_path / "base.png", "when the tests pass", "on the first try", tmp_path / "meme.png")
    for name in ("m.png", "f.png", "meme.png"):
        assert Image.open(tmp_path / name).size[0] >= 400


# --- voice, connections and fun ---------------------------------------------------------------------------

def test_persona_goes_into_the_system_prompt(jarvis):
    jarvis.process("/persona coach")
    assert "coach" in jarvis.brain.system_prompt().lower() and "JARVIS" in jarvis.brain.system_prompt()
    jarvis.process("/persona off")
    assert "Persona for this session" not in jarvis.brain.system_prompt()


def test_whatsapp_opens_with_the_message_and_never_sends(jarvis, monkeypatch):
    opened = []
    import os

    monkeypatch.setattr(os, "startfile", lambda url: opened.append(url), raising=False)
    jarvis.process("/whatsapp add Ali 0555 111 22 33")
    reply = jarvis.process("whatsapp ali: running 10 min late").text
    assert "press Send" in reply
    assert opened[0].startswith("whatsapp://send?phone=905551112233&text=running%2010%20min%20late")


def test_twenty_questions_and_watchlist(jarvis, monkeypatch):
    moves = iter([{"question": "Is it alive?"}, {"guess": "a cat"}])
    monkeypatch.setattr(jarvis.brain, "ask_once", lambda *a, **k: json.dumps(next(moves)))
    assert "Is it alive?" in jarvis.process("/20q").text
    assert "a cat" in jarvis.process("yes").text
    assert "Got it in 2" in jarvis.process("yes").text
    jarvis.process("add Dune to my watchlist")
    assert "8/10" in jarvis.process("/watchlist done 1 8").text


def test_gcal_setup_checks_the_address(jarvis):
    assert "doesn't look like" in jarvis.setup_gcal("https://example.com/cal.ics")


# --- providers and languages -------------------------------------------------------------------------------

def test_mistral_and_cloudflare(monkeypatch):
    monkeypatch.setenv("MISTRAL_API_KEY", "x")
    monkeypatch.setenv("CLOUDFLARE_API_TOKEN", "y")
    monkeypatch.delenv("CLOUDFLARE_ACCOUNT_ID", raising=False)
    monkeypatch.setattr(providers, "load_config", lambda: None)
    assert providers.has_key(providers.MISTRAL)
    assert not providers.has_key(providers.CLOUDFLARE)       # the account ID is needed too
    monkeypatch.setenv("CLOUDFLARE_ACCOUNT_ID", "abc123")
    assert providers.has_key(providers.CLOUDFLARE)
    assert providers.base_url_for(providers.CLOUDFLARE) == "https://api.cloudflare.com/client/v4/accounts/abc123/ai/v1"
    names = [p.name for p in providers.AUTO_CHAT_PROVIDERS]
    assert names.index("mistral") < names.index("openai") and names.index("cloudflare") < names.index("openai")
    assert providers.MISTRAL in providers.STT_PROVIDERS


def test_every_language_is_complete():
    for code, (_name, instruction, values) in i18n_more.TABLES.items():
        assert len(values) == len(i18n_more.KEYS), code
        assert code in i18n.available() and i18n.ANSWER_IN[code]


def test_a_cut_off_reply_is_repaired():
    cut = '{"title": "Water", "slides": [{"title": "A", "bullets": ["x"]}, {"title": "B", "bullets": ["y"]}, {"title": "C", "notes": "half'
    assert kit.parse_json(cut) == {"title": "Water", "slides": [{"title": "A", "bullets": ["x"]}, {"title": "B", "bullets": ["y"]}]}
    assert kit.parse_json("```json\n[[1], [2], [3") == [[1], [2]]


def test_long_calls_get_more_room(jarvis, monkeypatch):
    seen = []
    monkeypatch.setattr(jarvis.brain, "ask_once", lambda *a, **k: (seen.append(jarvis.brain.min_tokens), "{}")[1])
    kit.ask_json(jarvis.brain, "x")
    assert seen == [6000] and jarvis.brain.min_tokens == 0
