"""10.0 Design 2.0: charts, pen strokes, groups, patterns, curved text,
transitions and animations (drawn and in PowerPoint), find and replace,
spell check, white backgrounds, Excel and PDF in, sketches, the new
templates and commands — then the page itself: several selected, groups,
rulers and guides, the pen, chart data, the slideshow and presenter view."""

from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from PIL import Image

from jarvis.design import chart, extras, model, pptxio, render, templates


def _slide() -> tuple[dict, dict]:
    design = model.new_design("slides", "slides", "Test")
    return design, design["pages"][0]


# --- charts --------------------------------------------------------------------------------------------

def test_chart_numbers_and_data():
    assert chart.number("1.250,5") == 1250.5 and chart.number("1,250.5") == 1250.5
    assert chart.number("₺40") == 40 and chart.number("12,5") == 12.5 and chart.number("1,200") == 1200
    assert chart.number("") is None and chart.number("n/a") is None
    labels, series = chart.parse([["", "2025", "2026", "Notes"], ["Q1", "12", "15", "ok"], ["Q2", "18", "", "-"]])
    assert labels == ["Q1", "Q2"] and [n for n, _ in series] == ["2025", "2026"]   # the text column is dropped
    assert series[1][1] == [15.0, None]
    assert chart.nice_step(87) == 20 and chart.short(12400) == "12.4k" and chart.short(3) == "3"


@pytest.mark.parametrize("kind", model.CHARTS)
def test_every_chart_kind_draws(kind):
    el = model.chart(kind, model.SAMPLE_CHART, 0, 0, 600, 400, title="Sales", labels=True)
    image = chart.render(el, 1.0)
    assert image.size == (600, 400) and image.getbbox() is not None
    assert len(image.convert("RGB").getcolors(1 << 20)) > 20                  # bars, text, colours
    empty = chart.render(model.chart(kind, [["", "A"], ["x", "?"]], 0, 0, 300, 200), 1.0)
    assert empty.getbbox() is not None                                         # says what's missing


def test_pen_strokes_patterns_and_curved_text_draw():
    design, page = _slide()
    stroke = model.add(design, page, model.path([(100, 100), (150, 160), (220, 90), (300, 200)], stroke="#ff0000",
                                                stroke_w=10))
    assert (stroke["x"], stroke["y"], stroke["w"], stroke["h"]) == (100, 90, 200, 110)
    assert stroke["points"][0] == [0.0, 10 / 110] and max(p[0] for p in stroke["points"]) == 1.0
    assert len(render.smooth_points([(0, 0), (10, 0), (10, 10)])) > 3
    image = render.render_page(design, 0, 1.0)
    assert image.getpixel((150, 160))[0] > 200                                 # red ink on the stroke
    plain = model.shape("rect", 0, 0, 300, 300, fill="#2563eb")
    striped = dict(plain, pattern="stripes", pattern_color="#ffffff")
    assert render.shape_layer(plain, 1.0)[0].tobytes() != render.shape_layer(striped, 1.0)[0].tobytes()
    straight = model.text("Curved words on a line", 0, 0, 600, 150, size=40)
    curved = dict(straight, curve=150)
    a, b = render.text_layer(straight, 1.0)[0], render.text_layer(curved, 1.0)[0]
    assert a.getbbox() and b.getbbox() and a.tobytes() != b.tobytes()
    assert render.text_layer(dict(curved, curve=-360, shadow=True, stroke="#000000", stroke_w=2), 1.0)[0].getbbox()


def test_a_missing_glyph_comes_from_segoe_ui():
    face = render.face("Constantia", 40)
    if face.font.path.lower().endswith("constan.ttf"):
        assert face.runs("₺650")[0] == ["₺", "plain"]
    assert face.getlength("₺650") > 0


def test_the_emoji_list():
    import time

    from ui import design_tools

    start = time.monotonic()
    found = dict(design_tools.all_emoji())
    assert time.monotonic() - start < 10
    if render.fonts.emoji(20)[1]:
        assert "🎉" in found and found["😀"] == "grinning face" and len(found) > 800


# --- groups, aligning, words ----------------------------------------------------------------------------

def test_groups_and_aligning():
    design, page = _slide()
    a = model.add(design, page, model.shape("rect", 0, 0, 100, 100))
    middle = model.add(design, page, model.shape("rect", 500, 500, 100, 100))
    b = model.add(design, page, model.shape("ellipse", 300, 100, 50, 80))
    group = model.make_group(design, page, [a, b])
    assert group and a["group"] == b["group"] == group
    order = [e["id"] for e in page["elements"]]
    assert order.index(a["id"]) + 1 == order.index(b["id"]) and order[0] == middle["id"]   # gathered together
    assert model.group_members(page, group) == [a, b]
    model.align([a, b], "left")
    assert a["x"] == b["x"] == 0
    model.align([a, b], "middle")
    assert model.center(a)[1] == pytest.approx(model.center(b)[1])
    model.align([middle], "center", (1920, 1080))
    assert model.center(middle)[0] == 960
    c = model.add(design, page, model.shape("rect", 900, 0, 100, 100))
    model.align([a, middle, c], "hspace")
    gaps = sorted(model.bbox(e)[0] for e in (a, middle, c))
    assert gaps[1] - gaps[0] - 100 == pytest.approx(gaps[2] - gaps[1] - 100)
    assert model.ungroup([a, b]) == 2 and not a["group"]
    copy = model.duplicate(design, page, b)
    assert copy["group"] == ""


def test_find_and_replace_everywhere():
    design, page = _slide()
    model.add(design, page, model.text("Kermes on Saturday", 0, 0, 500, 100))
    model.add(design, page, model.table([["Day", "Event"], ["Saturday", "kermes"]], 0, 200, 500, 200))
    model.add(design, page, model.chart("bar", model.SAMPLE_CHART, 0, 400, 500, 300, title="Kermes sales"))
    page["notes"] = "Remind them about the kermes."
    assert len(model.find_text(design, "kermes")) == 3 and model.find_text(design, "kermes", case=True) == [
        (0, page["elements"][1]["id"])]
    assert model.replace_text(design, "kermes", "bazaar") == 4
    assert "bazaar on Saturday" in page["elements"][0]["text"] and page["elements"][1]["rows"][1][1] == "bazaar"
    assert page["elements"][2]["title"] == "bazaar sales" and "bazaar" in page["notes"]


def test_cleaning_new_properties():
    el = model.clean_element({"type": "text", "curve": 999, "anim": "spin", "group": 7})
    assert el["curve"] == 360 and el["anim"] == "none" and el["group"] == "7"
    path = model.clean_element({"type": "path", "points": [[2, -1], ["x", 0]]})
    assert path["points"] == [[0.0, 0.0], [1.0, 1.0]]                          # junk falls back to the default
    design = model.normalize({"pages": [{"transition": "fade"}, {"transition": "explode"}],
                              "guides": {"v": [100, "x"], "h": [50.5]}})
    assert design["pages"][0]["transition"] == "fade" and "transition" not in design["pages"][1]
    assert design["guides"] == {"v": [100.0], "h": [50.5]}


# --- PowerPoint ---------------------------------------------------------------------------------------------

def test_powerpoint_keeps_charts_strokes_groups_transitions_and_animations(base):
    from pptx import Presentation
    from pptx.oxml.ns import qn

    design, page = _slide()
    page["transition"] = "push"
    model.add(design, page, model.text("Arch", 100, 50, 800, 200, curve=120, stroke="#000000", stroke_w=3,
                                       shadow=True, anim="fade"))
    model.add(design, page, model.chart("pie", model.SAMPLE_CHART, 100, 300, 700, 600, title="Share", labels=True,
                                        anim="zoom"))
    model.add(design, page, model.path([(1000, 300), (1100, 400), (1300, 320)], stroke="#e11d48", stroke_w=8,
                                       anim="wipe"))
    a = model.add(design, page, model.shape("rect", 1000, 600, 300, 200, pattern="dots", pattern_color="#ffffff"))
    b = model.add(design, page, model.shape("ellipse", 1350, 600, 200, 200, anim="fly"))
    model.make_group(design, page, [a, b])
    path = pptxio.export_pptx(design, base / "d2.pptx")
    slide = Presentation(str(path)).slides[0]
    xml = slide._element
    assert xml.find(qn("p:transition")).find(qn("p:push")) is not None
    timing = xml.find(qn("p:timing"))
    effects = [c for c in timing.iter(qn("p:cTn")) if c.get("presetClass") == "entr"]
    assert [e.get("presetID") for e in effects] == ["10", "53", "22", "2"]
    ids = [int(c.get("id")) for c in timing.iter(qn("p:cTn"))]
    assert len(ids) == len(set(ids))                                           # every timing node has its own id
    kinds = {shape.shape_type for shape in slide.shapes}
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    assert {MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.GROUP} <= kinds
    text = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text == "Arch")
    assert text.text_frame._bodyPr.find(qn("a:prstTxWarp")).get("prst") == "textArchUp"
    rPr = text.text_frame.paragraphs[0].runs[0]._r.rPr
    assert rPr[0].tag == qn("a:ln") and rPr.find(qn("a:effectLst")) is not None
    back = pptxio.import_pptx(path)
    els = back["pages"][0]["elements"]
    assert back["pages"][0]["transition"] == "push"
    pie = next(e for e in els if e["type"] == "chart")
    assert pie["chart"] == "pie" and pie["rows"][1][:2] == ["Q1", "12"] and pie["title"] == "Share"
    assert any(e["type"] == "path" for e in els)
    members = [e for e in els if e.get("group")]
    assert len(members) == 2 and members[0]["pattern"] == "dots"


def test_timing_ids_are_unique_for_many_shapes():
    xml = pptxio.timing_xml([(i, anim, i % 2 == 0) for i, anim in enumerate(["fade", "fly", "zoom", "wipe",
                                                                             "appear"] * 3, 2)])
    assert xml.count('presetClass="entr"') == 15 and xml.count("<p:bldP") == 8


# --- pictures, tables, PDFs, sketches, words ------------------------------------------------------------

def test_remove_white_keeps_the_white_inside():
    image = Image.new("RGB", (120, 120), "#ffffff")
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    draw.ellipse((20, 20, 100, 100), fill="#1d4ed8")
    draw.ellipse((50, 50, 70, 70), fill="#ffffff")                             # the white of an eye
    out = extras.remove_white(image)
    assert out.mode == "RGBA" and out.getpixel((2, 2))[3] == 0 and out.getpixel((60, 30))[3] == 255
    assert out.getpixel((60, 60))[3] > 200


def test_slides_from_a_table(base):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Şehir", "Nüfus", "Alan"])
    for row in (["İstanbul", 15655924, 5343.0], ["Ankara", 5803482, 25632.0], ["İzmir", 4479525, 11891.0]):
        sheet.append(row)
    sheet.append([None, None, None])
    path = base / "cities.xlsx"
    book.save(path)
    rows = extras.read_table(path)
    assert rows[0] == ["Şehir", "Nüfus", "Alan"] and rows[1] == ["İstanbul", "15655924", "5343"] and len(rows) == 4
    assert extras.numeric_columns(rows) == [1, 2]
    per_row = extras.slides_from_table(rows, "Cities", "rows", "ocean")
    assert len(per_row["pages"]) == 1 + 3 + 1                                  # cover, a slide per row, a chart
    assert "Nüfus: 15655924" in model.page_text(per_row["pages"][1])
    assert per_row["pages"][-1]["elements"][-1]["type"] == "chart"
    table = extras.slides_from_table(rows, "Cities", "table")
    assert any(e["type"] == "table" for e in table["pages"][1]["elements"])
    (base / "t.csv").write_text("Ad;Puan\nAli;90\nAyşe;95\n", encoding="utf-8")
    assert extras.read_table(base / "t.csv") == [["Ad", "Puan"], ["Ali", "90"], ["Ayşe", "95"]]
    with pytest.raises(extras.DesignAIError):
        extras.read_table(base / "x.docx")


def test_a_pdfs_pages_become_a_design(base):
    from jarvis import winrt

    if not winrt.available():
        pytest.skip("needs Windows' PDF renderer")
    from fpdf import FPDF

    pdf = FPDF(orientation="L")
    for n in (1, 2):
        pdf.add_page()
        pdf.set_font("Helvetica", size=40)
        pdf.cell(0, 40, f"Page {n}")
    path = base / "handout.pdf"
    pdf.output(str(path))
    design = extras.design_from_pdf(path)
    assert len(design["pages"]) == 2 and design["format"] == "a4_landscape"
    first = design["pages"][0]["elements"][0]
    assert first["type"] == "image" and first["locked"] and model.resolve_src(first["src"]) is not None


SKETCH = {"title": "Bake sale", "bg": "#fff7ed", "palette": ["#ea580c", "#16a34a"],
          "elements": [{"type": "text", "x": 0.1, "y": 0.05, "w": 0.8, "h": 0.15, "role": "title",
                        "text": "Bake Sale!", "size": 0.08},
                       {"type": "image", "x": 0.1, "y": 0.25, "w": 0.8, "h": 0.4},
                       {"type": "shape", "x": 0.3, "y": 0.75, "w": 0.4, "h": 0.08, "role": "button",
                        "text": "Saturday 10am", "shape": "round"},
                       {"type": "icon", "x": 0.85, "y": 0.85, "w": 0.1, "h": 0.1, "glyph": "🧁"},
                       {"type": "nonsense"}]}


def test_a_design_from_a_sketch(base, tmp_path):
    design = extras.build_from_sketch(SKETCH, "poster")
    kinds = [e["type"] for e in design["pages"][0]["elements"]]
    assert kinds[:4] == ["text", "image", "shape", "icon"] and design["w"] == 1240
    title = design["pages"][0]["elements"][0]
    assert title["text"] == "Bake Sale!" and title["size"] == pytest.approx(1754 * 0.08)
    with pytest.raises(extras.DesignAIError):
        extras.build_from_sketch({"elements": []}, "poster")
    photo = tmp_path / "sketch.jpg"
    Image.new("RGB", (600, 900), "white").save(photo)
    asked = []
    brain = SimpleNamespace(ask_once=lambda prompt, **k: (asked.append(k), json.dumps(SKETCH))[1])
    design = extras.design_from_sketch(brain, photo)
    assert design["format"] == "poster" and asked[0].get("image_b64")             # portrait photo -> A4


def test_rewrite_and_spelling(monkeypatch):
    brain = SimpleNamespace(ask_once=lambda prompt, **k: '"Big summer sale"')
    assert extras.rewrite(brain, "We are having a sale in the summer", "punchy") == "Big summer sale"
    with pytest.raises(extras.DesignAIError):
        extras.rewrite(SimpleNamespace(ask_once=lambda p, **k: "No AI provider is configured."), "x", "shorter")
    design, page = _slide()
    el = model.add(design, page, model.text("Teh the the best", 0, 0, 500, 100))
    issue = {"page": 0, "id": el["id"], "field": "text", "start": 0, "length": 3, "word": "Teh", "suggestions": ["The"]}
    repeat = {"page": 0, "id": el["id"], "field": "text", "start": 8, "length": 3, "word": "the", "suggestions": [""]}
    assert extras.apply_fix(design, issue, "The") and extras.apply_fix(design, repeat, "")
    assert el["text"] == "The the best"
    fixes = {"fixes": [{"id": f"0:{el['id']}:text", "wrong": "best", "right": "beast"}]}
    found = extras.spelling_issues_ai(SimpleNamespace(ask_once=lambda p, **k: json.dumps(fixes)), design)
    assert found[0]["word"] == "best" and found[0]["start"] == 8
    from jarvis import spelling

    if not spelling.available():
        pytest.skip("no Windows spell checker")
    issues = extras.spelling_issues(design | {"pages": [{"elements": [model.text("Ths is speling", 0, 0, 9, 9)]}]})
    assert {i["word"] for i in issues} == {"Ths", "speling"}
    if spelling.supported("tr-TR"):
        assert [i["word"] for i in spelling.check("Bugün hava çok güzl")] == ["güzl"]


# --- templates and commands ------------------------------------------------------------------------------

def test_the_new_templates():
    for kind in ("cv", "bookcover", "ticket", "infographic", "carousel", "calendar"):
        assert kind in templates.KINDS
        for variant in templates.KINDS[kind]["variants"]:
            design = templates.make(kind, variant)
            assert render.thumbnail(design, 200).getbbox(), (kind, variant)
    carousel = templates.make("carousel", "tips", {"title": "Tips", "slides": [{"title": "One", "text": "a"}] * 3})
    assert len(carousel["pages"]) == 5 and carousel["format"] == "instagram"
    from jarvis.design import more_templates

    rows = more_templates.month_rows(2026, 2, "tr", {"14": "Sevgililer Günü"})
    assert rows[0][0] == "Pzt" and "14 •" in [c for r in rows for c in r] and rows[1][6] == "1"
    cv = templates.make("cv", "modern", {"name": "Zeynep Ak", "title": "Engineer"})
    assert "Zeynep Ak" in model.page_text(cv["pages"][0]) and "Engineer" in model.page_text(cv["pages"][0])
    from jarvis.design import ai

    assert {"cv", "ticket", "calendar"} <= set(ai.FIELDS) and "calendar" in ai.ALL_KINDS


def test_design_commands(base, monkeypatch, tmp_path):
    from jarvis.ten import design10

    assert design10.calendar_content("kasım 2026 aile takvimi") == {"title": "aile takvimi", "month": 11,
                                                                    "year": 2026, "lang": "tr"}
    assert design10.calendar_content("family stuff") is None
    from jarvis.assistant import Jarvis

    jarvis = Jarvis.__new__(Jarvis)
    jarvis.brain = SimpleNamespace(ask_once=lambda *a, **k: "")
    reply = jarvis.calendarpage_cmd("march 2027 minimal")
    assert "March 2027" in reply.text and reply.design_path.exists()
    show = jarvis.slideshow_cmd("presenter")
    assert show.open_page == "design:presenter" and show.design_path == reply.design_path
    picture = tmp_path / "logo.jpg"
    image = Image.new("RGB", (100, 100), "white")
    image.paste((200, 0, 0), (30, 30, 70, 70))
    image.save(picture)
    out = jarvis.removebg_cmd(str(picture))
    assert out.image_path.name == "logo_nobg.png" and Image.open(out.image_path).getpixel((1, 1))[3] == 0
    csv_path = tmp_path / "s.csv"
    csv_path.write_text("Product,Sold\nTea,40\nCoffee,55\n", encoding="utf-8")
    made = jarvis.excelslides_cmd(f"{csv_path} | table ocean")
    assert "PowerPoint" in made.text
    assert "Usage" in jarvis.pdfdesign_cmd("nothing.txt") and "Usage" in jarvis.sketch_cmd("")


def test_slideshow_steps_and_frames():
    from ui import design_show

    design, page = _slide()
    a = model.add(design, page, model.text("A", 0, 0, 100, 100, anim="fade"))
    model.add(design, page, model.text("still", 0, 0, 100, 100))
    b = model.add(design, page, model.shape("rect", 0, 0, 100, 100, anim="fly"))
    c = model.add(design, page, model.shape("rect", 0, 0, 100, 100))
    model.make_group(design, page, [b, c])
    assert design_show.steps(page) == [("fade", [a["id"]]), ("fly", [b["id"], c["id"]])]
    before, after = Image.new("RGB", (64, 36), "black"), Image.new("RGB", (64, 36), "white")
    for kind in model.TRANSITIONS:
        frame = design_show.transition_frame(kind, before, after, 0.5)
        assert frame.size == (64, 36)
    assert design_show.transition_frame("wipe", before, after, 1.0).getpixel((63, 0)) == (255, 255, 255)
    layer = Image.new("RGBA", (64, 36), (0, 0, 0, 0))
    layer.paste((255, 0, 0, 255), (10, 10, 30, 30))
    for anim in ("fade", "fly", "zoom", "wipe", "appear"):
        assert design_show.effect_frame(anim, before, layer, (10, 10, 30, 30), 1.0).getpixel((20, 20)) == (255, 0, 0)


# --- the page ----------------------------------------------------------------------------------------------------

from tests.test_gui import _display_available, app, inline_workers  # noqa: E402,F401


gui = pytest.mark.skipif(not _display_available(), reason="no display")


@pytest.fixture
def page(app, monkeypatch):
    monkeypatch.setattr(app, "_open_path", lambda path: None)
    app.geometry("1400x860")
    app._show_tab("design")
    design_page = app.design_page()
    app.update()
    design_page.fit()
    app.update()
    return design_page


CLOCK = [500_000]


def _entries(widget) -> list:
    found = []
    for child in widget.winfo_children():
        if child.winfo_class() == "Entry":
            found.append(child)
        found += _entries(child)
    return found


def _fire(widget, sequence: str) -> None:
    """Run a widget's binding directly (generated key events go to whatever holds the OS focus)."""
    import re

    script = widget.bind(sequence)
    filled = re.sub(r"%[#bfhkstwxyAEKNWTXYD]", lambda m: str(widget) if m.group() == "%W" else "0", script)
    widget.tk.eval(filled)
    widget.update()


def _screen(page, x, y):
    cx, cy = page._to_canvas(x, y)
    return int(cx - page.canvas.canvasx(0)), int(cy - page.canvas.canvasy(0))


def _drag(page, start, end, state=0, steps=4):
    sx, sy = _screen(page, *start)
    CLOCK[0] += 1000
    page.canvas.event_generate("<ButtonPress-1>", x=sx, y=sy, state=state, time=CLOCK[0])
    for i in range(1, steps + 1):
        mx, my = _screen(page, start[0] + (end[0] - start[0]) * i / steps, start[1] + (end[1] - start[1]) * i / steps)
        page.canvas.event_generate("<B1-Motion>", x=mx, y=my, state=state | 0x100, time=CLOCK[0] + i)
    page.canvas.event_generate("<ButtonRelease-1>", x=mx, y=my, state=state, time=CLOCK[0] + 10)
    page.update()


def _click(page, x, y, state=0):
    sx, sy = _screen(page, x, y)
    CLOCK[0] += 1000
    page.canvas.event_generate("<ButtonPress-1>", x=sx, y=sy, state=state, time=CLOCK[0])
    page.canvas.event_generate("<ButtonRelease-1>", x=sx, y=sy, state=state, time=CLOCK[0] + 5)
    page.update()


@gui
def test_select_several_group_move_and_copy(page):
    page.load(templates.blank("slides"))
    a = page.add_shape("rect")
    model.move(a, -500, -250)
    b = page.add_shape("ellipse")
    model.move(b, 400, 200)
    page.redraw()
    _click(page, *model.center(a))
    _click(page, *model.center(b), state=0x1)                                   # Shift+click adds
    assert page.selection == [a["id"], b["id"]]
    page._set("fill", "#ff0000")
    assert a["fill"] == b["fill"] == "#ff0000"
    page.group_selected()
    assert a["group"] and a["group"] == b["group"]
    page.select(None)
    _click(page, *model.center(a))                                             # a click selects the whole group
    assert set(page.selection) == {a["id"], b["id"]}
    ax, bx = a["x"], b["x"]
    _drag(page, model.center(a), (model.center(a)[0] + 37, model.center(a)[1]))
    assert a["x"] - ax == pytest.approx(b["x"] - bx) and abs(a["x"] - ax) > 20
    page.duplicate_selected()
    copies = page.selected_all()
    assert len(copies) == 2 and copies[0]["group"] == copies[1]["group"] != a["group"]
    page.undo()
    a, b = model.find(page.page, a["id"]), model.find(page.page, b["id"])     # undo brings back a fresh copy
    assert len(page.page["elements"]) == 2
    page.select(a["id"], single=True)                                          # Ctrl+click: one member
    assert page.selection == [a["id"]]
    page.select(None)
    _drag(page, (5, 5), (page.design["w"] - 5, page.design["h"] - 5))         # a box around everything
    assert set(page.selection) == {a["id"], b["id"]}
    page.ungroup_selected()
    assert not a["group"]
    page.align_selected("top")
    assert model.bbox(a)[1] == pytest.approx(model.bbox(b)[1])
    page.delete_selected()
    assert not page.page["elements"]
    page.undo()
    assert len(page.page["elements"]) == 2


@gui
def test_rulers_guides_snapping_and_the_pen(page):
    page.load(templates.blank("slides"))
    assert page.ruler_top.find_all() and page.ruler_left.find_all()
    page._ruler_press("v")
    x0 = page.canvas.winfo_rootx() + _screen(page, 700, 300)[0]
    page._ruler_motion("v", SimpleNamespace(x_root=x0, y_root=page.canvas.winfo_rooty() + 200))
    page._ruler_release()
    assert page.design["guides"]["v"] == [pytest.approx(700, abs=3)]
    page.design["guides"]["v"] = [700]
    el = page.add_shape("rect")
    page.select(None)
    left = model.bbox(el)[0]
    _drag(page, model.center(el), (model.center(el)[0] + (700 - left) + 3, model.center(el)[1]))
    assert model.bbox(el)[0] == pytest.approx(700)                              # pulled onto the guide
    page.toggle_rulers()
    assert not page.ruler_top.winfo_ismapped()
    page.toggle_rulers()
    page.clear_guides()
    assert page.design["guides"] == {"v": [], "h": []}
    page.toggle_pen(True)
    count = len(page.page["elements"])
    _drag(page, (200, 800), (500, 900), steps=8)
    stroke = page.page["elements"][-1]
    assert len(page.page["elements"]) == count + 1 and stroke["type"] == "path" and len(stroke["points"]) >= 5
    page._escape()
    assert not page._pen_on


@gui
def test_charts_text_tools_and_words(page, monkeypatch, inline_workers):
    import ui.designui as designui
    from tests.test_design_gui import inline_workers_thread

    monkeypatch.setattr(designui.threading, "Thread", inline_workers_thread())
    page.load(templates.blank("slides"))
    graph = page.add_chart("line")
    assert graph["type"] == "chart" and page.selected() is graph
    page.edit_chart_data()
    dialog = page._chart_dialog
    page.update()
    cells = _entries(dialog)
    assert len(cells) == len(model.SAMPLE_CHART) * 3
    cells[4].delete(0, "end")
    cells[4].insert(0, "99")                                                   # Q1's 2025 value
    _fire(cells[4], "<KeyRelease>")
    assert graph["rows"][1][1] == "99"
    dialog.destroy()
    page._set("rows", [["", "Visits"], ["Mon", "4"], ["Tue", "9"]])
    assert graph["rows"][2] == ["Tue", "9"]
    text = page.add_text("body")
    page.apply_text_style("subheading")
    assert text["size"] == pytest.approx(page.design["h"] * 0.04, abs=0.1) and text["bold"]
    page._set("curve", 120)
    page._set("stroke", "#000000")
    page._set("anim", "fade")
    assert (text["curve"], text["stroke"], text["anim"]) == (120, "#000000", "fade")
    monkeypatch.setattr(page.app.jarvis.brain, "ask_once", lambda *a, **k: "Fresh words")
    page.rewrite_selected("shorter")
    assert text["text"] == "Fresh words"
    page.undo()
    assert model.find(page.page, text["id"])["text"] == "Write something here"
    assert page.replace_all("Write", "Type") == 1
    assert model.find(page.page, text["id"])["text"] == "Type something here"
    page.open_find()
    page._find_dialog.find_entry.insert(0, "Visits")
    page._find_next()
    assert page.selected_id == graph["id"]
    page._find_dialog.destroy()
    page.set_transition("fade", all_pages=True)
    assert page.page["transition"] == "fade"
    page._refresh_page_tab()
    assert page.transition_menu.get() == "Fade"


@gui
def test_remove_background_and_emoji(page, monkeypatch, inline_workers, tmp_path):
    import ui.designui as designui
    from tests.test_design_gui import inline_workers_thread

    monkeypatch.setattr(designui.threading, "Thread", inline_workers_thread())
    picture = tmp_path / "product.png"
    image = Image.new("RGB", (200, 200), "white")
    image.paste((0, 120, 0), (60, 60, 140, 140))
    image.save(picture)
    el = page.add_image_file(str(picture))
    old = el["src"]
    page.remove_background()
    assert el["src"] != old and Image.open(model.resolve_src(el["src"])).getpixel((2, 2))[3] == 0
    page.pick_emoji()
    dialog = page._emoji_dialog
    page.update()
    count = len(page.page["elements"])
    page.add_icon("🎉", False)
    assert len(page.page["elements"]) == count + 1
    dialog.destroy()


@gui
def test_the_slideshow_and_presenter_view(page):
    page.load(templates.build_deck({"title": "Show", "slides": [{"title": "One", "bullets": ["a"], "notes": "Say hi"},
                                                                   {"title": "Two", "bullets": ["b"]}]}, "ocean"))
    page.go_to(1)
    title = next(e for e in page.page["elements"] if e.get("role") == "title")
    title["anim"] = "fade"
    page.design["pages"][2]["transition"] = "push"
    show = page.present(0, animate=False, screen=(0, 0, 480, 270))
    assert show.size == (480, 270) and show.index == 0
    show.next()
    assert (show.index, show.step) == (1, 0)
    show.next()
    assert show.step == 1                                                     # the title came in on a click
    show.next()
    assert show.index == 2
    show.next()
    assert show.ended
    show.prev()
    assert not show.ended and show.index == 2
    show._digit("1")
    show._enter()
    assert show.index == 0
    show.toggle_blank("black")
    assert show.blank == "black"
    show.close()
    assert show.closed
    talk = page.present(0, presenter=True, animate=False, screen=(0, 0, 480, 270))
    page.update()
    assert talk.presenter is not None and talk.counter_label.cget("text").startswith("Slide 1 / 3")
    talk.next()
    assert "Say hi" in talk.notes_view.get("1.0", "end")
    talk.toggle_timer()
    assert talk.paused_since is not None
    talk.close()
    started = []
    page.present = lambda *a, **k: started.append((a, k))
    page.open_target("presenter")
    page.after(250)
    page.update()
    assert started == [((0,), {"presenter": True})]
