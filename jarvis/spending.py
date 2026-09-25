"""A spending log: "I spent 45 on lunch" — totals by month and category.

Stored sealed like notes (vault), never sent to a model: the totals are
exact sums done here. Amounts in different currencies are totalled
separately rather than converted at a rate you didn't choose.
"""

from __future__ import annotations

import re
import time
from datetime import datetime, timedelta

from . import vault
from .config import get_data_dir, get_setting

FILE = "spending.json"
MAX_ENTRIES = 20000

CURRENCIES = {
    "$": "USD", "usd": "USD", "dollar": "USD", "dollars": "USD", "dolar": "USD",
    "€": "EUR", "eur": "EUR", "euro": "EUR", "euros": "EUR", "avro": "EUR",
    "£": "GBP", "gbp": "GBP", "pound": "GBP", "pounds": "GBP", "sterlin": "GBP",
    "₺": "TRY", "tl": "TRY", "try": "TRY", "lira": "TRY", "liras": "TRY", "tlira": "TRY",
}
SYMBOL = {"USD": "$", "EUR": "€", "GBP": "£", "TRY": "₺"}
CATEGORIES = {
    "food": "lunch dinner breakfast food restaurant cafe coffee kahve yemek pizza burger döner snack meal takeaway",
    "groceries": "groceries grocery market migros bim a101 şok carrefour supermarket bakkal",
    "transport": "taxi uber bus metro train fuel gas petrol benzin taksi otobüs ticket bilet istanbulkart parking",
    "bills": "rent bill bills electricity water internet phone fatura kira doğalgaz subscription netflix spotify",
    "shopping": "clothes shoes amazon trendyol hepsiburada shopping gift kıyafet",
    "health": "pharmacy medicine doctor eczane ilaç hospital dentist gym",
    "fun": "cinema movie game games steam concert bar party sinema oyun",
    "education": "book books course school kitap kurs udemy",
}


class SpendingError(Exception):
    pass


def _path():
    return get_data_dir() / FILE


def load() -> list[dict]:
    path = _path()
    if not path.exists():
        return []
    try:
        data = vault.read_json(path)
    except (OSError, ValueError):
        return []
    return [e for e in data if isinstance(e, dict) and "amount" in e] if isinstance(data, list) else []


def _save(entries: list[dict]) -> None:
    vault.write_json(_path(), entries[-MAX_ENTRIES:])


def default_currency(entries: list[dict] | None = None) -> str:
    chosen = (get_setting("JARVIS_CURRENCY", "") or "").strip().upper()
    if chosen:
        return chosen
    entries = load() if entries is None else entries
    if entries:
        return entries[-1].get("currency") or "TRY"
    return "TRY"


def guess_category(text: str) -> str:
    words = set(re.findall(r"\w+", text.lower()))
    words |= {w[:-1] for w in words if w.endswith("s") and len(w) > 3}   # "coffees", "taxis"
    for category, keywords in CATEGORIES.items():
        if words & set(keywords.split()):
            return category
    return "other"


def _number(text: str) -> float:
    """'1,250.50', '1.250,50', '1,250', '12,5' -> a float."""
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        return float(text.replace(thousands, "").replace(decimal, "."))
    for sep in ",.":
        if sep in text:
            head, _, tail = text.rpartition(sep)
            if len(tail) == 3:
                return float(text.replace(sep, ""))
            return float(head.replace(sep, "") + "." + tail)
    return float(text)


def parse(text: str, now: datetime | None = None) -> dict:
    """'45.50 tl on lunch #food yesterday' -> an entry. Raises SpendingError."""
    now = now or datetime.now()
    raw = " ".join(text.strip().split())
    when = now
    if re.search(r"\b(yesterday|dün)\b", raw, re.I):
        when = now - timedelta(days=1)
        raw = re.sub(r"\s*\b(yesterday|dün)\b", "", raw, flags=re.I).strip()
    tag = re.search(r"#(\w+)", raw)
    if tag:
        raw = raw.replace(tag.group(0), "").strip()
    found = list(re.finditer(r"([$€£₺])?\s*(\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?|\d+(?:[.,]\d{1,2})?)"
                             r"\s*([$€£₺]|[a-zA-Z]+\b)?", raw))
    # "2 coffees 9 dollars": the number next to a currency is the price.
    priced = [f for f in found if f.group(1) or (f.group(3) or "").lower() in CURRENCIES]
    m = (priced or found or [None])[0]
    if not m:
        raise SpendingError("How much? e.g. /spent 45 lunch  or  /spent 120 tl groceries")
    amount = _number(m.group(2))
    if amount <= 0 or amount > 10_000_000:
        raise SpendingError("That amount doesn't look right.")
    currency = None
    for token in (m.group(1), m.group(3)):
        if token and token.lower() in CURRENCIES:
            currency = CURRENCIES[token.lower()]
    rest = raw[:m.start()] + " " + raw[m.end():]
    if m.group(3) and m.group(3).lower() not in CURRENCIES:
        rest = raw[:m.start()] + " " + m.group(3) + " " + raw[m.end():]
    what = re.sub(r"^\s*(?:on|for|at|in|için)\s+", "", " ".join(rest.split()), flags=re.I).strip(" .,-")
    what = re.sub(r"\s+(?:for|on)$", "", what)
    return {
        "amount": round(amount, 2),
        "currency": currency,
        "what": what or "something",
        "category": (tag.group(1).lower() if tag else guess_category(what)),
        "at": when.timestamp(),
    }


def add(text: str) -> tuple[dict, list[dict]]:
    entry = parse(text)
    entries = load()
    entry["currency"] = entry["currency"] or default_currency(entries)
    entries.append(entry)
    _save(entries)
    return entry, entries


def money(amount: float, currency: str) -> str:
    symbol = SYMBOL.get(currency)
    number = f"{amount:,.2f}".replace(".00", "")
    return f"{symbol}{number}" if symbol else f"{number} {currency}"


def period(name: str, now: datetime | None = None) -> tuple[float, float, str]:
    """(start, end, label) for today | week | month | last month | year."""
    now = now or datetime.now()
    day = now.replace(hour=0, minute=0, second=0, microsecond=0)
    name = name.strip().lower()
    if name in {"today", "bugün"}:
        return day.timestamp(), now.timestamp() + 1, "today"
    if name in {"week", "this week", "bu hafta"}:
        start = day - timedelta(days=day.weekday())
        return start.timestamp(), now.timestamp() + 1, "this week"
    if name in {"last month", "geçen ay"}:
        first = day.replace(day=1)
        previous = (first - timedelta(days=1)).replace(day=1)
        return previous.timestamp(), first.timestamp(), f"{previous:%B %Y}"
    if name in {"year", "this year", "bu yıl"}:
        return day.replace(month=1, day=1).timestamp(), now.timestamp() + 1, f"{now:%Y}"
    first = day.replace(day=1)
    return first.timestamp(), now.timestamp() + 1, f"{now:%B %Y}"


def totals(entries: list[dict], start: float, end: float) -> dict[str, dict[str, float]]:
    """{currency: {category: total}}"""
    out: dict[str, dict[str, float]] = {}
    for e in entries:
        if start <= float(e.get("at", 0)) < end:
            bucket = out.setdefault(e.get("currency") or "TRY", {})
            bucket[e.get("category") or "other"] = bucket.get(e.get("category") or "other", 0) + float(e["amount"])
    return out


def budget() -> float:
    try:
        return float(get_setting("JARVIS_BUDGET", "") or 0)
    except ValueError:
        return 0.0


def summary(name: str = "month") -> str:
    entries = load()
    start, end, label = period(name)
    by_currency = totals(entries, start, end)
    if not by_currency:
        return f"Nothing logged for {label}. Say 'I spent 45 on lunch' or /spent 45 lunch."
    lines = [f"Spending — {label}"]
    for currency, cats in by_currency.items():
        total = sum(cats.values())
        lines.append(f"  Total: {money(total, currency)}")
        for category, amount in sorted(cats.items(), key=lambda kv: -kv[1]):
            share = amount / total * 100 if total else 0
            bar = "█" * max(1, round(share / 5))
            lines.append(f"    {category:<10} {money(amount, currency):>12}  {bar} {share:.0f}%")
    limit = budget()
    home = default_currency(entries)
    if limit and label == period("month")[2] and home in by_currency:
        used = sum(by_currency[home].values())
        lines.append(f"  Budget: {money(used, home)} of {money(limit, home)} ({used / limit * 100:.0f}%)")
    recent = [e for e in entries if start <= float(e.get("at", 0)) < end][-5:]
    if recent:
        lines.append("  Latest:")
        for e in reversed(recent):
            lines.append(f"    {datetime.fromtimestamp(e['at']):%d %b}  {money(e['amount'], e['currency'])}  {e['what']}")
    lines.append("/spent today|week|last month · /spent undo · /spent export · /spent budget 5000")
    return "\n".join(lines)


def budget_note(entries: list[dict], currency: str) -> str:
    limit = budget()
    if not limit or currency != default_currency(entries):
        return ""
    start, end, _ = period("month")
    used = sum(totals(entries, start, end).get(currency, {}).values())
    if used >= limit:
        return f"⚠ Over this month's budget: {money(used, currency)} of {money(limit, currency)}."
    if used >= limit * 0.8:
        return f"Heads up: {used / limit * 100:.0f}% of this month's budget used."
    return ""


def undo() -> str:
    entries = load()
    if not entries:
        return "Nothing to undo."
    last = entries.pop()
    _save(entries)
    return f"Removed {money(last['amount'], last['currency'])} {last['what']}."


def export(path) -> int:
    """Every entry to an Excel file, with a sheet of monthly totals."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise SpendingError("Excel export needs openpyxl (pip install openpyxl).")
    entries = load()
    book = Workbook()
    sheet = book.active
    sheet.title = "Spending"
    sheet.append(["Date", "Amount", "Currency", "Category", "What"])
    for e in entries:
        sheet.append([datetime.fromtimestamp(e["at"]).strftime("%Y-%m-%d %H:%M"), e["amount"],
                      e["currency"], e["category"], e["what"]])
    months = book.create_sheet("By month")
    months.append(["Month", "Currency", "Category", "Total"])
    grouped: dict[tuple, float] = {}
    for e in entries:
        key = (datetime.fromtimestamp(e["at"]).strftime("%Y-%m"), e["currency"], e["category"])
        grouped[key] = grouped.get(key, 0) + float(e["amount"])
    for key in sorted(grouped):
        months.append([*key, round(grouped[key], 2)])
    for ws in (sheet, months):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["E" if ws is sheet else "C"].width = 28
    book.save(str(path))
    return len(entries)


def stamp() -> str:
    return time.strftime("%Y%m%d_%H%M%S")
