"""10.0 documents and office: PDF tools, Excel from a description or plain
words, formulas, charts, mail merge, text checks, handouts, printing,
templates and recent documents.

PDFs are edited with pypdf, with overlays (page numbers, watermarks,
signatures) drawn by fpdf2 and merged on top, so the original text stays
text. Excel is written with openpyxl. Nothing here sends a document anywhere
unless the command is one that asks the AI, and those say so.
"""

from __future__ import annotations

import csv
import difflib
import io
import math
import os
import re
import shutil
import time
from collections import Counter
from pathlib import Path

from .. import kit, security, shield
from ..registry import command, field, split

G = "Documents"
PDF_TYPES = (("PDF", "*.pdf"),)
XL_TYPES = (("Spreadsheets", "*.xlsx *.xlsm *.csv"),)
DOC_TYPES = (("Word", "*.docx"),)


class OfficeError(Exception):
    pass


def out_path(source: Path, tag: str, suffix: str | None = None) -> Path:
    folder = kit.output_dir("documents")
    return folder / f"{source.stem}_{tag}_{kit.stamp()}{suffix or source.suffix}"


def need_file(text: str, *suffixes: str) -> Path:
    path = kit.path_arg(text)
    if path is None or not path.is_file():
        raise OfficeError(f"I can't find that file: {text.strip()[:120] or '(none given)'}")
    if suffixes and path.suffix.lower() not in suffixes:
        raise OfficeError(f"That needs a {' or '.join(suffixes)} file.")
    return path


def page_list(spec: str, count: int) -> list[int]:
    """'1-3,7' → [1, 2, 3, 7] (1-based, within 1..count). 'all' or '' → every page."""
    spec = (spec or "").strip().lower().replace(" ", "")
    if spec in {"", "all"}:
        return list(range(1, count + 1))
    pages: list[int] = []
    for part in spec.split(","):
        if not part:
            continue
        if "-" in part:
            a, _, b = part.partition("-")
            start = int(a) if a else 1
            end = int(b) if b else count
            step = 1 if end >= start else -1
            pages += list(range(start, end + step, step))
        elif part == "last":
            pages.append(count)
        else:
            pages.append(int(part))
    bad = [p for p in pages if not 1 <= p <= count]
    if bad:
        raise OfficeError(f"Page {bad[0]} doesn't exist — the PDF has {count} page(s).")
    return pages


# --- PDF ---------------------------------------------------------------------------

def pdf_reader(path: Path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise OfficeError("That PDF is password-protected.") from None
    return reader


def pdf_compress(path: Path, level: str = "medium") -> tuple[Path, int, int]:
    from pypdf import PdfWriter

    quality, longest = {"light": (85, 2400), "medium": (70, 1600), "strong": (50, 1100)}.get(level, (70, 1600))
    writer = PdfWriter(clone_from=str(path))
    for page in writer.pages:
        try:
            for image in page.images:
                picture = image.image
                if picture is None:
                    continue
                if max(picture.size) > longest:
                    picture.thumbnail((longest, longest))
                if picture.mode not in {"RGB", "L"}:
                    picture = picture.convert("RGB")
                image.replace(picture, quality=quality)
        except Exception:
            pass
        try:
            page.compress_content_streams()
        except Exception:
            pass
    try:
        writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
    except Exception:
        pass
    target = out_path(path, "small")
    with open(target, "wb") as handle:
        writer.write(handle)
    before, after = path.stat().st_size, target.stat().st_size
    if after >= before:
        shutil.copyfile(path, target)
        after = before
    return target, before, after


def pdf_rotate(path: Path, pages_spec: str, angle: int) -> Path:
    from pypdf import PdfWriter

    reader = pdf_reader(path)
    pages = set(page_list(pages_spec, len(reader.pages)))
    writer = PdfWriter()
    for number, page in enumerate(reader.pages, 1):
        if number in pages:
            page.rotate(angle)
        writer.add_page(page)
    target = out_path(path, "rotated")
    with open(target, "wb") as handle:
        writer.write(handle)
    return target


def pdf_arrange(path: Path, spec: str) -> tuple[Path, int]:
    """'3,1,2' reorders; 'delete 2,4' removes; 'keep 1-5' keeps only those."""
    from pypdf import PdfWriter

    reader = pdf_reader(path)
    count = len(reader.pages)
    spec = spec.strip().lower()
    if spec.startswith(("delete", "remove")):
        gone = set(page_list(spec.split(None, 1)[1] if " " in spec else "", count))
        order = [p for p in range(1, count + 1) if p not in gone]
    else:
        order = page_list(spec.replace("keep", "").strip(), count)
    if not order:
        raise OfficeError("That would leave no pages.")
    writer = PdfWriter()
    for number in order:
        writer.add_page(reader.pages[number - 1])
    target = out_path(path, "pages")
    with open(target, "wb") as handle:
        writer.write(handle)
    return target, len(order)


def _overlay(reader, draw) -> "PdfReader":
    """An fpdf2 PDF with one page per page of `reader`, same sizes, drawn by draw(pdf, number, w, h)."""
    from fpdf import FPDF
    from pypdf import PdfReader

    pdf = FPDF(unit="pt")
    pdf.set_auto_page_break(False)
    for number, page in enumerate(reader.pages, 1):
        box = page.mediabox
        w, h = float(box.width), float(box.height)
        rotation = (page.get("/Rotate") or 0) % 360
        if rotation in (90, 270):
            w, h = h, w
        pdf.add_page(format=(w, h))
        draw(pdf, number, w, h)
    return PdfReader(io.BytesIO(bytes(pdf.output())))


def _stamp(path: Path, overlay, tag: str, pages: set[int] | None = None) -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=pdf_reader(path))
    for number, page in enumerate(writer.pages, 1):
        if pages is None or number in pages:
            stamp = overlay.pages[number - 1]
            rotation = (page.get("/Rotate") or 0) % 360
            if rotation:
                page.transfer_rotation_to_content()
            page.merge_page(stamp)
    target = out_path(path, tag)
    with open(target, "wb") as handle:
        writer.write(handle)
    return target


def pdf_numbers(path: Path, where: str = "bottom-center", start: int = 1, style: str = "n") -> Path:
    reader = pdf_reader(path)
    total = len(reader.pages)

    def draw(pdf, number, w, h):
        pdf.set_font("helvetica", size=10)
        pdf.set_text_color(70, 70, 70)
        n = number + start - 1
        text = {"n": f"{n}", "of": f"{n} / {total + start - 1}", "page": f"Page {n} of {total + start - 1}",
                "sayfa": f"Sayfa {n} / {total + start - 1}"}.get(style, f"{n}")
        tw = pdf.get_string_width(text)
        x = {"left": 36, "right": w - 36 - tw}.get(where.split("-")[-1], (w - tw) / 2)
        y = 24 if where.startswith("top") else h - 30
        pdf.text(x, y + 10, text)

    return _stamp(path, _overlay(reader, draw), "numbered")


def pdf_watermark(path: Path, text: str, opacity: float = 0.15, size: int = 60) -> Path:
    reader = pdf_reader(path)
    text = text.strip()[:60] or "CONFIDENTIAL"

    def draw(pdf, number, w, h):
        pdf.set_font("helvetica", "B", size=size)
        pdf.set_text_color(120, 120, 120)
        tw = pdf.get_string_width(_latin(text))
        with pdf.local_context(fill_opacity=opacity, stroke_opacity=opacity):
            with pdf.rotation(angle=45, x=w / 2, y=h / 2):
                pdf.text(w / 2 - tw / 2, h / 2 + size / 3, _latin(text))

    return _stamp(path, _overlay(reader, draw), "watermarked")


def _latin(text: str) -> str:
    """fpdf2's core fonts are Latin-1: Turkish letters folded to their nearest."""
    return text.translate(str.maketrans("ğĞşŞıİ", "gGsSiI")).encode("latin-1", "replace").decode("latin-1")


def pdf_sign(path: Path, signature: Path, page: int = 0, position: str = "bottom-right", width_pt: float = 150,
             box: tuple[float, float] | None = None) -> Path:
    """Stamp a signature picture. `box` = (x, y) as fractions of the page, top-left of the signature."""
    from PIL import Image

    reader = pdf_reader(path)
    count = len(reader.pages)
    page = count if page in (0, -1) or page > count else page
    with Image.open(signature) as im:
        ratio = im.height / max(1, im.width)

    def draw(pdf, number, w, h):
        if number != page:
            return
        sw = min(width_pt, w * 0.6)
        sh = sw * ratio
        if box is not None:
            x, y = box[0] * w, box[1] * h
        else:
            x = {"left": 50, "center": (w - sw) / 2}.get(position.split("-")[-1], w - sw - 50)
            y = 60 if position.startswith("top") else h - sh - 70
        pdf.image(str(signature), x=x, y=y, w=sw, h=sh)

    return _stamp(path, _overlay(reader, draw), "signed", {page})


def pdf_fields(path: Path) -> dict[str, str]:
    reader = pdf_reader(path)
    fields = reader.get_fields() or {}
    out = {}
    for name, info in fields.items():
        value = info.get("/V", "")
        out[name] = "" if value is None else str(value)
    return out


def pdf_fill(path: Path, values: dict[str, str]) -> Path:
    from pypdf import PdfWriter
    from pypdf.generic import BooleanObject, NameObject

    reader = pdf_reader(path)
    writer = PdfWriter(clone_from=reader)
    known = set((reader.get_fields() or {}).keys())
    unknown = [k for k in values if k not in known]
    if unknown:
        raise OfficeError(f"No field called {unknown[0]!r}. Fields: {', '.join(sorted(known))[:300]}")
    for page in writer.pages:
        try:
            writer.update_page_form_field_values(page, values, auto_regenerate=False)
        except Exception:
            continue
    try:
        writer._root_object["/AcroForm"][NameObject("/NeedAppearances")] = BooleanObject(True)
    except Exception:
        pass
    target = out_path(path, "filled")
    with open(target, "wb") as handle:
        writer.write(handle)
    return target


def pdf_to_word(path: Path) -> tuple[Path, str]:
    """Word itself when installed (best layout); otherwise text and pictures."""
    target = out_path(path, "word", ".docx")
    if os.name == "nt":
        try:
            import comtypes.client

            word = comtypes.client.CreateObject("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            try:
                doc = word.Documents.Open(str(path.resolve()), False, True, False)
                doc.SaveAs2(str(target.resolve()), 16)
                doc.Close(False)
            finally:
                word.Quit()
            if target.exists():
                return target, "Microsoft Word"
        except Exception:
            pass
    from docx import Document
    from docx.shared import Inches

    reader = pdf_reader(path)
    document = Document()
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        for block in re.split(r"\n\s*\n", text):
            block = " ".join(block.split())
            if block:
                document.add_paragraph(block)
        try:
            for image in page.images[:6]:
                stream = io.BytesIO()
                image.image.convert("RGB").save(stream, "PNG")
                stream.seek(0)
                document.add_picture(stream, width=Inches(5))
        except Exception:
            pass
        if number < len(reader.pages):
            document.add_page_break()
    document.save(target)
    return target, "text and pictures (install Word for exact layout)"


# --- Excel -------------------------------------------------------------------------

XL_TR = {
    "SUM": "TOPLA", "IF": "EĞER", "VLOOKUP": "DÜŞEYARA", "HLOOKUP": "YATAYARA", "XLOOKUP": "ÇAPRAZARA",
    "COUNTIF": "EĞERSAY", "COUNTIFS": "ÇOKEĞERSAY", "SUMIF": "ETOPLA", "SUMIFS": "ÇOKETOPLA", "AVERAGE": "ORTALAMA",
    "AVERAGEIF": "EĞERORTALAMA", "IFERROR": "EĞERHATA", "INDEX": "İNDİS", "MATCH": "KAÇINCI", "CONCAT": "BİRLEŞTİR",
    "CONCATENATE": "BİRLEŞTİR", "TODAY": "BUGÜN", "NOW": "ŞİMDİ", "ROUND": "YUVARLA", "ROUNDUP": "YUKARIYUVARLA",
    "ROUNDDOWN": "AŞAĞIYUVARLA", "LEFT": "SOLDAN", "RIGHT": "SAĞDAN", "MID": "PARÇAAL", "LEN": "UZUNLUK",
    "TEXT": "METNEÇEVİR", "AND": "VE", "OR": "YADA", "NOT": "DEĞİL", "MAX": "MAK", "MIN": "MİN",
    "COUNT": "BAĞ_DEĞ_SAY", "COUNTA": "BAĞ_DEĞ_DOLU_SAY", "COUNTBLANK": "BOŞLUKSAY", "UNIQUE": "BENZERSİZ",
    "FILTER": "FİLTRE", "SORT": "SIRALA", "YEAR": "YIL", "MONTH": "AY", "DAY": "GÜN", "TRIM": "KIRP",
    "UPPER": "BÜYÜKHARF", "LOWER": "KÜÇÜKHARF", "PROPER": "YAZIM.DÜZENİ", "SUBSTITUTE": "YERİNEKOY", "ABS": "MUTLAK",
    "NETWORKDAYS": "TAMİŞGÜNÜ", "DATEDIF": "ETARİHLİ", "IFS": "ÇOKEĞER", "SUMPRODUCT": "TOPLA.ÇARPIM",
    "RANK": "RANK", "LARGE": "BÜYÜK", "SMALL": "KÜÇÜK", "WEEKDAY": "HAFTANINGÜNÜ", "EOMONTH": "SERİAY",
    "DATE": "TARİH", "VALUE": "SAYIYAÇEVİR", "PMT": "DEVRESEL_ÖDEME", "RAND": "S_SAYI_ÜRET",
    "RANDBETWEEN": "RASTGELEARADA", "SQRT": "KAREKÖK", "POWER": "KUVVET", "MOD": "MOD", "INT": "TAMSAYI",
}


def to_turkish_formula(formula: str) -> str:
    """=SUM(A1:A5, B2) → =TOPLA(A1:A5; B2): Turkish Excel's names and separators."""
    out = []
    in_string = False
    i = 0
    while i < len(formula):
        ch = formula[i]
        if ch == '"':
            in_string = not in_string
            out.append(ch)
            i += 1
            continue
        if not in_string:
            m = re.match(r"([A-Z][A-Z0-9.]*)\(", formula[i:])
            if m and (i == 0 or not formula[i - 1].isalnum()):
                out.append(XL_TR.get(m.group(1), m.group(1)) + "(")
                i += len(m.group(0))
                continue
            if ch == ",":
                out.append(";")
                i += 1
                continue
            if ch == "." and i > 0 and formula[i - 1].isdigit() and i + 1 < len(formula) and formula[i + 1].isdigit():
                out.append(",")
                i += 1
                continue
        out.append(ch)
        i += 1
    return "".join(out)


def read_table(path: Path, sheet: str = "", limit: int = 100000) -> tuple[list[str], list[list], str]:
    """(headers, rows, sheet name) from a .xlsx or .csv."""
    if path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        text = raw.decode("utf-8-sig", "replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
        rows = [[_cell(c) for c in r] for r in rows if any(str(c).strip() for c in r)][:limit + 1]
        return [str(h) for h in (rows[0] if rows else [])], rows[1:], path.stem
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in row):
            rows.append(list(row))
        if len(rows) > limit:
            break
    wb.close()
    headers = [str(h) if h is not None else f"Column {i + 1}" for i, h in enumerate(rows[0])] if rows else []
    width = len(headers)
    body = [(list(r) + [None] * width)[:width] for r in rows[1:]]
    return headers, body, ws.title


def _cell(value: str):
    text = str(value).strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+[.,]\d+", text):
        return float(text.replace(",", "."))
    return text


def col_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def write_table(path: Path, sheets: list[dict], accent: str = "1F6FEB") -> Path:
    """sheets: [{name, headers, rows, totals: [col names], formats: {col: fmt}}] → a styled workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        ws = wb.create_sheet(str(spec.get("name") or "Sheet")[:31])
        headers = [str(h) for h in spec.get("headers", [])]
        ws.append(headers)
        for row in spec.get("rows", []):
            ws.append([_formula_or_value(v, ws.max_row + 1) for v in row])
        for i, _h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=accent)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        totals = [t for t in spec.get("totals", []) if t in headers]
        last = ws.max_row
        if totals and last > 1:
            ws.append([""] * len(headers))
            total_row = ws.max_row
            ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
            for name in totals:
                col = headers.index(name) + 1
                letter = get_column_letter(col)
                cell = ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{last})")
                cell.font = Font(bold=True)
        for name, fmt in (spec.get("formats") or {}).items():
            if name in headers:
                col = headers.index(name) + 1
                for r in range(2, ws.max_row + 1):
                    ws.cell(row=r, column=col).number_format = str(fmt)
        for i, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or ""))
                                            for r in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(9, longest + 3))
        ws.freeze_panes = "A2"
        if headers:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    wb.save(path)
    return path


def _formula_or_value(value, row: int):
    if isinstance(value, str) and value.startswith("="):
        return value.replace("{r}", str(row))
    return value


def apply_ops(headers: list[str], rows: list[list], ops: list[dict]) -> tuple[list[str], list[list], list[str],
                                                                          dict, list[str]]:
    """Plain-words Excel edits, as data. Returns headers, rows, notes, formats, totals."""
    headers = list(headers)
    rows = [list(r) for r in rows]
    notes: list[str] = []
    formats: dict[str, str] = {}
    totals: list[str] = []

    def col(name) -> int:
        if name in headers:
            return headers.index(name)
        low = str(name).lower()
        for i, h in enumerate(headers):
            if h.lower() == low:
                return i
        raise OfficeError(f"No column called {name!r}.")

    def num(v):
        try:
            return float(str(v).replace(",", "."))
        except (TypeError, ValueError):
            return None

    for op in ops:
        kind = str(op.get("op", "")).lower()
        if kind == "add_column":
            name = str(op.get("name") or "New")
            formula = str(op.get("formula") or "")
            headers.append(name)
            for r in rows:
                r.append(formula if formula.startswith("=") else op.get("value", ""))
            notes.append(f"added column {name}")
        elif kind == "delete_column":
            i = col(op.get("name"))
            headers.pop(i)
            for r in rows:
                if i < len(r):
                    r.pop(i)
            notes.append(f"deleted column {op.get('name')}")
        elif kind == "rename_column":
            i = col(op.get("from"))
            headers[i] = str(op.get("to"))
            notes.append(f"renamed {op.get('from')} to {op.get('to')}")
        elif kind == "sort":
            i = col(op.get("by"))
            rows.sort(key=lambda r: (num(r[i]) is None, num(r[i]) if num(r[i]) is not None else str(r[i]).lower()),
                      reverse=bool(op.get("desc")))
            notes.append(f"sorted by {op.get('by')}{' (largest first)' if op.get('desc') else ''}")
        elif kind == "filter":
            i = col(op.get("column"))
            before = len(rows)
            if "equals" in op:
                rows = [r for r in rows if str(r[i]).strip().lower() == str(op["equals"]).strip().lower()]
            elif "contains" in op:
                rows = [r for r in rows if str(op["contains"]).lower() in str(r[i]).lower()]
            elif "gt" in op:
                rows = [r for r in rows if num(r[i]) is not None and num(r[i]) > float(op["gt"])]
            elif "lt" in op:
                rows = [r for r in rows if num(r[i]) is not None and num(r[i]) < float(op["lt"])]
            notes.append(f"kept {len(rows)} of {before} rows where {op.get('column')} matches")
        elif kind == "delete_duplicates":
            seen, kept = set(), []
            for r in rows:
                key = tuple(str(c) for c in r)
                if key not in seen:
                    seen.add(key)
                    kept.append(r)
            notes.append(f"removed {len(rows) - len(kept)} duplicate row(s)")
            rows = kept
        elif kind == "set":
            m = re.fullmatch(r"([A-Z]+)(\d+)", str(op.get("cell", "")).upper())
            if m:
                from openpyxl.utils import column_index_from_string

                c, r = column_index_from_string(m.group(1)) - 1, int(m.group(2))
                if r == 1:
                    while len(headers) <= c:
                        headers.append(f"Column {len(headers) + 1}")
                    headers[c] = str(op.get("value"))
                else:
                    while len(rows) < r - 1:
                        rows.append([None] * len(headers))
                    row = rows[r - 2]
                    while len(row) <= c:
                        row.append(None)
                    row[c] = op.get("value")
                notes.append(f"set {m.group(0)}")
        elif kind == "replace":
            i = col(op.get("column")) if op.get("column") else None
            count = 0
            for r in rows:
                for j, v in enumerate(r):
                    if (i is None or j == i) and isinstance(v, str) and str(op.get("find")) in v:
                        r[j] = v.replace(str(op.get("find")), str(op.get("with", "")))
                        count += 1
            notes.append(f"replaced {count} value(s)")
        elif kind == "format":
            name = headers[col(op.get("column"))]
            formats[name] = str(op.get("number_format") or "#,##0.00")
            notes.append(f"formatted {name}")
        elif kind == "total_row":
            for name in op.get("columns", []):
                totals.append(headers[col(name)])
            notes.append("added a total row")
        elif kind == "uppercase" or kind == "lowercase" or kind == "trim":
            i = col(op.get("column"))
            for r in rows:
                if isinstance(r[i], str):
                    r[i] = r[i].upper() if kind == "uppercase" else r[i].lower() if kind == "lowercase" else \
                        " ".join(r[i].split())
            notes.append(f"{kind} {op.get('column')}")
    return headers, rows, notes, formats, totals


EDIT_PROMPT = """You edit spreadsheets. Here are the column headers and the first rows of a sheet:
{sample}

The user wants: {request}

Answer with JSON only: {{"ops": [ ... ]}} using these operations (columns by header name; formulas use {{r}} for
the row number, e.g. "=B{{r}}*C{{r}}"):
  {{"op":"add_column","name":str,"formula":str}}   {{"op":"delete_column","name":str}}
  {{"op":"rename_column","from":str,"to":str}}     {{"op":"sort","by":str,"desc":bool}}
  {{"op":"filter","column":str,"equals"|"contains"|"gt"|"lt": value}}   {{"op":"delete_duplicates"}}
  {{"op":"replace","column":str or null,"find":str,"with":str}}   {{"op":"format","column":str,"number_format":str}}
  {{"op":"total_row","columns":[str]}}   {{"op":"uppercase"|"lowercase"|"trim","column":str}}
  {{"op":"set","cell":"B2","value":any}}"""


# --- text checks -----------------------------------------------------------------

def text_of(arg: str) -> tuple[str, str]:
    """(label, text) from a file path or the text itself."""
    path = kit.path_arg(arg.strip()) if len(arg.strip()) < 300 else None
    if path is not None and path.is_file():
        suffix = path.suffix.lower()
        if suffix == ".docx":
            from docx import Document

            return path.name, "\n".join(p.text for p in Document(str(path)).paragraphs)
        if suffix == ".pdf":
            return path.name, "\n".join((p.extract_text() or "") for p in pdf_reader(path).pages)
        return path.name, path.read_text(encoding="utf-8", errors="replace")
    return "your text", arg


def syllables_en(word: str) -> int:
    word = word.lower().strip(".,!?;:'\"")
    if len(word) <= 3:
        return 1
    word = re.sub(r"(?:es|ed|e)$", "", word)
    groups = re.findall(r"[aeiouy]+", word)
    return max(1, len(groups))


def syllables_tr(word: str) -> int:
    return max(1, len(re.findall(r"[aeıioöuüâîû]", word.lower())))


def is_turkish(text: str) -> bool:
    letters = sum(text.count(c) for c in "çğıöşüÇĞİÖŞÜ")
    common = len(re.findall(r"\b(ve|bir|bu|için|ile|da|de|ne|çok)\b", text.lower()))
    return letters + common * 2 > max(3, len(text) // 400)


def stats(text: str) -> dict:
    words = re.findall(r"[\wçğıöşüÇĞİÖŞÜ'’-]+", text)
    sentences = [s for s in re.split(r"(?<=[.!?…])\s+|\n{2,}", text) if re.search(r"\w", s)]
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    stop = {"the", "and", "a", "to", "of", "in", "is", "it", "that", "for", "on", "with", "as", "was", "are",
            "be", "this", "ve", "bir", "bu", "da", "de", "için", "ile", "çok", "ne", "o", "mi"}
    top = Counter(w.lower() for w in words if w.lower() not in stop and len(w) > 2).most_common(8)
    return {"words": len(words), "characters": len(text), "no_spaces": len(re.sub(r"\s", "", text)),
            "sentences": len(sentences), "paragraphs": len(paragraphs),
            "reading_min": len(words) / 238, "speaking_min": len(words) / 130, "top": top}


def readability(text: str) -> dict:
    words = re.findall(r"[\wçğıöşüÇĞİÖŞÜ'’]+", text)
    sentences = max(1, len([s for s in re.split(r"(?<=[.!?…])\s+", text) if re.search(r"\w", s)]))
    n = max(1, len(words))
    if is_turkish(text):
        syl = sum(syllables_tr(w) for w in words)
        score = 198.825 - 40.175 * (syl / n) - 2.610 * (n / sentences)   # Ateşman
        band = ("çok kolay" if score >= 90 else "kolay" if score >= 70 else "orta" if score >= 50 else
                "zor" if score >= 30 else "çok zor")
        return {"language": "Turkish (Ateşman)", "score": round(score, 1), "band": band,
                "words_per_sentence": round(n / sentences, 1), "syllables_per_word": round(syl / n, 2)}
    syl = sum(syllables_en(w) for w in words)
    ease = 206.835 - 1.015 * (n / sentences) - 84.6 * (syl / n)
    grade = 0.39 * (n / sentences) + 11.8 * (syl / n) - 15.59
    band = ("very easy" if ease >= 90 else "easy" if ease >= 70 else "plain English" if ease >= 60 else
            "fairly hard" if ease >= 50 else "hard" if ease >= 30 else "very hard")
    return {"language": "English (Flesch)", "score": round(ease, 1), "band": band, "grade": round(grade, 1),
            "words_per_sentence": round(n / sentences, 1), "syllables_per_word": round(syl / n, 2)}


def similarity(a: str, b: str, run: int = 8) -> dict:
    wa = re.findall(r"\w+", a.lower())
    wb = re.findall(r"\w+", b.lower())
    ratio = difflib.SequenceMatcher(None, wa, wb, autojunk=False).ratio() if wa and wb else 0.0

    def shingles(words):
        return {" ".join(words[i:i + 5]) for i in range(max(0, len(words) - 4))}

    sa, sb = shingles(wa), shingles(wb)
    jaccard = len(sa & sb) / len(sa | sb) if sa | sb else 0.0
    matcher = difflib.SequenceMatcher(None, wa, wb, autojunk=False)
    shared = [" ".join(wa[m.a:m.a + m.size]) for m in matcher.get_matching_blocks() if m.size >= run]
    return {"ratio": ratio, "overlap": jaccard, "shared": shared}


# --- mail merge ------------------------------------------------------------------

PLACEHOLDER = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")


def _replace_in_paragraph(paragraph, values: dict) -> None:
    text = "".join(run.text for run in paragraph.runs)
    if "{{" not in text:
        return
    new = PLACEHOLDER.sub(lambda m: str(values.get(m.group(1), values.get(m.group(1).lower(), m.group(0)))), text)
    if new == text or not paragraph.runs:
        return
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""


def _merge_document(document, values: dict) -> None:
    for paragraph in document.paragraphs:
        _replace_in_paragraph(paragraph, values)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, values)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                _replace_in_paragraph(paragraph, values)


def mail_merge(template: Path, data: Path) -> tuple[Path, Path, int]:
    from docx import Document

    headers, rows, _ = read_table(data)
    if not rows:
        raise OfficeError("The data file has no rows under its headers.")
    folder = kit.output_dir("documents") / f"{template.stem}_merged_{kit.stamp()}"
    folder.mkdir(parents=True, exist_ok=True)
    placeholders = set(PLACEHOLDER.findall("\n".join(p.text for p in Document(str(template)).paragraphs)))
    missing = [p for p in placeholders if p not in headers and p.lower() not in {h.lower() for h in headers}]
    if missing:
        raise OfficeError(f"The template uses {{{{{missing[0]}}}}} but the data has no such column. "
                          f"Columns: {', '.join(headers)}")
    combined = Document(str(template))
    body = combined.element.body
    for child in list(body):
        if not child.tag.endswith("sectPr"):
            body.remove(child)
    for n, row in enumerate(rows, 1):
        values = {h: ("" if v is None else v) for h, v in zip(headers, row)}
        values.update({h.lower(): v for h, v in values.items()})
        single = Document(str(template))
        _merge_document(single, values)
        name = kit.slug(str(row[0]) if row and row[0] else f"letter-{n}")
        single.save(folder / f"{n:03d}_{name}.docx")
        for child in list(single.element.body):
            if not child.tag.endswith("sectPr"):
                body.insert(len(body) - 1, child)
        if n < len(rows):
            from docx.enum.text import WD_BREAK

            paragraph = combined.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            body.insert(len(body) - 1, paragraph._p)
    combined_path = folder / f"{template.stem}_all.docx"
    combined.save(combined_path)
    return folder, combined_path, len(rows)


# --- email templates ---------------------------------------------------------------

EMAIL_TEMPLATES = {
    "thank-you": ("Thank you", "Subject: Thank you, {name}\n\nHi {name},\n\nThank you for {reason}. {detail}\n\n"
                               "Best regards,\n{me}"),
    "follow-up": ("Follow-up", "Subject: Following up on {topic}\n\nHi {name},\n\nI wanted to follow up on {topic} "
                               "from {when}. {detail}\n\nCould you let me know {question}?\n\nThanks,\n{me}"),
    "meeting-request": ("Meeting request", "Subject: Could we meet about {topic}?\n\nHi {name},\n\nWould you have "
                                           "{length} for a quick meeting about {topic}? I'm free {times}.\n\n"
                                           "Thanks,\n{me}"),
    "apology": ("Apology", "Subject: My apologies\n\nHi {name},\n\nI'm sorry for {what}. {fix}\n\nThank you for your "
                           "patience,\n{me}"),
    "complaint": ("Complaint", "Subject: Problem with {product}\n\nHello,\n\nOn {when} I {what}. Unfortunately "
                               "{problem}. I would like {wanted}.\n\nOrder/reference: {reference}\n\nRegards,\n{me}"),
    "leave-request": ("Leave request", "Subject: Leave request — {dates}\n\nHi {name},\n\nI'd like to request leave "
                                       "from {dates} for {reason}. {cover}\n\nThank you,\n{me}"),
    "job-follow-up": ("Job application follow-up", "Subject: Application for {role}\n\nDear {name},\n\nI applied for "
                                                   "the {role} position on {when} and wanted to confirm my continued "
                                                   "interest. {detail}\n\nKind regards,\n{me}"),
    "payment-reminder": ("Payment reminder", "Subject: Invoice {number} — friendly reminder\n\nHi {name},\n\nThis is a "
                                             "friendly reminder that invoice {number} for {amount} was due on {due}. "
                                             "{detail}\n\nThank you,\n{me}"),
    "introduction": ("Introduction", "Subject: Introducing {person}\n\nHi {name},\n\nI'd like to introduce {person}, "
                                     "{who}. {why}\n\nI'll let you two take it from here.\n\nBest,\n{me}"),
    "tesekkur": ("Teşekkür (TR)", "Konu: Teşekkürler {name}\n\nMerhaba {name},\n\n{reason} için çok teşekkür ederim. "
                                  "{detail}\n\nSaygılarımla,\n{me}"),
    "izin-talebi": ("İzin talebi (TR)", "Konu: İzin talebi — {dates}\n\nSayın {name},\n\n{dates} tarihleri arasında "
                                        "{reason} nedeniyle izin kullanmak istiyorum. {cover}\n\nGereğini bilgilerinize "
                                        "arz ederim.\n{me}"),
    "toplanti": ("Toplantı talebi (TR)", "Konu: {topic} hakkında görüşme\n\nMerhaba {name},\n\n{topic} hakkında {length} "
                                         "kısa bir görüşme yapabilir miyiz? {times} uygunum.\n\nTeşekkürler,\n{me}"),
}
CUSTOM_TEMPLATES = kit.Store("email_templates.json", {})


def all_email_templates() -> dict[str, tuple[str, str]]:
    found = dict(EMAIL_TEMPLATES)
    for key, body in CUSTOM_TEMPLATES.load().items():
        found[key] = (key.replace("-", " ").capitalize(), body)
    return found


def template_fields(body: str) -> list[str]:
    seen: list[str] = []
    for name in re.findall(r"\{(\w+)\}", body):
        if name not in seen:
            seen.append(name)
    return seen


# --- document templates --------------------------------------------------------------

DOC_TEMPLATES = {
    "letter": ("Formal letter", "a formal letter with the sender's address, date, recipient, subject line, a "
                                "clear body and a sign-off"),
    "report": ("Report", "a structured report: title, executive summary, background, findings with a table, "
                         "recommendations and next steps"),
    "agenda": ("Meeting agenda", "a meeting agenda: title, date/time/place, attendees, numbered agenda items with "
                                 "owners and times, and a notes section"),
    "minutes": ("Meeting minutes", "meeting minutes: attendees, summary, decisions, an action-items table "
                                   "(what, who, when)"),
    "proposal": ("Project proposal", "a project proposal: problem, proposed solution, scope, timeline table, budget "
                                     "table, risks, and a call to action"),
    "memo": ("Memo", "an internal memo: To, From, Date, Subject, then a short body with bullet points"),
    "recommendation": ("Recommendation letter", "a recommendation letter: relationship, strengths with examples, "
                                                "a clear endorsement"),
    "business-plan": ("Business plan (short)", "a short business plan: summary, problem, solution, market, "
                                               "competition table, revenue model, plan, financial table"),
    "lesson-plan": ("Lesson plan", "a lesson plan: objectives, materials, timed activities table, assessment, "
                                   "homework"),
    "resignation": ("Resignation letter", "a polite resignation letter with last working day and handover offer"),
    "cover-page": ("Cover page", "a cover page: title, subtitle, author, course or company, date"),
    "dilekce": ("Dilekçe", "a formal Turkish dilekçe: addressee in capitals, body in formal Turkish, date, name "
                           "and signature lines, address and attachments (ekler)"),
}


# --- recent documents ----------------------------------------------------------------

DOC_SUFFIXES = {".docx", ".doc", ".pdf", ".xlsx", ".xls", ".csv", ".pptx", ".ppt", ".txt", ".md", ".odt", ".rtf"}


def recent_documents(limit: int = 40) -> list[tuple[float, Path]]:
    found: dict[str, tuple[float, Path]] = {}
    recent = Path(os.environ.get("APPDATA", "")) / "Microsoft" / "Windows" / "Recent"
    if recent.is_dir():
        for link in recent.glob("*.lnk"):
            target_name = link.stem
            if Path(target_name).suffix.lower() in DOC_SUFFIXES:
                found[target_name.lower()] = (link.stat().st_mtime, link)
    for folder in (kit.output_dir("documents"), kit.output_dir()):
        if folder.is_dir():
            for path in folder.iterdir():
                if path.suffix.lower() in DOC_SUFFIXES:
                    found[path.name.lower()] = (path.stat().st_mtime, path)
    return sorted(found.values(), key=lambda item: -item[0])[:limit]


def resolve_link(path: Path) -> Path:
    """The file a .lnk points to (or the path itself)."""
    if path.suffix.lower() != ".lnk" or os.name != "nt":
        return path
    # pywin32, not comtypes: comtypes hands back a bare IDispatch for WScript.Shell, without
    # .TargetPath, so every shortcut came back unresolved. Each thread needs COM set up.
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return path
    started = False
    try:
        pythoncom.CoInitialize()
        started = True
    except Exception:
        pass
    try:
        target = win32com.client.Dispatch("WScript.Shell").CreateShortcut(str(path)).TargetPath
        return Path(target) if target else path
    except Exception:
        return path
    finally:
        if started:
            pythoncom.CoUninitialize()


class Office:
    # --- PDF -------------------------------------------------------------------------
    @command("pdfcompress", group=G, usage="/pdfcompress <file.pdf> [light|medium|strong]",
             help="make a PDF smaller (shrinks big pictures)", title="Compress a PDF", icon="🗜", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES),
                     field("level", "choice", "How much", "medium", ("light", "medium", "strong"))),
             template="{file} | {level}", keywords="smaller reduce size shrink küçült")
    def pdf_compress_cmd(self, args: str, routed: bool = False):
        path_text, level = split(args, 2)
        if not level and path_text.rsplit(" ", 1)[-1].lower() in {"light", "medium", "strong"}:
            path_text, level = path_text.rsplit(" ", 1)
        try:
            path = need_file(path_text, ".pdf")
            target, before, after = pdf_compress(path, (level or "medium").lower())
        except OfficeError as exc:
            return str(exc)
        saved = 100 * (before - after) / before if before else 0
        if after >= before:
            return f"That PDF is already as small as I can make it ({kit.size(before)})."
        return f"🗜 {kit.size(before)} → {kit.size(after)} ({saved:.0f}% smaller)\n  {target}"

    @command("pdf2word", "pdftoword", group=G, usage="/pdf2word <file.pdf>", help="turn a PDF into an editable Word file",
             title="PDF to Word", icon="📝", page="documents", fields=(field("file", "file", "PDF", types=PDF_TYPES),),
             keywords="convert docx edit pdf")
    def pdf_to_word_cmd(self, args: str, routed: bool = False):
        try:
            target, how = pdf_to_word(need_file(args, ".pdf"))
        except OfficeError as exc:
            return str(exc)
        return f"📝 Word file ({how}):\n  {target}"

    @command("pdfrotate", group=G, usage="/pdfrotate <file.pdf> | <pages: all, 1-3,7> | <90|180|270>",
             help="rotate PDF pages", title="Rotate PDF pages", icon="🔄", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES), field("pages", "text", "Pages", "all"),
                     field("angle", "choice", "Turn by", "90", ("90", "180", "270"))))
    def pdf_rotate_cmd(self, args: str, routed: bool = False):
        path_text, pages, angle = split(args, 3)
        try:
            target = pdf_rotate(need_file(path_text, ".pdf"), pages or "all", int(angle or 90))
        except (OfficeError, ValueError) as exc:
            return str(exc)
        return f"🔄 Rotated:\n  {target}"

    @command("pdfpages", group=G, usage="/pdfpages <file.pdf> | <3,1,2> · delete 2,4 · keep 1-5",
             help="reorder, delete or keep PDF pages", title="Reorder or delete PDF pages", icon="📑", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES),
                     field("order", "text", "New order, or delete 2,4", hint="3,1,2,4-6  ·  delete 2,4  ·  keep 1-5")))
    def pdf_pages_cmd(self, args: str, routed: bool = False):
        path_text, spec = split(args, 2)
        try:
            target, count = pdf_arrange(need_file(path_text, ".pdf"), spec)
        except (OfficeError, ValueError) as exc:
            return str(exc)
        return f"📑 {count} page(s):\n  {target}"

    @command("pdfnumber", "pagenumbers", group=G, usage="/pdfnumber <file.pdf> | [bottom-center] | [n|of|page|sayfa] | [start]",
             help="add page numbers to a PDF", title="PDF page numbers", icon="🔢", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES),
                     field("where", "choice", "Where", "bottom-center",
                           ("bottom-center", "bottom-right", "bottom-left", "top-center", "top-right")),
                     field("style", "choice", "Style", "of", ("n", "of", "page", "sayfa")),
                     field("start", "number", "Start at", "1")))
    def pdf_number_cmd(self, args: str, routed: bool = False):
        path_text, where, style, start = split(args, 4)
        try:
            target = pdf_numbers(need_file(path_text, ".pdf"), where or "bottom-center", int(start or 1), style or "of")
        except (OfficeError, ValueError) as exc:
            return str(exc)
        return f"🔢 Numbered:\n  {target}"

    @command("pdfwatermark", group=G, usage="/pdfwatermark <file.pdf> | <text> | [opacity 0.15]",
             help="a diagonal watermark on every page", title="PDF watermark", icon="💧", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES), field("text", "text", "Watermark", "CONFIDENTIAL"),
                     field("opacity", "choice", "Strength", "0.15", ("0.08", "0.15", "0.3"))))
    def pdf_watermark_cmd(self, args: str, routed: bool = False):
        path_text, text, opacity = split(args, 3)
        try:
            target = pdf_watermark(need_file(path_text, ".pdf"), text or "CONFIDENTIAL", float(opacity or 0.15))
        except (OfficeError, ValueError) as exc:
            return str(exc)
        return f"💧 Watermarked:\n  {target}"

    @command("pdfsign", group=G, usage="/pdfsign <file.pdf> | <signature.png> | [page, last] | [bottom-right]",
             help="put your signature picture on a PDF (draw one on the Documents page)", title="Sign a PDF",
             icon="✍", page="documents",
             fields=(field("file", "file", "PDF", types=PDF_TYPES),
                     field("signature", "file", "Signature picture", types=(("Pictures", "*.png *.jpg"),)),
                     field("page", "text", "Page", "last"),
                     field("where", "choice", "Where", "bottom-right", ("bottom-right", "bottom-left", "bottom-center",
                                                                       "top-right"))))
    def pdf_sign_cmd(self, args: str, routed: bool = False):
        path_text, sig_text, page, where = split(args, 4)
        try:
            path = need_file(path_text, ".pdf")
            sig = need_file(sig_text, ".png", ".jpg", ".jpeg")
            number = 0 if page.strip().lower() in {"", "last"} else int(page)
            target = pdf_sign(path, sig, number, where or "bottom-right")
        except (OfficeError, ValueError) as exc:
            return str(exc)
        security.audit.record("pdfsign", path.name)
        return f"✍ Signed:\n  {target}"

    @command("pdfform", group=G, usage="/pdfform <file.pdf> · /pdfform <file.pdf> | name=Ali; date=01.10.2026",
             help="see and fill a PDF form's fields", title="Fill a PDF form", icon="🧾", page="documents",
             fields=(field("file", "file", "PDF form", types=PDF_TYPES),
                     field("values", "long", "field=value; field=value (empty: list the fields)", optional=True)))
    def pdf_form_cmd(self, args: str, routed: bool = False):
        path_text, values = split(args, 2)
        try:
            path = need_file(path_text, ".pdf")
            if not values.strip():
                fields = pdf_fields(path)
                if not fields:
                    return "That PDF has no fillable fields."
                return "🧾 Fields:\n" + "\n".join(f"  {k} = {v}" for k, v in fields.items()) + \
                    "\nFill with: /pdfform <file> | field=value; field=value"
            pairs = {}
            for part in re.split(r";|\n", values):
                if "=" in part:
                    k, _, v = part.partition("=")
                    pairs[k.strip()] = v.strip()
            target = pdf_fill(path, pairs)
        except OfficeError as exc:
            return str(exc)
        return f"🧾 Filled {len(pairs)} field(s):\n  {target}"

    # --- Excel -----------------------------------------------------------------------
    @command("excel", "spreadsheet", group=G, usage="/excel <describe the spreadsheet>",
             help="a ready-to-use Excel file from a description", title="Excel from a description", icon="📗",
             page="documents", fields=(field("what", "long", "Describe the spreadsheet",
                                             hint="a monthly budget with categories, planned vs actual"),),
             keywords="xlsx sheet table budget tracker")
    def excel_cmd(self, args: str, routed: bool = False):
        request = args.strip()
        if not request:
            return "Usage: /excel <what the spreadsheet is for>   e.g. /excel a weekly study timetable"
        with kit.more_room(self.brain):
            data = kit.ask_json(self.brain, (
                "Design a practical Excel workbook for this request. JSON only: {\"title\": str, \"sheets\": [{\"name\": "
                "str, \"headers\": [str], \"rows\": [[values]], \"totals\": [header names to sum], \"formats\": "
                "{header: excel number format}}]}. Use realistic example rows (8-20), numbers as numbers, and Excel "
                "formulas as strings starting with '=' using {r} for the current row (e.g. \"=C{r}-B{r}\").\n\n"
                f"Request: {request}"))
        if not isinstance(data, dict) or not data.get("sheets"):
            return "The AI didn't return a usable spreadsheet — try describing it differently."
        path = kit.output_dir("documents") / f"{kit.slug(data.get('title') or request)}_{kit.stamp()}.xlsx"
        write_table(path, data["sheets"])
        sheets = ", ".join(f"{s.get('name')} ({len(s.get('rows', []))} rows)" for s in data["sheets"])
        return f"📗 {data.get('title') or 'Spreadsheet'} — {sheets}\n  {path}"

    @command("exceledit", "editexcel", group=G, usage="/exceledit <file.xlsx|csv> | <what to change>",
             help="change a spreadsheet in plain words (saves a copy)", title="Edit Excel in plain words", icon="✏",
             page="documents", fields=(field("file", "file", "Spreadsheet", types=XL_TYPES),
                                       field("change", "long", "What to change",
                                             hint="sort by total, largest first, and add a VAT column at 20%")))
    def excel_edit_cmd(self, args: str, routed: bool = False):
        path_text, request = split(args, 2)
        try:
            path = need_file(path_text, ".xlsx", ".xlsm", ".csv")
            headers, rows, sheet = read_table(path)
        except OfficeError as exc:
            return str(exc)
        if not request:
            return "What should change? /exceledit <file> | <what to change>"
        sample = " | ".join(headers) + "\n" + "\n".join(" | ".join("" if c is None else str(c) for c in r)
                                                       for r in rows[:15])
        wrapped, _ = shield.wrap(sample, "the spreadsheet")
        data = kit.ask_json(self.brain, EDIT_PROMPT.format(sample=wrapped, request=request))
        ops = data.get("ops") if isinstance(data, dict) else None
        if not ops:
            return "I couldn't turn that into spreadsheet changes — try saying it more specifically."
        try:
            headers, rows, notes, formats, totals = apply_ops(headers, rows, ops)
        except OfficeError as exc:
            return f"{exc} The columns are: {', '.join(headers)}"
        target = out_path(path, "edited", ".xlsx")
        write_table(target, [{"name": sheet, "headers": headers, "rows": rows, "formats": formats, "totals": totals}])
        return "✏ " + "; ".join(notes) + f"\n  {target}"

    @command("formula", group=G, usage="/formula <what you want the cell to do>",
             help="the Excel formula for it, in English and Turkish Excel", title="Excel formula helper", icon="ƒ",
             page="documents", fields=(field("need", "long", "What should the formula do?",
                                             hint="total of column C where column A says Paid"),),
             keywords="excel function vlookup sumif formül")
    def formula_cmd(self, args: str, routed: bool = False):
        need = args.strip()
        if not need:
            return "Usage: /formula <what the cell should do>"
        data = kit.ask_json(self.brain, (
            "Write the Excel formula for this. JSON only: {\"formula\": str (English function names, comma "
            "separators, starting with =), \"explain\": str (2-3 short sentences), \"example\": str}.\n\n" + need))
        if not isinstance(data, dict) or not str(data.get("formula", "")).startswith("="):
            return "No formula came back — try describing the columns you have."
        english = data["formula"].strip()
        return (f"ƒ {english}\n   Turkish Excel: {to_turkish_formula(english)}\n\n{data.get('explain', '').strip()}"
                + (f"\n\nExample: {data['example']}" if data.get("example") else ""))

    @command("xlchart", "excelchart", group=G, usage="/xlchart <file.xlsx|csv> | <column|bar|line|pie> | [label column] | [value columns]",
             help="a chart from spreadsheet data — in Excel and as a picture", title="Chart from Excel data",
             icon="📊", page="documents",
             fields=(field("file", "file", "Spreadsheet", types=XL_TYPES),
                     field("kind", "choice", "Chart", "column", ("column", "bar", "line", "area", "pie", "doughnut")),
                     field("labels", "text", "Label column (blank: first)", optional=True),
                     field("values", "text", "Value columns, comma-separated (blank: all numbers)", optional=True)))
    def xlchart_cmd(self, args: str, routed: bool = False):
        from openpyxl import load_workbook
        from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference

        from .. import charts

        path_text, kind, label_col, value_cols = split(args, 4)
        kind = (kind or "column").lower()
        try:
            path = need_file(path_text, ".xlsx", ".xlsm", ".csv")
            headers, rows, sheet = read_table(path, limit=500)
        except OfficeError as exc:
            return str(exc)
        if not rows:
            return "That sheet has no data rows."
        li = headers.index(label_col) if label_col in headers else 0
        wanted = [v.strip() for v in value_cols.split(",") if v.strip()] if value_cols else [
            h for i, h in enumerate(headers) if i != li and all(isinstance(r[i], (int, float)) or r[i] is None
                                                                for r in rows[:50]) and any(
                isinstance(r[i], (int, float)) for r in rows[:50])]
        missing = [w for w in wanted if w not in headers]
        if missing or not wanted:
            return f"No number columns found{' called ' + missing[0] if missing else ''}. Columns: {', '.join(headers)}"
        labels = [str(r[li]) for r in rows]
        series = {w: [float(r[headers.index(w)] or 0) if isinstance(r[headers.index(w)], (int, float)) else 0.0
                      for r in rows] for w in wanted}
        if kind in {"pie", "doughnut"}:
            series = {wanted[0]: series[wanted[0]]}
        image = charts.render(kind, labels, series, title=Path(path).stem)
        png = out_path(path, "chart", ".png")
        image.save(png)
        target = out_path(path, "chart", ".xlsx")
        write_table(target, [{"name": sheet, "headers": headers, "rows": rows}])
        wb = load_workbook(target)
        ws = wb.active
        n = len(rows) + 1
        chart = {"column": BarChart, "bar": BarChart, "line": LineChart, "area": AreaChart, "pie": PieChart,
                 "doughnut": DoughnutChart}[kind]()
        if kind in {"column", "bar"}:
            chart.type = "col" if kind == "column" else "bar"
        chart.title = Path(path).stem
        for w in series:
            col = headers.index(w) + 1
            chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=n), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=li + 1, min_row=2, max_row=n))
        chart.width, chart.height = 22, 12
        ws.add_chart(chart, f"{col_letter(len(headers) + 1)}2")
        wb.save(target)
        return f"📊 {kind.capitalize()} chart of {', '.join(series)}:\n  {target}\n  {png}"

    # --- text checks -----------------------------------------------------------------
    @command("wordcount", "words", group=G, usage="/wordcount <text or file>",
             help="words, characters, sentences, reading and speaking time", title="Word count and reading time",
             icon="🔢", page="documents", fields=(field("text", "long", "Text, or a file path"),),
             keywords="count characters karakter kelime sayısı")
    def wordcount_cmd(self, args: str, routed: bool = False):
        if not args.strip():
            return "Usage: /wordcount <text or a file>"
        name, text = text_of(args)
        s = stats(text)
        top = ", ".join(f"{w} ({n})" for w, n in s["top"])
        return (f"🔢 {name}: {s['words']:,} words · {s['characters']:,} characters ({s['no_spaces']:,} without spaces)\n"
                f"   {s['sentences']:,} sentences · {s['paragraphs']:,} paragraphs\n"
                f"   Reading ≈ {_minutes(s['reading_min'])} · speaking ≈ {_minutes(s['speaking_min'])}\n"
                + (f"   Most used: {top}" if top else ""))

    @command("readability", group=G, usage="/readability <text or file>",
             help="how easy a text is to read (English or Turkish)", title="Readability score", icon="📖",
             page="documents", fields=(field("text", "long", "Text, or a file path"),),
             keywords="flesch okunabilirlik ateşman")
    def readability_cmd(self, args: str, routed: bool = False):
        if not args.strip():
            return "Usage: /readability <text or a file>"
        name, text = text_of(args)
        r = readability(text)
        extra = f", grade level {r['grade']}" if "grade" in r else ""
        tip = ("Shorter sentences and shorter words raise the score." if r["score"] < 60 else
               "Easy to read — good for a general audience.")
        return (f"📖 {name} — {r['language']}: {r['score']} ({r['band']}{extra})\n"
                f"   {r['words_per_sentence']} words per sentence · {r['syllables_per_word']} syllables per word\n   {tip}")

    @command("similarity", "compare-texts", group=G, usage="/similarity <text or file> | <text or file>",
             help="how alike two texts are, and the passages they share", title="Similarity of two texts", icon="🪞",
             page="documents", fields=(field("first", "long", "First text or file"),
                                       field("second", "long", "Second text or file")),
             keywords="plagiarism copy compare benzerlik")
    def similarity_cmd(self, args: str, routed: bool = False):
        first, second = split(args, 2)
        if not first or not second:
            return "Usage: /similarity <text or file> | <text or file>"
        (na, a), (nb, b) = text_of(first), text_of(second)
        r = similarity(a, b)
        lines = [f"🪞 {na} vs {nb}: {100 * r['ratio']:.0f}% alike in wording, {100 * r['overlap']:.0f}% of 5-word "
                 "phrases shared."]
        if r["shared"]:
            lines.append("Shared passages (8+ words):")
            lines += [f"  “{s[:160]}”" for s in r["shared"][:8]]
        else:
            lines.append("No long passages in common.")
        return "\n".join(lines)

    # --- Word --------------------------------------------------------------------------
    @command("mailmerge", group=G, usage="/mailmerge <template.docx> | <data.xlsx|csv>",
             help="one letter per row: {{Name}} in the template is filled from the column 'Name'",
             title="Mail merge letters", icon="📨", page="documents",
             fields=(field("template", "file", "Word template with {{Column}} placeholders", types=DOC_TYPES),
                     field("data", "file", "Data (Excel or CSV)", types=XL_TYPES)))
    def mailmerge_cmd(self, args: str, routed: bool = False):
        template_text, data_text = split(args, 2)
        try:
            template = need_file(template_text, ".docx")
            data = need_file(data_text, ".xlsx", ".xlsm", ".csv")
            folder, combined, count = mail_merge(template, data)
        except OfficeError as exc:
            return str(exc)
        return f"📨 {count} letters:\n  {combined}\n  (and one file each in {folder})"

    @command("emailtpl", "emailtemplates", group=G, usage="/emailtpl [name] | key=value; … · /emailtpl save <name> | <text with {fields}>",
             help="ready email templates (English and Turkish) to fill in", title="Email templates", icon="📧",
             page="documents", fields=(field("name", "choice", "Template", "follow-up", tuple(EMAIL_TEMPLATES)),
                                       field("values", "long", "name=Ayşe; topic=the budget; me=Ahmet",
                                             optional=True)))
    def emailtpl_cmd(self, args: str, routed: bool = False):
        text = args.strip()
        templates = all_email_templates()
        if text.lower().startswith("save "):
            name, body = split(text[5:], 2)
            if not name or not body:
                return "Usage: /emailtpl save <name> | <text with {fields}>"
            custom = CUSTOM_TEMPLATES.load()
            custom[kit.slug(name).lower()] = body
            CUSTOM_TEMPLATES.save(custom)
            return f"📧 Saved the template {name} ({', '.join(template_fields(body)) or 'no fields'})."
        name, values = split(text, 2)
        if not name or name.lower() not in templates:
            return "📧 Email templates:\n" + "\n".join(f"  {k} — {v[0]}  ({', '.join(template_fields(v[1]))})"
                                                       for k, v in templates.items()) + \
                "\nFill one: /emailtpl follow-up | name=Ayşe; topic=the budget; me=Ahmet"
        body = templates[name.lower()][1]
        pairs = {}
        for part in re.split(r";|\n", values):
            if "=" in part:
                k, _, v = part.partition("=")
                pairs[k.strip()] = v.strip()
        filled = re.sub(r"\{(\w+)\}", lambda m: pairs.get(m.group(1), f"[{m.group(1)}]"), body)
        filled = re.sub(r" \[detail\]| \[cover\]| \[fix\]", "", filled)
        left = [f for f in template_fields(body) if f not in pairs and f not in {"detail", "cover", "fix"}]
        return filled + (f"\n\n(Still to fill: {', '.join(left)})" if left else "")

    @command("handout", group=G, usage="/handout <slides.pptx|.jdesign> [notes|lines]",
             help="a Word handout of a slide deck: pictures, notes or lines to write on", title="Slides to Word handout",
             icon="🗂", page="documents",
             fields=(field("file", "file", "Slides", types=(("Slides", "*.pptx *.jdesign"),)),
                     field("style", "choice", "Beside each slide", "notes", ("notes", "lines"))),
             template="{file} | {style}")
    def handout_cmd(self, args: str, routed: bool = False):
        from docx import Document
        from docx.shared import Inches, Pt

        from ..design import model, pptxio, render

        path_text, style = split(args, 2)
        if not style and path_text.rsplit(" ", 1)[-1].lower() in {"notes", "lines"}:
            path_text, style = path_text.rsplit(" ", 1)
        try:
            path = need_file(path_text, ".pptx", ".jdesign")
        except OfficeError as exc:
            return str(exc)
        design = model.load(path) if path.suffix.lower() == ".jdesign" else pptxio.import_pptx(path)
        document = Document()
        document.add_heading(design.get("title") or path.stem, 0)
        temp = kit.output_dir("documents") / f".handout_{kit.stamp()}"
        temp.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(design["pages"]):
            image = render.render_page(design, i, scale=900 / max(1, design["w"]))
            png = temp / f"slide_{i + 1}.png"
            image.convert("RGB").save(png)
            table = document.add_table(rows=1, cols=2)
            left, right = table.rows[0].cells
            left.paragraphs[0].add_run().add_picture(str(png), width=Inches(3.4))
            notes = (page.get("notes") or "").strip()
            if (style or "notes") == "lines" or not notes:
                for _ in range(8):
                    right.add_paragraph("_" * 38)
            else:
                run = right.paragraphs[0].add_run(f"Slide {i + 1}\n{notes}")
                run.font.size = Pt(10)
            document.add_paragraph()
        target = out_path(path, "handout", ".docx")
        document.save(target)
        shutil.rmtree(temp, ignore_errors=True)
        return f"🗂 Handout with {len(design['pages'])} slide(s):\n  {target}"

    @command("print", group=G, usage="/print <file>", help="send a document to your default printer (asks first)",
             title="Print a file", icon="🖨", page="documents", fields=(field("file", "file", "File to print"),))
    def print_cmd(self, args: str, routed: bool = False):
        try:
            path = need_file(args)
        except OfficeError as exc:
            return str(exc)
        if os.name != "nt":
            return "Printing from JARVIS needs Windows."
        if not security.permissions.ask(security.RUN_COMMAND, f"print {path}", context="/print"):
            return "Not printed."
        try:
            os.startfile(str(path), "print")  # noqa: S606 — the user's own file, after asking
        except OSError as exc:
            return f"Windows couldn't print that: {exc}"
        security.audit.record("print", path.name)
        return f"🖨 Sent {path.name} to your default printer."

    @command("doctemplate", "template", group=G, usage="/doctemplate <letter|report|agenda|…> | <details>",
             help="a document from a template: letter, report, agenda, minutes, proposal, memo, dilekçe…",
             title="Document templates", icon="🗃", page="documents",
             fields=(field("kind", "choice", "Template", "letter", tuple(DOC_TEMPLATES)),
                     field("details", "long", "Details to put in it")))
    def doctemplate_cmd(self, args: str, routed: bool = False):
        kind, details = split(args, 2)
        kind = kind.lower().strip()
        if kind not in DOC_TEMPLATES:
            return "🗃 Templates: " + ", ".join(f"{k} ({v[0]})" for k, v in DOC_TEMPLATES.items()) + \
                "\nUse: /doctemplate letter | to the landlord about the broken heating"
        if not details:
            return f"What should the {DOC_TEMPLATES[kind][0].lower()} say? /doctemplate {kind} | <details>"
        return self.make_document(f"{DOC_TEMPLATES[kind][1]}. Use these details: {details}", ("docx", "pdf"))

    @command("recentdocs", group=G, usage="/recentdocs [search]", help="documents you opened or made recently",
             title="Recent documents", icon="🕘", page="documents",
             fields=(field("search", "text", "Filter", optional=True),))
    def recentdocs_cmd(self, args: str, routed: bool = False):
        words = args.lower().split()
        items = [(t, p) for t, p in recent_documents(80) if all(w in p.name.lower() for w in words)]
        if not items:
            return "No recent documents found."
        lines = ["🕘 Recent documents:"]
        for t, p in items[:30]:
            real = resolve_link(p)
            lines.append(f"  {time.strftime('%d %b %H:%M', time.localtime(t))}  {real}")
        return "\n".join(lines)


def _minutes(value: float) -> str:
    if value < 1:
        return f"{max(1, round(value * 60))} s"
    whole = int(value)
    seconds = round((value - whole) * 60)
    return f"{whole} min" + (f" {seconds} s" if seconds else "")
