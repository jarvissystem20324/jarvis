"""8.0 — Your PC and files: duplicates, batch rename, image and file
conversion, PDF merge and split, startup apps, Wi-Fi, dark mode, wallpaper,
specs, a text expander, screen time, cleanup and do-not-disturb.

Anything that changes the machine asks first, and is undoable where Windows
allows: extra copies go to the Recycle Bin, renames are logged for
/rename undo, startup apps are switched the same way Task Manager does it
(StartupApproved), so either can switch them back.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from . import kit, reminders, security
from .registry import command

G = "PC and files"
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
RENAMES = kit.Store("renames.json", [])
EXPANSIONS = kit.Store("expansions.json", {})
SCREEN_TIME = kit.Store("screentime.json", {})
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)


def _windows_only() -> str | None:
    return None if sys.platform == "win32" else "That works on Windows only."


def powershell(script: str, timeout: float = 30) -> str:
    result = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                            capture_output=True, text=True, timeout=timeout, creationflags=NO_WINDOW,
                            encoding="utf-8", errors="replace")
    return result.stdout.strip()


# --- duplicates --------------------------------------------------------------------------

def find_duplicates(folder: Path, limit: int = 60000) -> list[list[Path]]:
    by_size: dict[int, list[Path]] = defaultdict(list)
    seen = 0
    for root, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if not d.startswith(".") and d not in {"node_modules", "__pycache__", "$RECYCLE.BIN"}]
        for name in files:
            path = Path(root) / name
            try:
                size = path.stat().st_size
            except OSError:
                continue
            if size > 0:
                by_size[size].append(path)
            seen += 1
            if seen >= limit:
                break
        if seen >= limit:
            break

    def digest(path: Path, head: bool) -> str:
        h = hashlib.sha256()
        with path.open("rb") as handle:
            if head:
                h.update(handle.read(65536))
            else:
                for block in iter(lambda: handle.read(1 << 20), b""):
                    h.update(block)
        return h.hexdigest()

    groups: list[list[Path]] = []
    for size, paths in by_size.items():
        if len(paths) < 2:
            continue
        for stage in (True, False):
            buckets: dict[str, list[Path]] = defaultdict(list)
            for path in paths:
                try:
                    buckets[digest(path, stage)].append(path)
                except OSError:
                    continue
            paths = [p for bucket in buckets.values() if len(bucket) > 1 for p in bucket]
            if stage:
                continue
            groups += [sorted(b, key=lambda p: (p.stat().st_mtime, len(str(p)))) for b in buckets.values() if len(b) > 1]
    return sorted(groups, key=lambda g: -g[0].stat().st_size * (len(g) - 1))


def recycle(paths: list[Path]) -> int:
    """Send files to the Recycle Bin (undoable), not delete them."""
    if sys.platform != "win32":
        raise OSError("The Recycle Bin is Windows-only.")
    import ctypes
    from ctypes import wintypes

    class SHFILEOPSTRUCTW(ctypes.Structure):
        _fields_ = [("hwnd", wintypes.HWND), ("wFunc", wintypes.UINT), ("pFrom", wintypes.LPCWSTR),
                    ("pTo", wintypes.LPCWSTR), ("fFlags", ctypes.c_uint16), ("fAnyOperationsAborted", wintypes.BOOL),
                    ("hNameMappings", ctypes.c_void_p), ("lpszProgressTitle", wintypes.LPCWSTR)]

    FO_DELETE, FOF_ALLOWUNDO, FOF_NOCONFIRMATION, FOF_SILENT, FOF_NOERRORUI = 3, 0x40, 0x10, 0x4, 0x400
    done = 0
    for path in paths:
        op = SHFILEOPSTRUCTW(None, FO_DELETE, str(path) + "\0", None,
                             FOF_ALLOWUNDO | FOF_NOCONFIRMATION | FOF_SILENT | FOF_NOERRORUI, False, None, None)
        if ctypes.windll.shell32.SHFileOperationW(ctypes.byref(op)) == 0:
            done += 1
    return done


# --- renaming ---------------------------------------------------------------------------------

def rename_plan(folder: Path, rule: str) -> list[tuple[Path, Path]]:
    files = sorted([p for p in folder.iterdir() if p.is_file()], key=lambda p: p.name.lower())
    rule = rule.strip()
    low = rule.lower()
    plan: list[tuple[Path, Path]] = []
    for n, path in enumerate(files, 1):
        stem, suffix = path.stem, path.suffix
        if low.startswith("prefix "):
            new = rule[7:] + stem
        elif low.startswith("suffix "):
            new = stem + rule[7:]
        elif m := re.match(r"replace\s+(.+?)\s+with\s+(.*)$", rule, re.I):
            new = stem.replace(m.group(1), m.group(2))
        elif low.startswith(("number", "numbered")):
            base = rule.split(None, 1)[1] if " " in rule else folder.name
            new = f"{base}_{n:0{max(3, len(str(len(files))))}d}"
        elif low == "date":
            new = f"{datetime.fromtimestamp(path.stat().st_mtime):%Y-%m-%d}_{stem}"
        elif low in {"lower", "lowercase"}:
            new = stem.lower()
        elif low in {"upper", "uppercase"}:
            new = stem.upper()
        elif low in {"clean", "tidy"}:
            new = re.sub(r"[_\s]+", " ", re.sub(r"\(\d+\)|copy|kopya", "", stem, flags=re.I)).strip()
        else:
            raise ValueError("Rules: prefix X · suffix X · replace A with B · number [name] · date · lower · upper · clean")
        target = path.with_name(new + suffix)
        if target != path:
            plan.append((path, target))
    names = [t.name.lower() for _, t in plan]
    existing = {p.name.lower() for p in files} - {s.name.lower() for s, _ in plan}
    if len(set(names)) != len(names) or existing & set(names):
        raise ValueError("That rule would give two files the same name. Try number or date.")
    return plan


# --- conversion -------------------------------------------------------------------------------

def convert_file(source: Path, target_format: str, out_dir: Path) -> Path:
    fmt = target_format.lower().lstrip(".")
    fmt = {"jpeg": "jpg", "word": "docx", "excel": "xlsx", "text": "txt", "markdown": "md"}.get(fmt, fmt)
    suffix = source.suffix.lower()
    out = out_dir / f"{source.stem}.{fmt}"
    if suffix in IMAGE_SUFFIXES and fmt in {"jpg", "png", "webp", "bmp", "gif", "tiff", "ico"}:
        from PIL import Image

        with Image.open(source) as image:
            image = image.convert("RGB") if fmt in {"jpg", "bmp"} else image
            image.save(out, quality=90) if fmt in {"jpg", "webp"} else image.save(out)
        return out
    if suffix in IMAGE_SUFFIXES and fmt == "pdf":
        from PIL import Image

        with Image.open(source) as image:
            image.convert("RGB").save(out, "PDF", resolution=150)
        return out
    if suffix == ".csv" and fmt == "xlsx":
        import csv

        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        text = source.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        for row in csv.reader(text.splitlines(), dialect):
            sheet.append([_cell(v) for v in row])
        book.save(out)
        return out
    if suffix in {".xlsx", ".xlsm"} and fmt == "csv":
        import csv

        from openpyxl import load_workbook

        sheet = load_workbook(source, read_only=True, data_only=True).active
        with out.open("w", newline="", encoding="utf-8-sig") as handle:
            writer_ = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer_.writerow(["" if v is None else v for v in row])
        return out
    if suffix == ".pdf" and fmt in {"txt", "md"}:
        from pypdf import PdfReader

        out.write_text("\n\n".join((p.extract_text() or "") for p in PdfReader(str(source)).pages), encoding="utf-8")
        return out
    if suffix == ".docx" and fmt == "pdf":
        if sys.platform == "win32":
            try:
                import win32com.client

                word = win32com.client.DispatchEx("Word.Application")
                word.Visible = False
                try:
                    document = word.Documents.Open(str(source.resolve()), ReadOnly=True)
                    document.SaveAs2(str(out.resolve()), FileFormat=17)
                    document.Close(False)
                finally:
                    word.Quit()
                return out
            except Exception:
                pass
        fmt_text = _docx_markdown(source)
        from . import writer

        return writer.to_pdf(fmt_text, out)
    if suffix == ".docx" and fmt in {"txt", "md"}:
        out.write_text(_docx_markdown(source), encoding="utf-8")
        return out
    if suffix in {".md", ".txt"} and fmt in {"docx", "pdf"}:
        from . import writer

        text = source.read_text(encoding="utf-8", errors="replace")
        if not text.lstrip().startswith("#"):
            text = f"# {source.stem}\n\n{text}"
        return writer.to_docx(text, out) if fmt == "docx" else writer.to_pdf(text, out)
    raise ValueError(f"I can't convert {suffix or 'that'} to {fmt}. I do images ↔ jpg/png/webp/pdf, csv ↔ xlsx, "
                     "docx → pdf/txt, pdf → txt, md/txt → docx/pdf.")


def _cell(value: str):
    try:
        return int(value) if re.fullmatch(r"-?\d+", value) else float(value) if re.fullmatch(r"-?\d+\.\d+", value) else value
    except ValueError:
        return value


def _docx_markdown(source: Path) -> str:
    from docx import Document

    lines = []
    for paragraph in Document(str(source)).paragraphs:
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") or style == "title":
            level = int(re.sub(r"\D", "", style) or 1)
            lines.append("#" * min(level, 3) + " " + paragraph.text)
        elif paragraph.text.strip():
            lines.append(("- " if "list" in style else "") + paragraph.text)
        lines.append("")
    return "\n".join(lines)


# --- Windows settings ----------------------------------------------------------------------------

PERSONALIZE = r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize"
RUN = r"Software\Microsoft\Windows\CurrentVersion\Run"
APPROVED = r"Software\Microsoft\Windows\CurrentVersion\Explorer\StartupApproved"


def _broadcast_setting_change(area: str = "ImmersiveColorSet") -> None:
    import ctypes

    HWND_BROADCAST, WM_SETTINGCHANGE, SMTO_ABORTIFHUNG = 0xFFFF, 0x1A, 0x2
    result = ctypes.c_ulong()
    ctypes.windll.user32.SendMessageTimeoutW(HWND_BROADCAST, WM_SETTINGCHANGE, 0, area, SMTO_ABORTIFHUNG, 1000,
                                             ctypes.byref(result))


def dark_mode(on: bool | None = None) -> bool:
    """Read (on=None) or set Windows' dark mode for apps and the system. Returns the new state."""
    import winreg

    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, PERSONALIZE, 0, winreg.KEY_READ | winreg.KEY_SET_VALUE) as key:
        current = winreg.QueryValueEx(key, "AppsUseLightTheme")[0] == 0
        if on is None:
            return current
        for name in ("AppsUseLightTheme", "SystemUsesLightTheme"):
            winreg.SetValueEx(key, name, 0, winreg.REG_DWORD, 0 if on else 1)
    _broadcast_setting_change()
    return on


def startup_entries() -> list[dict]:
    import winreg

    entries: list[dict] = []
    for hive, scope in ((winreg.HKEY_CURRENT_USER, "you"), (winreg.HKEY_LOCAL_MACHINE, "all users")):
        try:
            with winreg.OpenKey(hive, RUN) as key:
                for i in range(winreg.QueryInfoKey(key)[1]):
                    name, command_line, _ = winreg.EnumValue(key, i)
                    entries.append({"name": name, "command": str(command_line), "scope": scope, "where": "Run", "hive": hive})
        except OSError:
            continue
    folders = [(Path(os.environ.get("APPDATA", "")) / r"Microsoft\Windows\Start Menu\Programs\Startup", "you", winreg.HKEY_CURRENT_USER),
               (Path(os.environ.get("PROGRAMDATA", "")) / r"Microsoft\Windows\Start Menu\Programs\StartUp", "all users", winreg.HKEY_LOCAL_MACHINE)]
    for folder, scope, hive in folders:
        if folder.is_dir():
            for item in folder.iterdir():
                if item.name.lower() != "desktop.ini":
                    entries.append({"name": item.name, "command": str(item), "scope": scope, "where": "StartupFolder", "hive": hive})
    for entry in entries:
        entry["enabled"] = _approved(entry)
    return entries


def _approved(entry: dict) -> bool:
    import winreg

    try:
        with winreg.OpenKey(entry["hive"], f"{APPROVED}\\{entry['where']}") as key:
            data = winreg.QueryValueEx(key, entry["name"])[0]
            return not data or data[0] % 2 == 0     # 02/06 enabled, 03/07 disabled — as Task Manager writes it
    except OSError:
        return True


def set_startup(entry: dict, enabled: bool) -> None:
    import struct
    import winreg

    if entry["hive"] != winreg.HKEY_CURRENT_USER:
        raise PermissionError("That one starts for all users; switch it in Task Manager → Startup apps (needs admin).")
    value = (b"\x02" if enabled else b"\x03") + b"\x00" * 3 + (b"\x00" * 8 if enabled else struct.pack("<Q", int((time.time() + 11644473600) * 1e7)))
    with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, f"{APPROVED}\\{entry['where']}", 0, winreg.KEY_SET_VALUE) as key:
        winreg.SetValueEx(key, entry["name"], 0, winreg.REG_BINARY, value)


def wifi_info() -> dict:
    out = subprocess.run(["netsh", "wlan", "show", "interfaces"], capture_output=True, text=True,
                         creationflags=NO_WINDOW, encoding="utf-8", errors="replace").stdout
    fields: dict[str, str] = {}
    for line in out.splitlines():
        if ":" in line:
            key, _, value = line.partition(":")
            fields.setdefault(key.strip().lower(), value.strip())
    return fields


def set_wallpaper(path: Path) -> None:
    import ctypes

    SPI_SETDESKWALLPAPER, SPIF_UPDATEINIFILE, SPIF_SENDCHANGE = 20, 0x1, 0x2
    if not ctypes.windll.user32.SystemParametersInfoW(SPI_SETDESKWALLPAPER, 0, str(path), SPIF_UPDATEINIFILE | SPIF_SENDCHANGE):
        raise OSError("Windows refused the wallpaper.")


def notifications(on: bool) -> None:
    import winreg

    with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\PushNotifications",
                            0, winreg.KEY_SET_VALUE) as key:
        winreg.SetValueEx(key, "ToastEnabled", 0, winreg.REG_DWORD, 1 if on else 0)
    _broadcast_setting_change("Policy")


def recycle_bin_size() -> tuple[int, int]:
    import ctypes

    class SHQUERYRBINFO(ctypes.Structure):
        _fields_ = [("cbSize", ctypes.c_ulong), ("i64Size", ctypes.c_int64), ("i64NumItems", ctypes.c_int64)]

    info = SHQUERYRBINFO(ctypes.sizeof(SHQUERYRBINFO), 0, 0)
    ctypes.windll.shell32.SHQueryRecycleBinW(None, ctypes.byref(info))
    return int(info.i64Size), int(info.i64NumItems)


def temp_files(older_than_hours: float = 24) -> tuple[list[Path], int]:
    cutoff = time.time() - older_than_hours * 3600
    found, total = [], 0
    for root, dirs, files in os.walk(tempfile.gettempdir()):
        for name in files:
            path = Path(root) / name
            try:
                stat = path.stat()
            except OSError:
                continue
            if stat.st_mtime < cutoff:
                found.append(path)
                total += stat.st_size
    return found, total


# --- screen time --------------------------------------------------------------------------------

def foreground_app() -> str:
    import ctypes
    from ctypes import wintypes

    hwnd = ctypes.windll.user32.GetForegroundWindow()
    if not hwnd:
        return ""
    pid = wintypes.DWORD()
    ctypes.windll.user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
    try:
        import psutil

        name = psutil.Process(pid.value).name()
    except Exception:
        return ""
    return re.sub(r"\.exe$", "", name, flags=re.I)


def idle_seconds() -> float:
    import ctypes

    class LASTINPUTINFO(ctypes.Structure):
        _fields_ = [("cbSize", ctypes.c_uint), ("dwTime", ctypes.c_uint)]

    info = LASTINPUTINFO(ctypes.sizeof(LASTINPUTINFO), 0)
    ctypes.windll.user32.GetLastInputInfo(ctypes.byref(info))
    return (ctypes.windll.kernel32.GetTickCount() - info.dwTime) / 1000


class ScreenTime:
    """Counts which app is in front, every few seconds, while JARVIS runs."""

    def __init__(self, every: float = 5.0):
        self.every = every
        self.pending: dict[str, float] = defaultdict(float)
        self._stop = None

    def start(self) -> None:
        import threading

        if sys.platform != "win32" or self._stop is not None:
            return
        self._stop = threading.Event()
        threading.Thread(target=self._loop, daemon=True).start()

    def stop(self) -> None:
        if self._stop:
            self._stop.set()
            self.flush()

    def sample(self, app: str, idle: float) -> None:
        if app and idle < 120:      # away from the keyboard isn't screen time
            self.pending[app] += self.every

    def flush(self) -> None:
        if not self.pending:
            return
        data = SCREEN_TIME.load()
        day = data.setdefault(date.today().isoformat(), {})
        for app, seconds in self.pending.items():
            day[app] = day.get(app, 0) + seconds
        for old in sorted(data)[:-60]:
            data.pop(old, None)
        SCREEN_TIME.save(data)
        self.pending.clear()

    def _loop(self) -> None:
        ticks = 0
        while not self._stop.wait(self.every):
            try:
                self.sample(foreground_app(), idle_seconds())
            except Exception:
                pass
            ticks += 1
            if ticks % 12 == 0:
                try:
                    self.flush()
                except Exception:
                    pass


screen_time = ScreenTime()


class PCTools:
    @command("dupes", "duplicates", group=G, usage="/dupes [folder] · /dupes clean",
             help="finds duplicate files; extra copies go to the Recycle Bin")
    def dupes(self, args: str, routed: bool = False):
        from . import pc

        text = args.strip()
        clean = text.lower().startswith("clean")
        if clean:
            groups = getattr(self, "_dupes", None)
            if not groups:
                return "Run /dupes <folder> first."
            extra = [p for g in groups for p in g[1:] if p.exists()]
            wasted = sum(p.stat().st_size for p in extra)
            if not security.permissions.ask(security.DELETE_FILE, f"send {len(extra)} duplicate file(s) ({kit.size(wasted)}) "
                                            "to the Recycle Bin, keeping the oldest copy of each", context="/dupes"):
                return "Denied. Nothing was moved."
            try:
                moved = recycle(extra)
            except OSError as exc:
                return str(exc)
            self._dupes = None
            security.audit.record("dupes clean", f"{moved} files", kit.size(wasted))
            return f"🗑 Sent {moved} duplicate(s) to the Recycle Bin, freeing {kit.size(wasted)}. Restore them from the bin if needed."
        folder = kit.path_arg(text) if text else pc.known_folder("downloads")
        if folder is None or not folder.is_dir():
            return f"No such folder: {text or 'Downloads'}"
        groups = find_duplicates(folder)
        self._dupes = groups
        if not groups:
            return f"No duplicate files in {folder}. 👌"
        wasted = sum(g[0].stat().st_size * (len(g) - 1) for g in groups)
        lines = [f"Found {len(groups)} set(s) of duplicates in {folder}, wasting {kit.size(wasted)}:"]
        for g in groups[:12]:
            lines.append(f"  {kit.size(g[0].stat().st_size)} × {len(g)}: {g[0].name}")
            lines += [f"      also {p.relative_to(folder) if p.is_relative_to(folder) else p}" for p in g[1:3]]
        lines.append("/dupes clean keeps the oldest copy of each and sends the rest to the Recycle Bin.")
        return "\n".join(lines)

    @command("rename", "batchrename", group=G, usage="/rename <folder> number holiday · /rename go · /rename undo",
             help="renames every file in a folder by a rule, with undo")
    def rename(self, args: str, routed: bool = False):
        text = args.strip()
        if text.lower() == "undo":
            log = RENAMES.load()
            if not log:
                return "Nothing to undo."
            last = log.pop()
            undone = 0
            for new, old in reversed(last["moves"]):
                if Path(new).exists() and not Path(old).exists():
                    Path(new).rename(old)
                    undone += 1
            RENAMES.save(log)
            return f"Undid {undone} rename(s)."
        if text.lower() == "go":
            plan = getattr(self, "_rename_plan", None)
            if not plan:
                return "Make a plan first: /rename <folder> <rule>"
            if not security.permissions.ask(security.WRITE_FILE, f"rename {len(plan)} file(s) in {plan[0][0].parent}", context="/rename"):
                return "Denied. Nothing was renamed."
            done = []
            for old, new in plan:
                try:
                    old.rename(new)
                    done.append((str(new), str(old)))
                except OSError:
                    continue
            log = RENAMES.load()
            log.append({"at": time.time(), "moves": done})
            RENAMES.save(log[-20:])
            self._rename_plan = None
            return f"Renamed {len(done)} file(s). /rename undo puts the names back."
        m = re.match(r'^("[^"]+"|\S+)\s+(.+)$', text)
        folder = kit.path_arg(m.group(1)) if m else None
        if not m or folder is None or not folder.is_dir():
            return ("Usage: /rename <folder> <rule>, then /rename go\n"
                    "Rules: prefix X · suffix X · replace A with B · number [name] · date · lower · upper · clean")
        try:
            plan = rename_plan(folder, m.group(2))
        except ValueError as exc:
            return str(exc)
        if not plan:
            return "That rule changes nothing."
        self._rename_plan = plan
        preview = "\n".join(f"  {a.name}  →  {b.name}" for a, b in plan[:15])
        more = f"\n  … and {len(plan) - 15} more" if len(plan) > 15 else ""
        return f"Plan ({len(plan)} files):\n{preview}{more}\n/rename go does it · /rename undo afterwards puts them back."

    @command("imgs", "bulkimages", group=G, usage="/imgs <folder> to jpg [1600px] [quality 80]",
             help="converts or shrinks every image in a folder")
    def bulk_images(self, args: str, routed: bool = False):
        m = re.match(r'^("[^"]+"|\S+)\s+(?:to\s+)?(jpg|jpeg|png|webp)\b(.*)$', args.strip(), re.I)
        folder = kit.path_arg(m.group(1)) if m else None
        if not m or folder is None or not folder.is_dir():
            return "Usage: /imgs <folder> to jpg|png|webp [1600px] [quality 80]   — results go in a 'converted' subfolder."
        fmt = m.group(2).lower().replace("jpeg", "jpg")
        size = re.search(r"(\d{2,5})\s*px", m.group(3))
        quality = re.search(r"quality\s*(\d{1,3})", m.group(3))
        images = [p for p in folder.iterdir() if p.suffix.lower() in IMAGE_SUFFIXES]
        if not images:
            return f"No images in {folder}."
        if not security.permissions.ask(security.WRITE_FILE, f"write {len(images)} converted image(s) to {folder / 'converted'}", context="/imgs"):
            return "Denied."
        from PIL import Image

        out_dir = folder / "converted"
        out_dir.mkdir(exist_ok=True)
        before = after = 0
        for path in images:
            try:
                with Image.open(path) as image:
                    if size:
                        image.thumbnail((int(size.group(1)), int(size.group(1))))
                    if fmt == "jpg":
                        image = image.convert("RGB")
                    target = out_dir / f"{path.stem}.{fmt}"
                    image.save(target, quality=int(quality.group(1)) if quality else 85, optimize=True)
                before += path.stat().st_size
                after += target.stat().st_size
            except Exception:
                continue
        return f"🖼 Converted {len(images)} image(s) to {fmt.upper()} in {out_dir}\n  {kit.size(before)} → {kit.size(after)}"

    @command("pdfmerge", group=G, usage="/pdfmerge a.pdf b.pdf … (or a folder)", help="joins PDFs into one")
    def pdf_merge(self, args: str, routed: bool = False):
        from pypdf import PdfWriter

        parts = re.findall(r'"([^"]+)"|(\S+)', args)
        paths = [Path(a or b).expanduser() for a, b in parts]
        if len(paths) == 1 and paths[0].is_dir():
            paths = sorted(paths[0].glob("*.pdf"), key=lambda p: p.name.lower())
        paths = [p for p in paths if p.suffix.lower() == ".pdf" and p.exists()]
        if len(paths) < 2:
            return "Usage: /pdfmerge a.pdf b.pdf c.pdf   or   /pdfmerge <folder of PDFs>"
        writer_ = PdfWriter()
        for path in paths:
            writer_.append(str(path))
        out = kit.output_dir() / f"merged_{kit.stamp()}.pdf"
        with out.open("wb") as handle:
            writer_.write(handle)
        return f"📎 Merged {len(paths)} PDFs ({sum(1 for _ in writer_.pages)} pages):\n  {out}"

    @command("pdfsplit", group=G, usage="/pdfsplit file.pdf [1-3,7]", help="extracts pages, or splits every page")
    def pdf_split(self, args: str, routed: bool = False):
        from pypdf import PdfReader, PdfWriter

        m = re.match(r'^("[^"]+"|.+?\.pdf)\s*([\d,\s-]*)$', args.strip(), re.I)
        path = kit.path_arg(m.group(1)) if m else None
        if path is None:
            return "Usage: /pdfsplit <file.pdf> [pages like 1-3,7]   — no pages = one file per page"
        reader = PdfReader(str(path))
        total = len(reader.pages)
        ranges = m.group(2).strip()
        out_dir = kit.output_dir()
        if ranges:
            pages: list[int] = []
            for part in re.split(r"\s*,\s*", ranges):
                if "-" in part:
                    a, b = (int(x) for x in part.split("-", 1))
                    pages += list(range(a, b + 1))
                elif part:
                    pages.append(int(part))
            pages = [p for p in pages if 1 <= p <= total]
            if not pages:
                return f"{path.name} has {total} pages."
            writer_ = PdfWriter()
            for p in pages:
                writer_.add_page(reader.pages[p - 1])
            out = out_dir / f"{path.stem}_pages_{ranges.replace(',', '_').replace(' ', '')}.pdf"
            with out.open("wb") as handle:
                writer_.write(handle)
            return f"✂ Pages {ranges} of {path.name}:\n  {out}"
        folder = out_dir / f"{path.stem}_pages"
        folder.mkdir(exist_ok=True)
        for i, page in enumerate(reader.pages, 1):
            writer_ = PdfWriter()
            writer_.add_page(page)
            with (folder / f"{path.stem}_{i:03d}.pdf").open("wb") as handle:
                writer_.write(handle)
        return f"✂ Split {path.name} into {total} files in\n  {folder}"

    @command("fileconvert", "convertfile", group=G, usage="/fileconvert <file> to <pdf|docx|xlsx|csv|jpg|png|txt>",
             help="converts documents, sheets and images")
    def file_convert(self, args: str, routed: bool = False):
        m = re.match(r'^("[^"]+"|.+?)\s+(?:to|->|into)\s+\.?(\w+)$', args.strip(), re.I)
        path = kit.path_arg(m.group(1)) if m else None
        if path is None:
            return "Usage: /fileconvert <file> to <format>   e.g. /fileconvert report.docx to pdf · /fileconvert data.csv to xlsx"
        try:
            out = convert_file(path, m.group(2), kit.output_dir())
        except ValueError as exc:
            return str(exc)
        except Exception as exc:
            return f"Conversion failed: {exc}"
        return f"🔁 {path.name} → {out.name}\n  {out}"

    # --- Windows ----------------------------------------------------------------------------

    @command("startup", "startupapps", group=G, usage="/startup · /startup disable <name> · /startup enable <name>",
             help="what starts with Windows, and switching it off")
    def startup(self, args: str, routed: bool = False):
        if (problem := _windows_only()):
            return problem
        text = args.strip()
        verb, _, name = text.partition(" ")
        entries = startup_entries()
        if verb.lower() in {"disable", "off", "enable", "on"}:
            match = [e for e in entries if name.strip().lower() in e["name"].lower()] if name.strip() else []
            if len(match) != 1:
                return f"Which one? {len(match)} entries match '{name}'. /startup lists them." if match else f"Nothing called '{name}'."
            entry, enable = match[0], verb.lower() in {"enable", "on"}
            if not security.permissions.ask(security.CHANGE_SETTING, f"{'enable' if enable else 'disable'} {entry['name']} at Windows start-up",
                                            context="/startup"):
                return "Denied. Nothing changed."
            try:
                set_startup(entry, enable)
            except PermissionError as exc:
                return str(exc)
            security.audit.record("startup", entry["name"], "enabled" if enable else "disabled")
            return f"{entry['name']} will {'now' if enable else 'no longer'} start with Windows (Task Manager shows the same)."
        if not entries:
            return "Nothing is set to start with Windows (from the usual places)."
        lines = ["Starts with Windows:"]
        for e in sorted(entries, key=lambda e: (not e["enabled"], e["name"].lower())):
            lines.append(f"  {'●' if e['enabled'] else '○'} {e['name']}  ({e['scope']}{'' if e['enabled'] else ', off'})")
        lines.append("/startup disable <name> · /startup enable <name>   (fewer = faster boot)")
        return "\n".join(lines)

    @command("wifi", group=G, usage="/wifi", help="your Wi-Fi network, signal and speed")
    def wifi(self, args: str, routed: bool = False):
        if (problem := _windows_only()):
            return problem
        if args.strip().lower() in {"check", "safe", "safety", "secure"}:
            return self.wifi_safety("")
        info = wifi_info()
        if not info.get("ssid"):
            return "Not connected to Wi-Fi (or no Wi-Fi adapter)."
        return (f"📶 {info.get('ssid')}  ·  signal {info.get('signal', '?')}  ·  {info.get('radio type', '')}  ·  channel {info.get('channel', '?')}\n"
                f"  link {info.get('receive rate (mbps)', '?')} Mbps down / {info.get('transmit rate (mbps)', '?')} Mbps up  ·  "
                f"security {info.get('authentication', '?')}\n  /speedtest for real speed · /wifi check for a safety check")

    @command("darkmode", "theme-windows", group=G, usage="/darkmode on|off", help="switches Windows dark mode")
    def darkmode(self, args: str, routed: bool = False):
        if (problem := _windows_only()):
            return problem
        word = args.strip().lower()
        try:
            current = dark_mode()
            wanted = {"on": True, "dark": True, "off": False, "light": False}.get(word, not current)
            if wanted == current:
                return f"Windows is already in {'dark' if current else 'light'} mode."
            if not security.permissions.ask(security.CHANGE_SETTING, f"switch Windows to {'dark' if wanted else 'light'} mode", context="/darkmode"):
                return "Denied."
            dark_mode(wanted)
        except OSError as exc:
            return f"Couldn't change it: {exc}"
        return f"{'🌙 Dark' if wanted else '☀ Light'} mode on."

    @command("wallpaper", group=G, usage="/wallpaper <what to paint> · /wallpaper last",
             help="paints an AI wallpaper and sets it")
    def wallpaper(self, args: str, routed: bool = False):
        from .assistant import JarvisResponse

        if (problem := _windows_only()):
            return problem
        text = args.strip()
        if text.lower() in {"last", "this", "current"}:
            path = getattr(self, "current_image", None)
            if not path:
                return "No image yet. /wallpaper <description> paints one."
        elif text:
            try:
                path = self.images.generate(f"{text}, desktop wallpaper, 16:9, highly detailed, no text", size="1792x1024")
            except Exception as exc:
                return f"Couldn't paint it: {exc}"
        else:
            return "Usage: /wallpaper <description>   e.g. /wallpaper misty Istanbul skyline at dawn, cinematic"
        if not security.permissions.ask(security.CHANGE_SETTING, f"set {Path(path).name} as your desktop wallpaper", context="/wallpaper"):
            return "Denied. The image is saved but not set."
        try:
            set_wallpaper(Path(path).resolve())
        except OSError as exc:
            return str(exc)
        self.current_image = Path(path)
        return JarvisResponse(text=f"🖼 New wallpaper set.\n{path}", image_path=Path(path), image_paths=[Path(path)])

    @command("specs", "sysinfo", group=G, usage="/specs", help="your PC's hardware and Windows version")
    def specs(self, args: str, routed: bool = False):
        import platform

        import psutil

        lines = ["🖥 This PC"]
        details = {}
        if sys.platform == "win32":
            try:
                raw = powershell("$o=@{}; $c=Get-CimInstance Win32_Processor|Select -First 1; $o.cpu=$c.Name; "
                                 "$o.gpu=@(Get-CimInstance Win32_VideoController|%{ $_.Name + ' (driver ' + $_.DriverVersion + ')'}); "
                                 "$b=Get-CimInstance Win32_BaseBoard; $o.board=$b.Manufacturer + ' ' + $b.Product; "
                                 "$s=Get-CimInstance Win32_ComputerSystem; $o.model=$s.Manufacturer + ' ' + $s.Model; "
                                 "$v=Get-ItemProperty 'HKLM:\\SOFTWARE\\Microsoft\\Windows NT\\CurrentVersion'; "
                                 "$o.os=$v.ProductName + ' ' + $v.DisplayVersion + ' (build ' + $v.CurrentBuild + ')'; "
                                 "$o.disks=@(Get-PhysicalDisk|%{ $_.FriendlyName + ' ' + $_.MediaType + ' ' + [math]::Round($_.Size/1GB) + ' GB'}); "
                                 "$o|ConvertTo-Json -Compress", timeout=40)
                details = json.loads(raw) if raw else {}
            except Exception:
                details = {}
        mem = psutil.virtual_memory()
        boot = datetime.fromtimestamp(psutil.boot_time())
        lines += [f"  Model     {details.get('model', platform.node())}",
                  f"  CPU       {details.get('cpu') or platform.processor()} — {psutil.cpu_count(logical=False)} cores / {psutil.cpu_count()} threads",
                  f"  Memory    {kit.size(mem.total)} ({mem.percent:.0f}% in use)"]
        for gpu in details.get("gpu") or []:
            lines.append(f"  GPU       {gpu}")
        for disk in details.get("disks") or []:
            lines.append(f"  Disk      {disk}")
        for part in psutil.disk_partitions():
            try:
                usage = psutil.disk_usage(part.mountpoint)
                lines.append(f"  {part.mountpoint:<9} {kit.size(usage.free)} free of {kit.size(usage.total)}")
            except OSError:
                continue
        lines += [f"  Board     {details.get('board', '')}".rstrip(),
                  f"  Windows   {details.get('os') or platform.platform()}",
                  f"  Up since  {boot:%d %b %H:%M} ({kit.ago(boot.timestamp()).replace(' ago', '')})"]
        battery = psutil.sensors_battery() if hasattr(psutil, "sensors_battery") else None
        if battery:
            lines.append(f"  Battery   {battery.percent:.0f}%{' (charging)' if battery.power_plugged else ''}")
        return "\n".join(l for l in lines if l.strip())

    @command("expand", "expander", group=G, usage="/expand add addr = <text> · /expand",
             help="text shortcuts: type addr, press Ctrl+Alt+E anywhere")
    def expand(self, args: str, routed: bool = False):
        text = args.strip()
        items = EXPANSIONS.load()
        m = re.match(r"^(?:add|set)\s+(\S+)\s*=\s*(.+)$", text, re.I | re.S)
        if m:
            key = m.group(1).strip().lstrip(";").lower()
            items[key] = m.group(2).replace("\\n", "\n")
            EXPANSIONS.save(items)
            return f"⌨ '{key}' → {items[key][:80]}\nType {key} anywhere, then press Ctrl+Alt+E to expand it."
        if text.lower().startswith(("delete ", "remove ")):
            key = text.split(None, 1)[1].strip().lstrip(";").lower()
            if items.pop(key, None) is None:
                return f"No shortcut '{key}'."
            EXPANSIONS.save(items)
            return f"Deleted '{key}'."
        if not items:
            return "No text shortcuts yet.\n  /expand add addr = Moda Cd. No:12, Kadıköy, İstanbul\n  /expand add sig = Best regards,\\nAhmed\nThen type addr anywhere and press Ctrl+Alt+E."
        lines = ["Text shortcuts (type one, press Ctrl+Alt+E):"]
        lines += [f"  {k:<10} {v[:70].replace(chr(10), ' / ')}" for k, v in sorted(items.items())]
        return "\n".join(lines) + "\n/expand add <key> = <text> · /expand delete <key>"

    @command("screentime", group=G, usage="/screentime [week] · /screentime off",
             help="how long you spent in each app today")
    def screentime(self, args: str, routed: bool = False):
        word = args.strip().lower()
        if word in {"off", "on"}:
            self._write_setting("JARVIS_SCREENTIME", word)
            if word == "off":
                screen_time.stop()
            else:
                screen_time.start()
            return f"Screen time tracking {word}. It only ever counts while JARVIS is running, and stays on this PC."
        screen_time.flush()
        data = SCREEN_TIME.load()
        days = 7 if word in {"week", "7"} else 1
        totals: dict[str, float] = defaultdict(float)
        for offset in range(days):
            for app, seconds in data.get((date.today() - timedelta(days=offset)).isoformat(), {}).items():
                totals[app] += seconds
        if not totals:
            return "No screen time recorded yet. It's counted while JARVIS is open (and not idle). /screentime off stops it."
        total = sum(totals.values())
        lines = [f"⏳ Screen time {'this week' if days > 1 else 'today'}: {reminders.describe_seconds(total)}"]
        for app, seconds in sorted(totals.items(), key=lambda kv: -kv[1])[:12]:
            bar = "█" * max(1, round(seconds / total * 30))
            lines.append(f"  {app[:18]:<18} {reminders.describe_seconds(seconds):>8}  {bar}")
        return "\n".join(lines) + "\n(Counted while JARVIS runs; idle time is skipped. /screentime week)"

    @command("cleanup", "clean", group=G, usage="/cleanup · /cleanup go · /cleanup bin",
             help="frees space: old temp files and the Recycle Bin")
    def cleanup(self, args: str, routed: bool = False):
        word = args.strip().lower()
        files, size = temp_files()
        bin_size, bin_items = recycle_bin_size() if sys.platform == "win32" else (0, 0)
        if word == "go":
            if not files:
                return "No old temp files to remove."
            if not security.permissions.ask(security.DELETE_FILE, f"delete {len(files)} temp file(s) older than a day ({kit.size(size)})",
                                            context="/cleanup"):
                return "Denied."
            freed = removed = 0
            for path in files:
                try:
                    length = path.stat().st_size
                    path.unlink()
                    freed += length
                    removed += 1
                except OSError:
                    continue      # in use — that's fine
            security.audit.record("cleanup", f"{removed} temp files", kit.size(freed))
            return f"🧹 Removed {removed} temp file(s), freed {kit.size(freed)}. (Files in use were skipped.)"
        if word in {"bin", "recycle", "trash"}:
            if not bin_items:
                return "The Recycle Bin is already empty."
            if not security.permissions.ask(security.DELETE_FILE, f"empty the Recycle Bin — {bin_items} item(s), {kit.size(bin_size)}, permanently",
                                            context="/cleanup"):
                return "Denied."
            import ctypes

            ctypes.windll.shell32.SHEmptyRecycleBinW(None, None, 0x1 | 0x2 | 0x4)
            security.audit.record("cleanup", "recycle bin", kit.size(bin_size))
            return f"🧹 Emptied the Recycle Bin ({kit.size(bin_size)})."
        return (f"Space you can free:\n  Temp files older than a day   {kit.size(size)} ({len(files)} files)   → /cleanup go\n"
                f"  Recycle Bin                    {kit.size(bin_size)} ({bin_items} items)   → /cleanup bin\n"
                "/dupes finds duplicate files too.")

    @command("dnd", "donotdisturb", "focusmode", group=G, usage="/dnd on [for 1h] · /dnd off",
             help="silences Windows notifications and JARVIS")
    def dnd(self, args: str, routed: bool = False):
        if (problem := _windows_only()):
            return problem
        text = args.strip().lower()
        on = not text.startswith("off")
        if not security.permissions.ask(security.CHANGE_SETTING, f"turn Windows notifications {'off' if on else 'on'}", context="/dnd"):
            return "Denied."
        try:
            notifications(not on)
        except OSError as exc:
            return f"Couldn't change it: {exc}"
        from . import config

        os.environ["JARVIS_QUIET"] = "1" if on else ""
        reminders.board.cancel_kind("dnd")
        tail = ""
        m = re.search(r"for\s+(\d+)\s*(m|min|minutes?|h|hours?)", text)
        if on and m:
            seconds = int(m.group(1)) * (60 if m.group(2).startswith("m") else 3600)
            reminders.board.add("dnd", "Do not disturb ends", time.time() + seconds)
            tail = f" for {reminders.describe_seconds(seconds)}"
        return ("🔕 Do not disturb on" + tail + ". Windows banners are off and I'll stay quiet (reminders still show here). /dnd off"
                if on else "🔔 Do not disturb off.")
