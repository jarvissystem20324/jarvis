"""/data — ask questions about a spreadsheet and get exact numbers.

    /data sales.xlsx                      what is in it
    /data which region sold the most?     asks about the open file
    /data chart revenue by month          draws it

The model never computes anything and never runs code. It sees the column
names, their types and a handful of sample rows, and answers with a small
JSON query — filter, group, one of sum/mean/count/min/max/median. This module
runs that query itself over every row, so the numbers are exact rather than
estimated by a model reading a sample. Charts are drawn with Pillow.

That keeps the promise Sandbox mode made: code from an AI is analysed here,
never executed.
"""

from __future__ import annotations

import csv
import json
import math
import re
import statistics
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

MAX_ROWS = 200_000
SAMPLE_ROWS = 6
METRICS = ("count", "sum", "mean", "min", "max", "median")
OPS = ("==", "!=", ">", "<", ">=", "<=", "contains", "startswith")


class DataError(Exception):
    pass


@dataclass
class Table:
    path: Path
    headers: list[str]
    rows: list[list]
    types: dict[str, str] = field(default_factory=dict)   # number | date | text
    truncated: bool = False

    def column(self, name: str) -> int:
        lookup = {h.lower().strip(): i for i, h in enumerate(self.headers)}
        index = lookup.get(str(name).lower().strip())
        if index is None:
            raise DataError(f"There is no column called '{name}'. Columns: {', '.join(self.headers)}")
        return index


# --- loading -------------------------------------------------------------------

def _number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if isinstance(value, float) and math.isnan(value) else float(value)
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return None
    text = re.sub(r"^[€$£₺]|[€$£₺%]$", "", text).strip()
    # 1.234,56 (Turkish/European) and 1,234.56 (English) both occur.
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", text):
        text = text.replace(",", "")
    elif re.fullmatch(r"-?\d+,\d+", text):
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _sniff_types(headers, rows) -> dict[str, str]:
    types = {}
    for index, header in enumerate(headers):
        values = [r[index] for r in rows[:500] if index < len(r) and str(r[index]).strip() != ""]
        if not values:
            types[header] = "text"
            continue
        numeric = sum(1 for v in values if _number(v) is not None)
        dated = sum(1 for v in values if isinstance(v, datetime) or re.match(r"^\d{4}-\d{2}-\d{2}", str(v)))
        if numeric >= 0.9 * len(values):
            types[header] = "number"
        elif dated >= 0.9 * len(values):
            types[header] = "date"
        else:
            types[header] = "text"
    return types


def load(path: Path) -> Table:
    path = Path(path).expanduser()
    if not path.is_file():
        raise DataError(f"No such file:\n  {path}")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise DataError("Excel files need openpyxl (pip install openpyxl).")
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        raw = [list(r) for r in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif suffix in {".csv", ".tsv", ".txt"}:
        text = None
        for encoding in ("utf-8-sig", "cp1254", "latin-1"):
            try:
                text = path.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        raw = list(csv.reader(text.splitlines(), dialect))
    else:
        raise DataError("I read .csv, .tsv and .xlsx files.")

    raw = [r for r in raw if any(str(c or "").strip() for c in r)]
    if len(raw) < 2:
        raise DataError(f"{path.name} has no data rows.")
    headers = [str(h).strip() if h not in (None, "") else f"column {i + 1}" for i, h in enumerate(raw[0])]
    rows = [list(r) + [None] * (len(headers) - len(r)) for r in raw[1:MAX_ROWS + 1]]
    table = Table(path, headers, rows, truncated=len(raw) - 1 > MAX_ROWS)
    table.types = _sniff_types(headers, rows)
    return table


def profile(table: Table) -> str:
    """Columns, types and ranges — what the model sees, and what you see first."""
    lines = [f"{table.path.name}: {len(table.rows):,} rows × {len(table.headers)} columns"
             + (" (first 200,000 rows only)" if table.truncated else "")]
    for index, header in enumerate(table.headers):
        kind = table.types.get(header, "text")
        values = [r[index] for r in table.rows if str(r[index] or "").strip() != ""]
        if kind == "number":
            nums = [n for n in (_number(v) for v in values) if n is not None]
            detail = (f"{min(nums):,.2f} … {max(nums):,.2f}, mean {statistics.fmean(nums):,.2f}"
                      if nums else "empty")
        else:
            distinct = len({str(v) for v in values})
            common = sorted({str(v) for v in values}, key=lambda v: -sum(1 for x in values if str(x) == v))[:3] \
                if distinct <= 50 else []
            detail = f"{distinct:,} distinct" + (f" (e.g. {', '.join(common)})" if common else "")
        missing = len(table.rows) - len(values)
        lines.append(f"  {header:<24} {kind:<7} {detail}" + (f", {missing} blank" if missing else ""))
    return "\n".join(lines)


def sample(table: Table, n: int = SAMPLE_ROWS) -> str:
    rows = [" | ".join(str(c if c is not None else "") for c in r) for r in table.rows[:n]]
    return " | ".join(table.headers) + "\n" + "\n".join(rows)


# --- queries -------------------------------------------------------------------

QUERY_PROMPT = """You answer questions about a table by writing ONE JSON query.
You never see all the rows; the query is run exactly over every row.

Table columns and types:
{profile}

First rows:
{sample}

Reply with JSON only, this shape:
{{"filters": [{{"column": "...", "op": "==|!=|>|<|>=|<=|contains|startswith", "value": ...}}],
  "group_by": "column name or null",
  "metric": "count|sum|mean|min|max|median",
  "column": "column to measure (null for count)",
  "sort": "desc|asc",
  "limit": 10,
  "chart": "bar|line|none",
  "title": "short title for the answer"}}

If the question cannot be answered from these columns, reply
{{"error": "why not"}}.

Question: {question}"""


def _match(value, op: str, target) -> bool:
    if op in {"contains", "startswith"}:
        text, needle = str(value or "").lower(), str(target).lower()
        return needle in text if op == "contains" else text.startswith(needle)
    left, right = _number(value), _number(target)
    if left is None or right is None:
        left, right = str(value or "").strip().lower(), str(target).strip().lower()
    try:
        return {"==": left == right, "!=": left != right, ">": left > right,
                "<": left < right, ">=": left >= right, "<=": left <= right}[op]
    except TypeError:
        return False


def _aggregate(values: list, metric: str) -> float:
    if metric == "count":
        return float(len(values))
    nums = [n for n in (_number(v) for v in values) if n is not None]
    if not nums:
        return float("nan")
    return {"sum": sum(nums), "mean": statistics.fmean(nums), "min": min(nums),
            "max": max(nums), "median": statistics.median(nums)}[metric]


def run(table: Table, query: dict) -> tuple[list[tuple[str, float]], str]:
    """Execute a query. Returns ([(label, value)], description)."""
    if "error" in query:
        raise DataError(str(query["error"]))
    metric = str(query.get("metric") or "count").lower()
    if metric not in METRICS:
        raise DataError(f"Unknown measure '{metric}'.")
    rows = table.rows
    described = []
    for condition in query.get("filters") or []:
        op = str(condition.get("op") or "==")
        if op not in OPS:
            raise DataError(f"Unknown comparison '{op}'.")
        index = table.column(condition.get("column"))
        rows = [r for r in rows if _match(r[index], op, condition.get("value"))]
        described.append(f"{table.headers[index]} {op} {condition.get('value')}")
    measure = query.get("column")
    measure_index = table.column(measure) if measure and metric != "count" else None

    def pick(r):
        return r[measure_index] if measure_index is not None else 1

    group = query.get("group_by")
    if group:
        index = table.column(group)
        buckets: dict[str, list] = {}
        for r in rows:
            key = r[index]
            if isinstance(key, datetime):
                key = key.strftime("%Y-%m-%d")
            buckets.setdefault(str(key if key not in (None, "") else "(blank)"), []).append(pick(r))
        results = [(k, _aggregate(v, metric)) for k, v in buckets.items()]
        reverse = str(query.get("sort") or "desc").lower() != "asc"
        if query.get("chart") == "line":
            results.sort(key=lambda kv: kv[0])
        else:
            results.sort(key=lambda kv: (-math.inf if math.isnan(kv[1]) else kv[1]), reverse=reverse)
        limit = max(1, min(int(query.get("limit") or 10), 50))
        results = results[:limit] if query.get("chart") != "line" else results[:200]
    else:
        results = [(metric, _aggregate([pick(r) for r in rows], metric))]
    what = f"{metric}" + (f" of {table.headers[measure_index]}" if measure_index is not None else " of rows")
    how = what + (f" by {group}" if group else "") + (f", where {' and '.join(described)}" if described else "")
    return results, f"{how} ({len(rows):,} matching rows)"


def parse_query(text: str) -> dict:
    body = text.strip()
    fence = re.search(r"```(?:json)?\s*(.+?)```", body, re.S)
    if fence:
        body = fence.group(1)
    start, end = body.find("{"), body.rfind("}")
    if start == -1 or end <= start:
        raise DataError("The model didn't return a query I could run.")
    try:
        return json.loads(body[start:end + 1])
    except ValueError as exc:
        raise DataError(f"The query wasn't valid JSON: {exc}")


def format_results(results: list[tuple[str, float]]) -> str:
    def fmt(v: float) -> str:
        if math.isnan(v):
            return "—"
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"
    if len(results) == 1 and results[0][0] in METRICS:
        return fmt(results[0][1])
    width = min(max(len(k) for k, _ in results), 32)
    return "\n".join(f"  {k[:32]:<{width}}  {fmt(v):>14}" for k, v in results)


# --- charts --------------------------------------------------------------------

def chart(results: list[tuple[str, float]], title: str, kind: str, out: Path) -> Path:
    from PIL import Image, ImageDraw, ImageFont

    points = [(k, v) for k, v in results if not math.isnan(v)]
    if not points:
        raise DataError("Nothing to draw.")
    width, height, pad_l, pad_b, pad_t, pad_r = 1000, 560, 90, 120, 60, 30
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 15)
        big = ImageFont.truetype("arialbd.ttf", 20)
    except OSError:
        font = big = ImageFont.load_default()
    draw.text((pad_l, 18), title[:80], fill="#111827", font=big)
    top = max(v for _, v in points)
    low = min(0.0, min(v for _, v in points))
    span = (top - low) or 1.0
    plot_w, plot_h = width - pad_l - pad_r, height - pad_t - pad_b

    def y_of(v: float) -> float:
        return pad_t + plot_h - (v - low) / span * plot_h

    for i in range(5):
        value = low + span * i / 4
        y = y_of(value)
        draw.line([(pad_l, y), (width - pad_r, y)], fill="#e5e7eb")
        draw.text((8, y - 8), f"{value:,.0f}" if abs(value) >= 10 else f"{value:,.2f}", fill="#6b7280", font=font)
    step = plot_w / len(points)
    coords = []
    for index, (label, value) in enumerate(points):
        x = pad_l + step * index + step / 2
        coords.append((x, y_of(value)))
        if kind == "bar":
            draw.rectangle([x - step * 0.35, y_of(value), x + step * 0.35, y_of(0 if low <= 0 else low)],
                           fill="#3b82f6")
        if len(points) <= 24 or index % max(1, len(points) // 12) == 0:
            text = label[:14]
            draw.text((x - len(text) * 3.5, height - pad_b + 10), text, fill="#374151", font=font)
    if kind == "line" and len(coords) > 1:
        draw.line(coords, fill="#3b82f6", width=3)
        for x, y in coords:
            draw.ellipse([x - 3, y - 3, x + 3, y + 3], fill="#1d4ed8")
    image.save(out)
    return out
