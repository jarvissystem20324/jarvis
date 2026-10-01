"""8.0 — Make and write: slides, email, CV, cover letters, dilekçe, citations,
essay outlines, grammar fixes, invoices, memes, mind maps, flowcharts, social
posts and one-page websites.

The model writes content; the file formats, totals, layouts and citation
styles are done here, so a deck always opens, an invoice always adds up and
a DOI citation is the publisher's own record rather than a model's memory.
"""

from __future__ import annotations

import html
import json
import re
import urllib.parse
from datetime import date, datetime, timedelta
from pathlib import Path

from . import drawing, kit, security, shield, writer
from .registry import command

G = "Make and write"


def _deck_prompt(about: str, count: int) -> str:
    return f"""Plan a presentation of {count} content slides {about}.
JSON: {{"title": "...", "subtitle": "...", "slides": [{{"title": "...", "bullets": ["...", "..."], "notes": "what to say"}}]}}
- 3-5 bullets per slide, each under 12 words; speaker notes 2-4 sentences
- a clear story: intro, the main points, a conclusion slide last
- write in the language of the request"""


class Makers:
    # --- slides --------------------------------------------------------------

    @command("slides", "deck", "pptx", group=G, usage="/slides <topic|file> [12 slides]",
             help="a PowerPoint deck with speaker notes; /slides revise: <change>")
    def slides(self, args: str, routed: bool = False):
        text = args.strip()
        if not text:
            return "Usage: /slides <topic or file>   e.g. /slides 10 slides on renewable energy\n       /slides revise: add a slide about costs"
        revise = re.match(r"^(?:revise|edit|change)\s*:?\s*(.+)$", text, re.I | re.S)
        last = getattr(self, "_last_deck", None)
        try:
            if revise and last:
                deck = kit.ask_json(self.brain, "Here is a presentation as JSON. Change it as asked and return the whole "
                                    f"presentation in the same JSON form.\nRequest: {revise.group(1)}\n\n{json.dumps(last, ensure_ascii=False)}")
            else:
                count = 8
                m = re.search(r"\b(\d{1,2})\s*(?:slides?|slayt)\b", text, re.I)
                if m:
                    count = max(3, min(int(m.group(1)), 30))
                    text = (text[:m.start()] + text[m.end():]).strip(" ,") or text
                text = re.sub(r"^(?:on|about)\s+", "", text, flags=re.I)
                label, material = self._material(text)
                deck = kit.ask_json(self.brain, _deck_prompt(self._about(label, material, text), count))
        except (kit.KitError, Exception) as exc:   # a denied file read raises QuizError
            return str(exc)
        slides = [s for s in deck.get("slides", []) if isinstance(s, dict) and s.get("title")] if isinstance(deck, dict) else []
        if not slides:
            return "No slides came back. Try again."
        deck["slides"] = slides
        path = kit.output_dir() / f"{kit.slug(deck.get('title') or text)}_{kit.stamp()}.pptx"
        try:
            build_pptx(deck, path)
        except ImportError:
            return "Slides need python-pptx (pip install python-pptx)."
        self._last_deck = deck
        security.audit.record("slides", str(deck.get("title"))[:80], f"{len(slides)} slides")
        outline = "\n".join(f"  {i}. {s['title']}" for i, s in enumerate(slides, 1))
        return (f"📊 {deck.get('title', 'Presentation')} — {len(slides) + 1} slides\n{outline}\n\n{path}\n"
                "/slides revise: <change> edits it.")

    # --- email -----------------------------------------------------------------

    @command("email", group=G, usage="/email <to whom, about what>",
             help="drafts an email; /email open (mail app) or /email gmail")
    def email(self, args: str, routed: bool = False):
        text = args.strip()
        draft = getattr(self, "_email", None)
        if text.lower() in {"open", "send", "mail"} or text.lower() == "gmail":
            if not draft:
                return "Write one first: /email <to whom, about what>"
            fields = {"subject": draft["subject"], "body": draft["body"]}
            if text.lower() == "gmail":
                url = "https://mail.google.com/mail/?" + urllib.parse.urlencode(
                    {"view": "cm", "fs": "1", "to": draft.get("to", ""), "su": fields["subject"], "body": fields["body"]})
            else:
                url = f"mailto:{urllib.parse.quote(draft.get('to', ''))}?" + urllib.parse.urlencode(fields, quote_via=urllib.parse.quote)
            import webbrowser

            webbrowser.open(url)
            return "Opened it ready to send — check it over and press Send yourself."
        if not text:
            return "Usage: /email <who it's to and what it's about>   e.g. /email my landlord, the heating is broken since Monday"
        try:
            data = kit.ask_json(self.brain, f"""Write an email for this request: {text}
JSON: {{"subject": "...", "body": "..."}} — the body with greeting and sign-off, in the language of the request,
polite and to the point. Use [Your name] where the sender's name goes.""")
        except kit.KitError as exc:
            return str(exc)
        to = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
        self._email = {"to": to.group(0) if to else "", "subject": str(data.get("subject", "")), "body": str(data.get("body", ""))}
        return (f"Subject: {self._email['subject']}\n\n{self._email['body']}\n\n"
                "/email open puts it in your mail app · /email gmail opens Gmail · say what to change and /email again")

    # --- CV, cover letter, dilekçe ----------------------------------------------

    def _document(self, prompt: str, formats=("docx", "pdf")) -> str:
        with kit.more_room(self.brain):
            markdown = writer.clean_markdown(self.brain.ask_once(prompt))
        if not markdown.startswith("#"):
            return f"No document came back:\n{markdown[:300]}"
        try:
            paths = writer.save(markdown, formats)
        except writer.WriterError as exc:
            return str(exc)
        self._last_document = markdown
        security.audit.record("document", writer.title_of(markdown)[:80], ", ".join(p.suffix for p in paths))
        return f"📄 {writer.title_of(markdown)}\n" + "\n".join(f"  {p}" for p in paths) + \
            "\n/makedoc revise: <change> edits it."

    @command("cv", "resume", group=G, usage="/cv <your details>", help="a CV in Word and PDF from your details")
    def cv(self, args: str, routed: bool = False):
        if not args.strip():
            return ("Usage: /cv <everything about you>\n"
                    "e.g. /cv Ahmed, Istanbul, CS student at İTÜ (2027), Python, React, built JARVIS (desktop AI), "
                    "intern at X 2025, English C1, email a@b.com")
        return self._document(f"""{writer.PROMPT}
Write a one-page CV (résumé) from these details. Sections: name and contact line, a 2-line profile,
Experience, Projects, Education, Skills, Languages — omit any with nothing to say. Strong action verbs,
achievements with numbers where the details give them. Never invent employers, dates or grades.
Details: {args.strip()}""")

    @command("coverletter", "cover", group=G, usage="/coverletter <job ad or link>",
             help="a cover letter for a job, using your open CV if there is one")
    def cover_letter(self, args: str, routed: bool = False):
        text = args.strip()
        if not text:
            return "Usage: /coverletter <the job ad, or a link to it>   (open your CV first with /doc cv.pdf to tailor it)"
        if re.match(r"https?://", text):
            from . import websearch

            try:
                title, page = websearch.fetch(text.split()[0])
            except Exception as exc:
                return f"Couldn't read that page: {exc}"
            text = f"{title}\n{page[:8000]}"
        job, _ = shield.wrap(text[:8000], "the job ad")
        doc = self._document_addon()
        cv = shield.wrap(doc.text[:8000], doc._label())[0] if doc is not None and doc.text else "(no CV open — keep claims general)"
        return self._document(f"""{writer.PROMPT}
{shield.RULE}
Write a one-page cover letter for this job, matching the candidate's real experience to what the ad asks.
Never invent experience. Address the hiring manager; three or four short paragraphs; confident, not flowery.
Job ad: {job}
Candidate's CV: {cv}""")

    @command("dilekce", "dilekçe", "petition", group=G, usage="/dilekce <konu>",
             help="a formal Turkish petition (dilekçe) in Word and PDF")
    def dilekce(self, args: str, routed: bool = False):
        if not args.strip():
            return ("Kullanım: /dilekce <kime, ne için, ayrıntılar>\n"
                    "örn. /dilekce İTÜ öğrenci işlerine, sağlık raporum nedeniyle vize sınavına mazeret sınavı talebi")
        return self._document(f"""{writer.PROMPT}
Write a formal Turkish dilekçe (petition) in correct official format:
- "# <KURUM ADINA / MAKAMINA>" as the title line (the addressee, in capitals, e.g. "# İSTANBUL TEKNİK ÜNİVERSİTESİ ÖĞRENCİ İŞLERİ DAİRE BAŞKANLIĞINA")
- the request paragraph in formal Turkish, ending with "Gereğinin yapılmasını saygılarımla arz ederim."
- then the date ({date.today():%d.%m.%Y}), "Ad Soyad:", "İmza:", "Adres:", "Telefon:", "T.C. Kimlik No:" lines left
  blank for the person to fill, and an "Ekler:" line if documents are mentioned.
Request: {args.strip()}""")

    # --- citations ----------------------------------------------------------------

    @command("cite", "citation", group=G, usage="/cite <url|doi|isbn> [apa|mla|chicago|harvard]",
             help="a citation; DOIs come from the publisher's record")
    def cite(self, args: str, routed: bool = False):
        text = args.strip()
        style = "apa"
        m = re.search(r"\b(apa|mla|chicago|harvard)\b\s*$", text, re.I)
        if m:
            style, text = m.group(1).lower(), text[:m.start()].strip()
        if not text:
            return "Usage: /cite <link, DOI or ISBN> [apa|mla|chicago|harvard]   e.g. /cite 10.1038/nature14539 mla"
        try:
            return f"{cite(text, style)}\n\n({style.upper()})"
        except kit.KitError as exc:
            return str(exc)

    # --- essays and grammar --------------------------------------------------------

    @command("outline", group=G, usage="/outline <essay topic> [1500 words]", help="an essay outline with a thesis")
    def outline(self, args: str, routed: bool = False):
        if not args.strip():
            return "Usage: /outline <essay question or topic> [length]   e.g. /outline Should social media be regulated? 1500 words"
        return self.brain.ask_once(f"""Write an essay outline for: {args.strip()}
Give: a one-sentence thesis; an introduction plan (hook, context, thesis); 3-5 body sections, each with a topic
sentence, 2-3 supporting points and the kind of evidence to find; a counter-argument and rebuttal; a conclusion plan.
Suggest a word count per section if a length is given. Markdown, in the language of the request.""")

    @command("grammar", "proofread", group=G, usage="/grammar <file.docx|txt|md>",
             help="fixes spelling and grammar in a file, writes a corrected copy")
    def grammar(self, args: str, routed: bool = False):
        path = kit.path_arg(args)
        if path is None or path.suffix.lower() not in {".docx", ".txt", ".md"}:
            return "Usage: /grammar <file.docx, .txt or .md>   — a corrected copy is written; the original is untouched."
        if not security.permissions.ask(security.READ_FILE, str(path), context="/grammar"):
            return "Denied. The file was not read."
        if path.suffix.lower() == ".docx":
            from docx import Document

            document = Document(str(path))
            paragraphs = [p for p in document.paragraphs if p.text.strip()]
            texts = [p.text for p in paragraphs]
        else:
            document = None
            texts = path.read_text(encoding="utf-8", errors="replace").split("\n")
        fixed = list(texts)
        batch: list[int] = []
        for i, t in enumerate(texts + [None]):
            if t is not None and t.strip():
                batch.append(i)
            if batch and (t is None or sum(len(texts[j]) for j in batch) > 5000 or len(batch) >= 40):
                originals = [texts[j] for j in batch]
                try:
                    result = kit.ask_json(self.brain, "Correct the spelling, grammar and punctuation of each string. Keep "
                                          "the meaning, wording, tone and language; change nothing that is already correct. "
                                          f"Return a JSON list of exactly {len(originals)} strings, in order.\n"
                                          + json.dumps(originals, ensure_ascii=False))
                except kit.KitError:
                    result = None
                if isinstance(result, list) and len(result) == len(originals):
                    for j, new in zip(batch, result):
                        fixed[j] = str(new)
                batch = []
        changed = [(a, b) for a, b in zip(texts, fixed) if a.strip() != b.strip()]
        out = kit.output_dir() / f"{path.stem}_corrected{path.suffix}"
        if document is not None:
            for paragraph, old, new in zip(paragraphs, texts, fixed):
                if old != new:
                    style = paragraph.style
                    paragraph.text = new
                    paragraph.style = style
            document.save(str(out))
        else:
            out.write_text("\n".join(fixed), encoding="utf-8")
        security.audit.record("grammar", path.name, f"{len(changed)} changed")
        examples = "\n".join(f"  - {a.strip()[:70]}\n    → {b.strip()[:70]}" for a, b in changed[:5])
        return (f"✍ {len(changed)} of {len([t for t in texts if t.strip()])} paragraphs corrected in {path.name}.\n"
                f"{examples}\n\nCorrected copy: {out}")

    # --- invoices --------------------------------------------------------------------

    @command("invoice", group=G, usage="/invoice <who, what, prices>", help="an invoice PDF; totals are exact")
    def invoice(self, args: str, routed: bool = False):
        if not args.strip():
            return ("Usage: /invoice <from, to, items and prices>\n"
                    "e.g. /invoice from Ahmed Dilmen to Acme Ltd: website design 1 x 12000 TL, hosting 12 months x 150 TL, KDV 20%")
        try:
            data = kit.ask_json(self.brain, f"""Extract an invoice from this request: {args.strip()}
JSON: {{"from": "seller name and details", "to": "buyer name and details", "currency": "TRY|USD|EUR|...",
"items": [{{"description": "...", "quantity": 1, "unit_price": 0}}], "tax_percent": 0, "notes": "", "due_days": 14}}
Numbers as plain numbers. Do not compute totals.""")
            path = make_invoice(data)
        except (kit.KitError, ValueError) as exc:
            return f"Couldn't make the invoice: {exc}"
        return f"🧾 Invoice saved:\n  {path}"

    # --- pictures ----------------------------------------------------------------------

    @command("meme", group=G, usage="/meme <top> | <bottom>", help="meme text on the last image (or /meme <idea>)")
    def meme(self, args: str, routed: bool = False):
        from .assistant import JarvisResponse

        source = getattr(self, "current_image", None)
        text = args.strip()
        path = None
        if text and "|" not in text:
            first = text.split()[0]
            path = kit.path_arg(first)
            if path is not None:
                source, text = path, text[len(first):].strip()
        if not source or not Path(source).exists():
            return "Make or open an image first (/image <prompt>, or drop one in), then /meme <top> | <bottom>."
        if not text:
            return "Usage: /meme <top text> | <bottom text>   or /meme <idea> and I'll write the captions."
        if "|" in text:
            top, _, bottom = text.partition("|")
        else:
            try:
                captions = kit.ask_json(self.brain, f'Write a short, funny two-line meme caption about: {text}. '
                                        'JSON: {"top": "...", "bottom": "..."}, each under 8 words.')
                top, bottom = str(captions.get("top", "")), str(captions.get("bottom", ""))
            except kit.KitError as exc:
                return str(exc)
        out = kit.output_dir("images") / f"meme_{kit.stamp()}.png"
        drawing.meme(Path(source), top, bottom, out)
        self.current_image = out
        return JarvisResponse(text=f"Meme saved to {out}", image_path=out, image_paths=[out])

    @command("mindmap", group=G, usage="/mindmap <topic|file>", help="a mind map picture")
    def mind_map(self, args: str, routed: bool = False):
        from .assistant import JarvisResponse

        if not args.strip():
            return "Usage: /mindmap <topic or file>   e.g. /mindmap causes of World War I"
        try:
            label, material = self._material(args.strip())
            tree = kit.ask_json(self.brain, f"""Make a mind map {self._about(label, material, args.strip())}
JSON: {{"center": "short topic", "branches": [{{"name": "2-4 words", "children": ["2-5 words", ...]}}]}}
4-7 branches, 2-4 children each. Language of the request.""")
        except Exception as exc:
            return str(exc)
        out = kit.output_dir("images") / f"mindmap_{kit.slug(str(tree.get('center', 'map')))}_{kit.stamp()}.png"
        drawing.mind_map(tree, out)
        self.current_image = out
        return JarvisResponse(text=f"🧠 Mind map: {tree.get('center', '')}\n{out}", image_path=out, image_paths=[out])

    @command("flowchart", "flow", group=G, usage="/flowchart <process>", help="a flowchart picture from a description")
    def flowchart(self, args: str, routed: bool = False):
        from .assistant import JarvisResponse

        if not args.strip():
            return "Usage: /flowchart <describe the process>   e.g. /flowchart how a user logs in with 2FA"
        try:
            chart = kit.ask_json(self.brain, f"""Draw this process as a flowchart: {args.strip()}
JSON: {{"nodes": [{{"id": "a", "text": "short label", "type": "start|step|decision|io|end"}}],
"edges": [{{"from": "a", "to": "b", "label": "yes/no or empty"}}]}}
One start, at least one end, decisions have labelled yes/no edges, at most 16 nodes, labels under 6 words.""")
            out = kit.output_dir("images") / f"flowchart_{kit.stamp()}.png"
            drawing.flowchart(chart, out)
        except (kit.KitError, ValueError, KeyError) as exc:
            return f"Couldn't draw it: {exc}"
        self.current_image = out
        return JarvisResponse(text=f"Flowchart saved to {out}", image_path=out, image_paths=[out])

    # --- posts and websites ---------------------------------------------------------------

    @command("post", group=G, usage="/post <topic> [x|linkedin|instagram]", help="social media posts with hashtags")
    def post(self, args: str, routed: bool = False):
        text = args.strip()
        if not text:
            return "Usage: /post <what it's about> [x|linkedin|instagram]   e.g. /post I shipped JARVIS 8.0 linkedin"
        m = re.search(r"\b(x|twitter|linkedin|instagram|insta)\s*$", text, re.I)
        platforms = {"x": "X (under 280 characters)", "twitter": "X (under 280 characters)", "linkedin": "LinkedIn",
                     "instagram": "Instagram", "insta": "Instagram"}
        where = platforms[m.group(1).lower()] if m else "X (under 280 characters), LinkedIn and Instagram"
        topic = text[:m.start()].strip() if m else text
        return self.brain.ask_once(f"Write a social media post for {where} about: {topic}. "
                                   "Natural voice, a hook in the first line, 3-6 relevant hashtags, emojis only where "
                                   "they fit the platform. If several platforms, one post per platform under a heading. "
                                   "Same language as the request.")

    @command("website", "webpage", group=G, usage="/website <what it's for>",
             help="a one-page website (.html) that opens in your browser")
    def website(self, args: str, routed: bool = False):
        if not args.strip():
            return "Usage: /website <what it's for>   e.g. /website a portfolio for a CS student who builds AI apps"
        with kit.more_room(self.brain, 8000):
            answer = self.brain.ask_once(f"""Write a complete, beautiful one-page website for: {args.strip()}
One self-contained HTML file: inline <style>, no external scripts or images (use CSS shapes, gradients and emoji),
responsive, semantic HTML, a nav with anchors to sections, a footer. Real content, no lorem ipsum.
Same language as the request. Output only the HTML.""")
        body = answer.strip()
        fence = re.search(r"```(?:html)?\s*(.*?)```", body, re.S | re.I)
        if fence:
            body = fence.group(1).strip()
        if "<html" not in body.lower():
            return f"No page came back:\n{body[:300]}"
        title = re.search(r"<title>(.*?)</title>", body, re.I | re.S)
        out = kit.output_dir("websites") / f"{kit.slug(html.unescape(title.group(1)) if title else 'site')}_{kit.stamp()}.html"
        out.write_text(body, encoding="utf-8")
        from . import pc

        try:
            pc.open_path(out)
        except Exception:
            pass
        return f"🌐 Website saved and opened:\n  {out}\nAsk for changes and run /website again, or edit the file."


# --- PowerPoint ---------------------------------------------------------------------

def build_pptx(deck: dict, path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    navy, accent, light, muted = RGBColor(0x14, 0x1B, 0x2D), RGBColor(0x4F, 0x8E, 0xF7), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xB8, 0xC2, 0xD6)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def background(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def text(slide, left, top, width, height, value, size, color, bold=False):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = str(value)
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, color
                run.font.name = "Segoe UI"
        return frame

    def bar(slide, left, top, width, height, color):
        shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    first = prs.slides.add_slide(blank)
    background(first, navy)
    bar(first, 0.8, 3.05, 1.2, 0.08, accent)
    text(first, 0.8, 1.6, 11.5, 1.4, deck.get("title") or "Presentation", 44, light, bold=True)
    text(first, 0.8, 3.35, 11.5, 1.0, deck.get("subtitle") or "", 22, muted)
    for number, item in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(blank)
        background(slide, RGBColor(0xFF, 0xFF, 0xFF))
        bar(slide, 0, 0, 0.18, 7.5, accent)
        text(slide, 0.7, 0.45, 12, 1.0, item["title"], 32, navy, bold=True)
        bar(slide, 0.7, 1.45, 0.9, 0.06, accent)
        frame = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.8), Inches(5)).text_frame
        frame.word_wrap = True
        for i, bullet in enumerate([str(b) for b in item.get("bullets", [])][:7]):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"•  {bullet}"
            run.font.size, run.font.color.rgb, run.font.name = Pt(22), RGBColor(0x2B, 0x31, 0x40), "Segoe UI"
            paragraph.space_after = Pt(14)
        text(slide, 12.2, 6.85, 0.9, 0.4, str(number + 1), 12, muted)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(item["notes"])
    prs.save(str(path))
    return path


# --- citations ----------------------------------------------------------------------

CSL_STYLES = {"apa": "apa", "mla": "modern-language-association", "chicago": "chicago-author-date",
              "harvard": "harvard-cite-them-right"}


def cite(source: str, style: str = "apa") -> str:
    source = source.strip()
    doi = re.search(r"\b(10\.\d{4,9}/[^\s\"<>]+)", source)
    if doi:
        try:
            return kit.get_text(f"https://doi.org/{doi.group(1).rstrip('.')}",
                                headers={"Accept": f"text/x-bibliography; style={CSL_STYLES[style]}; locale=en-US"}).strip()
        except kit.KitError as exc:
            raise kit.KitError(f"doi.org didn't recognise that DOI ({exc}).")
    isbn = re.sub(r"[^\dX]", "", source.upper())
    if re.fullmatch(r"\d{9}[\dX]|\d{13}", isbn) and not source.lower().startswith("http"):
        book = kit.get_json(f"https://openlibrary.org/isbn/{isbn}.json")
        authors = []
        for ref in book.get("authors", [])[:4]:
            try:
                authors.append(kit.get_json(f"https://openlibrary.org{ref['key']}.json").get("name", ""))
            except kit.KitError:
                continue
        year = re.search(r"\d{4}", str(book.get("publish_date", "")))
        return format_citation({"authors": [a for a in authors if a], "year": year.group(0) if year else "n.d.",
                                "title": book.get("title", ""), "publisher": ", ".join(book.get("publishers", [])[:1]),
                                "kind": "book"}, style)
    if not re.match(r"https?://", source):
        raise kit.KitError("Give a link, a DOI (10.xxxx/...) or an ISBN.")
    page = kit.get_text(source, limit=600_000)

    def meta(*names):
        for name in names:
            m = re.search(rf'<meta[^>]+(?:name|property)=["\']{re.escape(name)}["\'][^>]*content=["\']([^"\']+)', page, re.I) \
                or re.search(rf'<meta[^>]+content=["\']([^"\']+)["\'][^>]*(?:name|property)=["\']{re.escape(name)}["\']', page, re.I)
            if m:
                return html.unescape(m.group(1)).strip()
        return ""

    authors = re.findall(r'<meta[^>]+name=["\']citation_author["\'][^>]*content=["\']([^"\']+)', page, re.I) or \
        [a for a in [meta("author", "article:author")] if a and not a.startswith("http")]
    title = meta("citation_title", "og:title", "twitter:title")
    if not title:
        t = re.search(r"<title[^>]*>(.*?)</title>", page, re.I | re.S)
        title = html.unescape(t.group(1)).strip() if t else source
    published = meta("citation_publication_date", "article:published_time", "date", "dc.date", "pubdate")
    year = re.search(r"(19|20)\d{2}", published)
    site = meta("og:site_name", "citation_journal_title") or urllib.parse.urlparse(source).netloc.removeprefix("www.")
    return format_citation({"authors": authors[:6], "year": year.group(0) if year else "n.d.", "title": title,
                            "site": site, "url": source, "date": published[:10], "kind": "web"}, style)


def _split_name(name: str) -> tuple[str, str]:
    name = name.strip()
    if "," in name:
        last, first = [p.strip() for p in name.split(",", 1)]
    else:
        parts = name.split()
        last, first = (parts[-1], " ".join(parts[:-1])) if len(parts) > 1 else (name, "")
    return last, first


def format_citation(item: dict, style: str) -> str:
    authors = [a for a in item.get("authors", []) if a]
    year, title = item.get("year") or "n.d.", item.get("title", "").strip()
    container = item.get("site") or item.get("publisher") or ""
    url = item.get("url", "")
    accessed = datetime.now().strftime("%d %b. %Y")
    if style == "mla":
        if authors:
            last, first = _split_name(authors[0])
            who = f"{last}, {first}" + (", et al." if len(authors) > 2 else f", and {authors[1]}" if len(authors) == 2 else "")
            who += ". "
        else:
            who = ""
        shown = f"*{title}*" if item.get("kind") == "book" else f"“{title}.”"
        tail = f" {container}," if container else ""
        return f"{who}{shown}{tail} {year}" + (f", {url.split('://', 1)[-1]}. Accessed {accessed}." if url else ".")
    initials = lambda first: " ".join(f"{p[0]}." for p in re.split(r"[\s-]+", first) if p)  # noqa: E731
    names = [(_split_name(a)) for a in authors]
    if style == "apa":
        shown = [f"{last}, {initials(first)}".strip(", ") for last, first in names]
        who = (", ".join(shown[:-1]) + ", & " + shown[-1]) if len(shown) > 1 else (shown[0] if shown else "")
        where = f"{container}. " if container else ""
        if who:
            return f"{who} ({year}). *{title}*. {where}{url}".strip()
        return f"*{title}*. ({year}). {where}{url}".strip()
    if style == "chicago":
        shown = [f"{names[0][0]}, {names[0][1]}".strip(", ")] + [f"{f} {l}".strip() for l, f in names[1:]] if names else []
        who = (", ".join(shown[:-1]) + ", and " + shown[-1]) if len(shown) > 1 else (shown[0] if shown else "")
        return f"{who + '. ' if who else ''}{year}. “{title}.” {container + '. ' if container else ''}{url}".strip()
    # harvard
    shown = [f"{last}, {initials(first)}".strip(", ") for last, first in names]
    who = (", ".join(shown[:-1]) + " and " + shown[-1]) if len(shown) > 1 else (shown[0] if shown else container)
    return f"{who} ({year}) *{title}*. {container + '. ' if container and who != container else ''}" + \
        (f"Available at: {url} (Accessed: {datetime.now():%d %B %Y})." if url else "")


# --- invoice -------------------------------------------------------------------------

INVOICES = kit.Store("invoices.json", {"next": 1})


def invoice_totals(items: list[dict], tax_percent: float) -> tuple[list[tuple[str, float, float, float]], float, float, float]:
    rows = []
    for item in items:
        qty = float(item.get("quantity") or 1)
        price = float(item.get("unit_price") or 0)
        rows.append((str(item.get("description") or "Item"), qty, price, round(qty * price, 2)))
    subtotal = round(sum(r[3] for r in rows), 2)
    tax = round(subtotal * float(tax_percent or 0) / 100, 2)
    return rows, subtotal, tax, round(subtotal + tax, 2)


def make_invoice(data: dict, out_dir: Path | None = None) -> Path:
    from fpdf import FPDF

    items = [i for i in data.get("items", []) if isinstance(i, dict)]
    if not items:
        raise ValueError("no items found")
    rows, subtotal, tax, total = invoice_totals(items, data.get("tax_percent") or 0)
    counter = INVOICES.load()
    number = f"INV-{date.today():%Y}-{int(counter.get('next', 1)):03d}"
    counter["next"] = int(counter.get("next", 1)) + 1
    INVOICES.save(counter)
    currency = str(data.get("currency") or "TRY")
    money = lambda v: f"{v:,.2f} {currency}"  # noqa: E731
    pdf = FPDF(format="A4")
    pdf.add_page()
    fonts = writer._font_files()
    if fonts:
        pdf.add_font("Body", "", str(fonts[0]))
        pdf.add_font("Body", "B", str(fonts[1]))
        family = "Body"
    else:
        family = "Helvetica"
    pdf.set_font(family, "B", 24)
    pdf.cell(0, 12, "INVOICE", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(family, "", 10)
    issued = date.today()
    due = issued + timedelta(days=int(data.get("due_days") or 14))
    pdf.cell(0, 6, f"{number}    Date: {issued:%d.%m.%Y}    Due: {due:%d.%m.%Y}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    y = pdf.get_y()
    for x, label, who in ((10, "From", data.get("from")), (110, "Bill to", data.get("to"))):
        pdf.set_xy(x, y)
        pdf.set_font(family, "B", 10)
        pdf.cell(90, 6, label, new_x="LEFT", new_y="NEXT")
        pdf.set_font(family, "", 10)
        pdf.multi_cell(90, 5, str(who or ""))
    pdf.ln(6)
    pdf.set_font(family, "", 10)
    with pdf.table(col_widths=(90, 25, 35, 40), text_align=("LEFT", "RIGHT", "RIGHT", "RIGHT")) as table:
        table.row(["Description", "Qty", "Unit price", "Amount"])
        for description, qty, price, amount in rows:
            table.row([description, f"{qty:g}", money(price), money(amount)])
    pdf.ln(4)
    lines = [("Subtotal", subtotal)] + ([(f"Tax ({float(data.get('tax_percent')):g}%)", tax)] if tax else []) + [("Total", total)]
    for label, value in lines:
        pdf.set_font(family, "B" if label == "Total" else "", 12 if label == "Total" else 10)
        pdf.cell(150, 7, label, align="R")
        pdf.cell(40, 7, money(value), align="R", new_x="LMARGIN", new_y="NEXT")
    if data.get("notes"):
        pdf.ln(6)
        pdf.set_font(family, "", 9)
        pdf.multi_cell(0, 5, str(data["notes"]))
    path = (out_dir or kit.output_dir()) / f"{number}.pdf"
    pdf.output(str(path))
    return path
