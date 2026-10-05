"""Design 2.0's helpers: taking the white background off a picture, slides
from an Excel table, a PDF's pages as a design, a design from a photo of a
sketch, rewriting words with AI, and spell checking a whole design.

Each is plain data in, design out, so the Design page, the chat commands
and the tests all use the same code.
"""

from __future__ import annotations

import base64
import bisect
import csv
import datetime as dt
import io
import shutil
from pathlib import Path

from PIL import Image, ImageFilter, ImageOps

from .. import kit
from . import chart, model, templates
from .ai import DesignAIError


def _failed(answer: str) -> bool:
    return not answer or answer.startswith(("No AI provider", "I couldn't reach", "Error", "Add a free"))


# --- remove a white background ------------------------------------------------------------------

def _border_connected(mask) -> "object":
    """The parts of a boolean mask joined to the picture's edge. Works on runs of pixels per row
    rather than pixel by pixel, so a 12-megapixel photo takes a fraction of a second."""
    import numpy as np

    height, width = mask.shape
    runs: list[tuple[list[int], list[int]]] = []
    for y in range(height):
        edges = np.diff(np.concatenate(([0], mask[y].astype(np.int8), [0])))
        runs.append((np.flatnonzero(edges == 1).tolist(), np.flatnonzero(edges == -1).tolist()))
    keep = [bytearray(len(starts)) for starts, _ in runs]
    stack = []
    for y, (starts, ends) in enumerate(runs):
        for i, (s, e) in enumerate(zip(starts, ends)):
            if y in (0, height - 1) or s == 0 or e == width:
                keep[y][i] = 1
                stack.append((y, i))
    while stack:
        y, i = stack.pop()
        s, e = runs[y][0][i], runs[y][1][i]
        for ny in (y - 1, y + 1):
            if not 0 <= ny < height:
                continue
            starts, ends = runs[ny]
            j = bisect.bisect_right(ends, s)
            while j < len(starts) and starts[j] < e:
                if not keep[ny][j]:
                    keep[ny][j] = 1
                    stack.append((ny, j))
                j += 1
    out = np.zeros_like(mask, dtype=bool)
    for y, (starts, ends) in enumerate(runs):
        for i, flag in enumerate(keep[y]):
            if flag:
                out[y, starts[i]:ends[i]] = True
    return out


def remove_white(image: Image.Image, tolerance: int = 28) -> Image.Image:
    """The white (or nearly white) background around a logo or a product photo made
    transparent. Only white that reaches the edge goes — the white of an eye or a letter's
    inside stays — and the rim is feathered so there's no white halo."""
    import numpy as np

    rgba = image.convert("RGBA")
    pixels = np.asarray(rgba)
    lightest = pixels[..., :3].min(axis=2)
    whiteish = (lightest >= 255 - tolerance) | (pixels[..., 3] == 0)
    background = _border_connected(whiteish)
    alpha = np.where(background, 0.0, 255.0)
    grown = np.asarray(Image.fromarray((background * 255).astype(np.uint8)).filter(ImageFilter.MaxFilter(5))) > 0
    rim = grown & ~background
    soft = np.clip((255.0 - lightest) / max(1.0, tolerance * 2.5), 0.0, 1.0) * 255.0
    alpha = np.where(rim, np.minimum(alpha, soft), alpha)
    alpha = np.minimum(alpha, pixels[..., 3]).astype(np.uint8)
    out = rgba.copy()
    out.putalpha(Image.fromarray(alpha).filter(ImageFilter.GaussianBlur(0.6)))
    return out


def remove_white_src(src: str, tolerance: int = 28) -> str:
    path = model.resolve_src(src)
    if path is None:
        raise model.DesignError("The picture's file is missing.")
    with Image.open(path) as image:
        image = ImageOps.exif_transpose(image)
        image.thumbnail((3000, 3000))
        return model.asset_image(remove_white(image, tolerance))


# --- slides from an Excel table -------------------------------------------------------------------

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "✓" if value else ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:,.2f}".rstrip("0").rstrip(".")
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def read_table(path: Path, limit: int = 400) -> list[list[str]]:
    """The first sheet of an .xlsx, or a .csv, as rows of text — header row first."""
    path = Path(path)
    suffix = path.suffix.lower()
    rows: list[list[str]] = []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise DesignAIError(f"Couldn't open {path.name}: {exc}") from None
        try:
            for row in book.active.iter_rows(values_only=True):
                rows.append([_cell(v) for v in row])
                if len(rows) > limit:
                    break
        finally:
            book.close()
    elif suffix in {".csv", ".tsv", ".txt"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        first = text.splitlines()[0] if text.strip() else ""
        delimiter = "\t" if "\t" in first else (";" if first.count(";") > first.count(",") else ",")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)][:limit + 1]
    else:
        raise DesignAIError("Choose an Excel (.xlsx) or CSV file.")
    rows = [r for r in rows if any(c for c in r)]
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    rows = [(r + [""] * width)[:width] for r in rows]
    if len(rows) < 2 or width < 1:
        raise DesignAIError("That sheet has no table in it — a header row, then rows of data.")
    return rows


def numeric_columns(rows: list[list[str]]) -> list[int]:
    body = rows[1:]
    out = []
    for c in range(1, len(rows[0])):
        filled = [r[c] for r in body if r[c]]
        if filled and sum(chart.number(v) is not None for v in filled) >= 0.6 * len(filled):
            out.append(c)
    return out


def _theme_table(el: dict, theme: dict) -> None:
    el.update(fill=theme["accent"], header_color=model.readable_on(theme["accent"]), stripe=theme["surface"],
              color=theme["body"], bg=theme["bg"], border=model.mix(theme["body"], theme["bg"], 0.75),
              font=templates._font(theme["body_font"]))


def slides_from_table(rows: list[list[str]], title: str, mode: str = "rows", theme_name: str = "midnight") -> dict:
    """mode 'rows': a slide per row (its first cell the title, the rest as 'Header: value').
    mode 'table': the table across slides, nine rows a slide. Either way, a chart slide too
    when there are numbers to chart."""
    theme_name = theme_name if theme_name in templates.SLIDE_THEMES else "midnight"
    theme = templates.SLIDE_THEMES[theme_name]
    header, body = rows[0], rows[1:]
    if mode == "rows":
        slides = [{"title": row[0] or f"{header[0] or 'Row'} {i}",
                   "bullets": [f"{header[c] or f'Column {c + 1}'}: {row[c]}" for c in range(1, len(row)) if row[c]]}
                  for i, row in enumerate(body[:40], 1)]
        design = templates.build_deck({"title": title, "subtitle": f"{len(body)} rows", "slides": slides}, theme_name)
    else:
        design = model.new_design("slides", "slides", title)
        design["pages"] = [templates.cover_page(design, title, f"{len(body)} rows · {len(header)} columns", theme)]
        x = templates._content_x(theme)
        W = design["w"]
        per = 9
        for start in range(0, min(len(body), 180), per):
            chunk = body[start:start + per]
            page = model.new_page(theme["bg"])
            page["layout"] = "content"
            part = f" ({start // per + 1})" if len(body) > per else ""
            model.add(design, page, model.text(f"{title}{part}", x, 64, W - x - 120, 150, role="title", size=68,
                                               bold=True, autofit=True, valign="bottom"))
            height = min(700.0, 78.0 * (len(chunk) + 1))
            model.add(design, page, model.table([header] + chunk, x, 280, W - x - 120, height, size=30,
                                                role="table"))
            design["pages"].append(page)
        templates.apply_theme(design, theme_name)
    numbers = numeric_columns(rows)
    if numbers and 2 <= len(body):
        picked = numbers[:3]
        data = [[header[0]] + [header[c] or f"Column {c + 1}" for c in picked]] + \
               [[r[0]] + [r[c] for c in picked] for r in body[:24]]
        page = model.new_page(theme["bg"])
        page["layout"] = "content"
        x = templates._content_x(theme)
        name = header[picked[0]] if len(picked) == 1 else ", ".join(header[c] for c in picked)
        model.add(design, page, model.text(f"{name}" + (f" by {header[0]}" if header[0] else ""), x, 64,
                                           design["w"] - x - 120, 150, role="title", size=68, bold=True,
                                           autofit=True, valign="bottom"))
        model.add(design, page, model.chart("bar" if len(data) <= 13 else "line", data, x, 270,
                                            design["w"] - x - 120, 740, role="chart", size=28,
                                            fill=theme["accent"], color=theme["body"],
                                            font=templates._font(theme["body_font"])))
        design["pages"].append(page)
        templates.decorate(design, page, theme)
        templates.style_text(page, theme)
    for page in design["pages"]:
        for el in page["elements"]:
            if el["type"] == "table":
                _theme_table(el, theme)
    return design


# --- a PDF's pages ------------------------------------------------------------------------------

def design_from_pdf(path: Path, limit: int = 60) -> dict:
    """Every page of a PDF as a picture on its own page (locked, so you can write and draw
    on top of it), the design shaped like the PDF."""
    from .. import winrt

    path = Path(path)
    if not winrt.available():
        raise DesignAIError("Opening a PDF's pages needs Windows 10 or 11.")
    try:
        count = winrt.pdf_page_count(path)
    except Exception as exc:
        raise DesignAIError(f"Couldn't open {path.name}: {exc}") from None
    if count < 1:
        raise DesignAIError(f"{path.name} has no pages.")
    folder = model.assets_dir() / f"pdf_{kit.stamp()}"
    try:
        pages = winrt.pdf_pages(path, folder, list(range(1, min(count, limit) + 1)), width=1920)
        if not pages:
            raise DesignAIError(f"Windows couldn't draw {path.name}'s pages.")
        ratio = pages[0]["width_pt"] / max(1.0, pages[0]["height_pt"])
        fmt = ("slides" if ratio >= 1.5 else "a4_landscape") if ratio > 1.05 else "poster"
        design = model.new_design("slides" if fmt == "slides" else fmt, fmt, path.stem, pages=0)
        design["pages"] = []
        design["w"] = model.FORMATS[fmt][1]
        design["h"] = int(round(design["w"] / ratio))
        for info in pages:
            page = model.new_page("#ffffff")
            page["layout"] = "content"
            model.add(design, page, model.picture(model.import_asset(info["path"]), 0, 0, design["w"], design["h"],
                                                  fit="contain", role="background", locked=True))
            design["pages"].append(page)
    except DesignAIError:
        raise
    except Exception as exc:
        raise DesignAIError(f"Couldn't read {path.name}: {exc}") from None
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    return design


# --- a design from a photo of a sketch --------------------------------------------------------------

SKETCH_PROMPT = """This photo is a hand-drawn sketch of a {label} layout. Turn it into a clean design, keeping
where things are in the sketch. {note}
JSON: {{"title": "short name", "bg": "#rrggbb", "palette": ["#rrggbb", "#rrggbb", "#rrggbb"],
 "elements": [{{"type": "text|shape|line|image|icon", "x": 0.1, "y": 0.1, "w": 0.5, "h": 0.1,
   "role": "title|subtitle|body|button|deco", "text": "...", "shape": "rect|round|ellipse|star|...",
   "fill": "#rrggbb", "color": "#rrggbb", "size": 0.06, "bold": true, "align": "left|center|right",
   "glyph": "one emoji, for icons"}}]}}
x, y, w, h are fractions of the page (0-1); size is the text height as a fraction of the page height.
A box with an X through it, or labelled photo/picture/image, is a picture (type image). Write the words that are
written in the sketch; where you can't read them, write sensible ones in the same language. Pick colours that
look good together, and keep text readable on what's behind it."""


def sketch_format(image: Image.Image) -> str:
    ratio = image.width / max(1, image.height)
    return "slides" if ratio > 1.25 else ("poster" if ratio < 0.8 else "instagram")


def build_from_sketch(plan: dict, fmt: str) -> dict:
    if not isinstance(plan, dict) or not isinstance(plan.get("elements"), list) or not plan["elements"]:
        raise DesignAIError("I couldn't make out a layout in that photo. Try a clearer, straight-on photo.")
    fmt = fmt if fmt in model.FORMATS else "slides"
    design = model.new_design("slides" if fmt == "slides" else fmt, fmt, str(plan.get("title") or "Sketch")[:80])
    W, H = design["w"], design["h"]
    page = design["pages"][0]
    page["bg"] = model.color(plan.get("bg"), "#ffffff")
    palette = [c for c in (model.color(p, None) for p in plan.get("palette") or []) if c]
    if palette:
        design["palette"] = palette[:5]
    accent = palette[0] if palette else "#2563eb"
    items = [i for i in plan["elements"] if isinstance(i, dict)][:60]
    fractions = all(float(i.get(k, 0) or 0) <= 1.01 for i in items for k in ("x", "y", "w", "h")
                    if isinstance(i.get(k, 0), (int, float)))

    def size_of(item, key, total):
        try:
            value = float(item.get(key) or 0)
        except (TypeError, ValueError):
            value = 0.0
        return value * total if fractions else value

    for item in items:
        kind = str(item.get("type") or "text").lower()
        x, y = size_of(item, "x", W), size_of(item, "y", H)
        w, h = max(8.0, size_of(item, "w", W)), max(8.0, size_of(item, "h", H))
        role = str(item.get("role") or "")
        try:
            if kind == "line":
                el = model.line(x, y, x + size_of(item, "w", W), y + size_of(item, "h", H),
                                stroke=model.color(item.get("color"), model.readable_on(page["bg"])),
                                stroke_w=max(3.0, H * 0.005), role="deco")
            elif kind == "image":
                el = model.picture("", x, y, w, h, role="image", radius=min(w, h) * 0.04)
            elif kind == "icon":
                el = model.icon(str(item.get("glyph") or "⭐")[:4], x, y, min(w, h), set="color", role="icon")
            elif kind == "shape":
                fill = model.color(item.get("fill"), accent)
                shape = str(item.get("shape") or "round").lower()
                text = str(item.get("text") or "")
                el = model.shape(shape if shape in model.SHAPES else "round", x, y, w, h, fill=fill, text=text,
                                 color=model.color(item.get("color"), model.readable_on(fill)),
                                 size=max(12.0, min(h * 0.45, H * 0.04)), bold=role == "button",
                                 radius=min(w, h) * 0.2, role=role or "deco")
            else:
                default = {"title": 0.08, "subtitle": 0.045, "button": 0.035}.get(role, 0.032)
                try:
                    size = float(item.get("size") or default)
                except (TypeError, ValueError):
                    size = default
                size = size * H if size <= 1 else size
                el = model.text(str(item.get("text") or "Text"), x, y, w, h, size=max(10.0, size),
                                color=model.color(item.get("color"), model.readable_on(page["bg"])),
                                bold=bool(item.get("bold", role in {"title", "button"})),
                                align=str(item.get("align") or "left"), autofit=True, role=role or "body")
        except (model.DesignError, ValueError, TypeError):
            continue
        model.add(design, page, el)
    if not page["elements"]:
        raise DesignAIError("I couldn't make out a layout in that photo. Try a clearer, straight-on photo.")
    return design


def design_from_sketch(brain, photo: Path, fmt: str = "auto", note: str = "") -> dict:
    try:
        with Image.open(photo) as image:
            image = ImageOps.exif_transpose(image).convert("RGB")
    except Exception as exc:
        raise DesignAIError(f"Couldn't open the photo: {exc}") from None
    if fmt not in model.FORMATS:
        fmt = sketch_format(image)
    image.thumbnail((1280, 1280))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    prompt = SKETCH_PROMPT.format(label=model.FORMATS[fmt][0], note=f"Also: {note}" if note else "")
    try:
        plan = kit.ask_json(brain, prompt, room=5000, image_b64=base64.b64encode(buffer.getvalue()).decode("ascii"))
    except kit.KitError as exc:
        raise DesignAIError(str(exc)) from None
    design = build_from_sketch(plan, fmt)
    from . import ai

    ai.polish(design, 0)
    return design


# --- rewriting words -----------------------------------------------------------------------------------

REWRITES = {"shorter": "Make it shorter", "longer": "Make it a little longer and more vivid",
            "formal": "Make it more formal and professional", "friendly": "Make it warmer and friendlier",
            "punchy": "Make it punchier, like a headline", "fix": "Fix the spelling, grammar and punctuation only",
            "tr": "Translate it into Turkish", "en": "Translate it into English"}
REWRITE_LABELS = {"shorter": "Shorter", "longer": "Longer", "formal": "More formal", "friendly": "Friendlier",
                  "punchy": "Punchier", "fix": "Fix grammar", "tr": "In Turkish", "en": "In English"}


def rewrite(brain, text: str, how: str) -> str:
    if not text.strip():
        raise DesignAIError("There are no words to rewrite.")
    instruction = REWRITES.get(how, how)
    answer = brain.ask_once(
        f"{instruction}. The text sits on a design (a slide, poster or card): keep it about as long as it is "
        "unless asked otherwise, keep its line breaks and bullet marks, and keep its language unless asked to "
        "translate. Reply with only the new text — no quotes, no notes.\n\nText:\n" + text)
    if _failed(answer):
        raise DesignAIError(answer.split("\n")[0])
    answer = answer.strip()
    if answer.startswith("```"):
        answer = answer.strip("`").split("\n", 1)[-1].strip()
    if len(answer) > 1 and answer[0] == answer[-1] and answer[0] in "\"'“”":
        answer = answer[1:-1].strip()
    return answer[:4000]


# --- spell checking ------------------------------------------------------------------------------------

def _fields(el: dict):
    if el["type"] in {"text", "shape"} and el.get("text", "").strip():
        yield "text", el["text"]
    if el["type"] == "chart" and el.get("title"):
        yield "title", el["title"]
    if el["type"] == "table":
        for r, row in enumerate(el["rows"]):
            for c, cell in enumerate(row):
                if cell.strip():
                    yield f"rows:{r}:{c}", cell


def _get(el: dict, field: str) -> str:
    if field.startswith("rows:"):
        _, r, c = field.split(":")
        return el["rows"][int(r)][int(c)]
    return el.get(field, "")


def _put(el: dict, field: str, value: str) -> None:
    if field.startswith("rows:"):
        _, r, c = field.split(":")
        el["rows"][int(r)][int(c)] = value
    else:
        el[field] = value


def spelling_issues(design: dict, pages: list[int] | None = None) -> list[dict]:
    """Every misspelling in the design: {page, id, field, start, length, word, suggestions}."""
    from .. import spelling

    out = []
    for index, page in enumerate(design["pages"]):
        if pages is not None and index not in pages:
            continue
        for el in page["elements"]:
            for field, text in _fields(el):
                for issue in spelling.check(text):
                    out.append({"page": index, "id": el["id"], "field": field, **issue})
    return out


def spelling_issues_ai(brain, design: dict) -> list[dict]:
    """The same, asked of the AI — for when Windows' spell checker isn't there."""
    texts = [{"id": f"{i}:{el['id']}:{field}", "text": text[:600]} for i, page in enumerate(design["pages"])
             for el in page["elements"] for field, text in _fields(el)][:120]
    if not texts:
        return []
    import json

    try:
        reply = kit.ask_json(brain, "Find spelling mistakes in these texts from a design (any language). Ignore "
                                    "names, brands and deliberate styling.\n" + json.dumps(texts, ensure_ascii=False) +
                             '\nJSON: {"fixes": [{"id": "...", "wrong": "the word as written", "right": "..."}]}',
                             room=3000)
    except kit.KitError as exc:
        raise DesignAIError(str(exc)) from None
    out = []
    for fix in (reply.get("fixes") if isinstance(reply, dict) else reply) or []:
        try:
            page, element_id, field = str(fix["id"]).split(":", 2)
            el = model.find(design["pages"][int(page)], element_id)
            text = _get(el, field) if el else ""
            start = text.find(str(fix["wrong"]))
        except (KeyError, ValueError, IndexError, TypeError, AttributeError):
            continue
        if start >= 0 and fix["wrong"] != fix.get("right"):
            out.append({"page": int(page), "id": element_id, "field": field, "start": start,
                        "length": len(fix["wrong"]), "word": fix["wrong"], "suggestions": [str(fix.get("right", ""))]})
    return out


def apply_fix(design: dict, issue: dict, replacement: str) -> bool:
    """Put one correction in. Earlier fixes in the same text may have moved the word, so it is
    found again if it isn't where it was."""
    try:
        el = model.find(design["pages"][issue["page"]], issue["id"])
    except IndexError:
        return False
    if el is None:
        return False
    text = _get(el, issue["field"])
    start, word = issue["start"], issue["word"]
    if text[start:start + len(word)] != word:
        start = text.find(word)
        if start < 0:
            return False
    end = start + len(word)
    if not replacement:                         # deleting a repeated word takes its space too
        if start > 0 and text[start - 1] == " ":
            start -= 1
        elif end < len(text) and text[end] == " ":
            end += 1
    _put(el, issue["field"], text[:start] + replacement + text[end:])
    return True
